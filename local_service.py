Exit code: 0
Wall time: 0.3 seconds
Total output lines: 2662
Output:
"""本机批量服务：按人员文件夹调用云端 OCR，自动回写并留存每页 OCR Excel。"""
from __future__ import annotations

import base64
import difflib
import json
import math
import re
import shutil
import time
import uuid
from collections import Counter, defaultdict
from copy import copy
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Annotated, Any
from urllib.parse import quote

import requests
from bs4 import BeautifulSoup
from fastapi import BackgroundTasks, FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from datetime import datetime
from nutrition_report import NutritionReportService

ROOT = Path(__file__).resolve().parent
JOB_ROOT = ROOT / "data" / "jobs"
JOB_ROOT.mkdir(parents=True, exist_ok=True)
TEMPLATE_ROOT = ROOT / "templates"
TEMPLATE_ROOT.mkdir(parents=True, exist_ok=True)
MEDICAL_TEMPLATE_ROOT = ROOT / "data" / "templates"
MEDICAL_TEMPLATE_ROOT.mkdir(parents=True, exist_ok=True)
LAST_MEDICAL_TEMPLATE = MEDICAL_TEMPLATE_ROOT / "medical_last_template.xlsx"
LAST_MEDICAL_TEMPLATE_META = MEDICAL_TEMPLATE_ROOT / "medical_last_template.json"
MEDICAL_NORMAL_RANGE_CONFIG = MEDICAL_TEMPLATE_ROOT / "medical_normal_ranges.json"
NUTRITION_TEMPLATE = TEMPLATE_ROOT / "食物频率调查_汇总模板.xlsx"
JOBS: dict[str, dict[str, Any]] = {}

app = FastAPI(title="本地 OCR 批量回写服务")
app.add_middleware(CORSMiddleware, allow_origins=["http://127.0.0.1:5173", "http://localhost:5173"], allow_methods=["*"], allow_headers=["*"])
app.mount("/files", StaticFiles(directory=JOB_ROOT), name="files")

NUTRITION_GENERAL_COLUMNS = ["人员文件夹", "受试者编号", "姓名", "调查日期", "访视号", "每日餐次", "每周在家吃饭天数", "早餐地点", "午餐地点", "晚餐地点", "每周户外日照天数", "每日户外日照时长(小时)", "晒太阳时段", "皮肤暴露部位", "去脂体重", "人工备注"]
NUTRITION_FOOD_COLUMNS = ["人员文件夹", "姓名", "食物编号", "食物名称", "平均每次食用量", "次数", "频率周期(请核对)", "是否不吃", "人工核对备注"]
NUTRITION_SUPPLEMENT_COLUMNS = ["人员文件夹", "姓名", "保健品种类", "保健品名称", "平均每次服用量", "次数", "频率周期(请核对)", "是否不吃", "备注"]
NUTRITION_PERIODS = ("每天", "每周", "每月", "每年", "不吃")
NUTRITION_FOOD_ITEMS = {
    "1.1": "米饭", "1.2": "大米粥", "2.1": "馒头", "2.2": "面包", "2.3": "面条", "2.4": "包子/饺子",
    "3.1": "玉米", "3.2": "小米", "3.3": "窝头", "4.1": "红薯", "4.2": "山药", "4.3": "芋头",
    "4.4": "土豆", "5.1": "油条", "5.2": "油饼", "6": "猪肉", "7.1": "牛肉", "7.2": "羊肉",
    "8": "鸡肉、鸭肉", "9": "内脏类", "10.1": "鱼", "10.2": "虾", "11": "鲜奶", "12": "奶粉",
    "13": "奶酪", "14": "酸奶", "15": "蛋类", "16": "豆腐", "17": "豆腐丝/千张/豆腐干",
    "18": "豆浆", "19": "新鲜蔬菜", "20": "糕点", "21": "新鲜水果", "22": "坚果",
    "23": "果汁饮料", "24": "咖啡", "25": "茶/茶饮料", "26": "其他饮料",
}
NUTRITION_SUPPLEMENT_ITEMS = ("维生素D", "钙", "鱼油", "酵素", "其他保健品(如蛋白粉)")
CLOUD_OCR_TIMEOUT = (30, 900)
CLOUD_OCR_MAX_ATTEMPTS = 4
CLOUD_OCR_RETRY_DELAYS = (2, 5, 10)


def ensure_nutrition_template() -> None:
    if NUTRITION_TEMPLATE.exists():
        # 迁移已在本机保存的旧默认模板：明细表不再导出 OCR 原始行，
        # 并在人员文件夹后增加姓名，方便跨表人工核对。
        workbook = load_workbook(NUTRITION_TEMPLATE)
        changed = False
        for sheet_name in ("食物频率明细", "营养保健品"):
            if sheet_name not in workbook.sheetnames:
                continue
            worksheet = workbook[sheet_name]
            headers = [str(worksheet.cell(1, column).value or "").strip() for column in range(1, worksheet.max_column + 1)]
            for column in reversed([index + 1 for index, header in enumerate(headers) if header == "OCR原始行"]):
                worksheet.delete_cols(column)
                changed = True
            headers = [str(worksheet.cell(1, column).value or "").strip() for column in range(1, worksheet.max_column + 1)]
            if "姓名" not in headers and "人员文件夹" in headers:
                folder_column = headers.index("人员文件夹") + 1
                insert_column = folder_column + 1
                worksheet.insert_cols(insert_column)
                for row in range(1, worksheet.max_row + 1):
                    worksheet.cell(row, insert_column)._style = copy(worksheet.cell(row, folder_column)._style)
                    if row == 1:
                        worksheet.cell(row, insert_column).value = "姓名"
                changed = True
            if changed:
                worksheet.auto_filter.ref = worksheet.dimensions
        if changed:
            workbook.save(NUTRITION_TEMPLATE)
        workbook.close()
        return
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = [("人员汇总", NUTRITION_GENERAL_COLUMNS), ("食物频率明细", NUTRITION_FOOD_COLUMNS), ("营养保健品", NUTRITION_SUPPLEMENT_COLUMNS)]
    for name, headers in sheets:
        worksheet = workbook.create_sheet(name)
        worksheet.append(headers)
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="167D73")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for index, header in enumerate(headers, start=1):
            worksheet.column_dimensions[chr(64 + min(index, 26))].width = max(14, min(28, len(header) + 5))
    note = workbook.create_sheet("填写说明")
    note.append(["说明"])
    note.append(["1. 本模板由系统自动生成。人员汇总是一人一行；食物频率明细与营养保健品是一人多行。"])
    note.append(["2. OCR 对手写数值、勾选项和频率所在列可能不稳定；“频率周期(请核对)”应在核对页人工确认。"])
    note.append(["3. 原始 OCR Excel 和 PDF 会按任务、人员分别留存，便于复核。"])
    note.column_dimensions["A"].width = 110
    workbook.save(NUTRITION_TEMPLATE)
    workbook.close()


ensure_nutrition_template()

ALIASES = {
    "姓名": ("姓名", "受检者", "患者姓名"), "年龄": ("年龄", "年龄岁"), "性别": ("性别",),
    "病历号": ("病历号", "门诊号", "住院号", "编号"), "报告日期": ("报告日期", "检验日期"),
    "转铁蛋白": ("转铁蛋白", "trf"), "25羟维生素d": ("25羟维生素d", "维生素d", "vd"), "类胰岛素样生长因子": ("类胰岛素样生长因子", "igf1"),
    "生长激素": ("生长激素", "hgh"), "游离三碘甲状腺原氨酸": ("游离三碘甲状腺原氨酸", "ft3"), "游离甲状腺素": ("游离甲状腺素", "ft4"),
    "促甲状腺素": ("促甲状腺素", "促甲状腺激素", "tsh"), "孕酮": ("孕酮",), "睾酮": ("睾酮",), "卵泡生成素": ("卵泡生成素", "fsh"),
    "雄烯二酮": ("雄烯二酮",), "抗缪勒管激素": ("抗缪勒管激素", "抗继勒管激素", "抗穆勒管激素", "amh"), "促黄体生成素": ("促黄体生成素", "lh"),
    "雌二醇": ("雌二醇", "e2"), "垂体泌乳素": ("垂体泌乳素", "prl"), "白细胞": ("白细胞", "白细胞计数", "wbc"),
    "红细胞": ("红细胞", "红细胞计数", "rbc"), "血红蛋白": ("血红蛋白", "hgb"), "红细胞压积": ("红细胞压积", "hct"),
    "平均红细胞体积": ("平均红细胞体积", "mcv"), "平均血红蛋白含量": ("平均血红蛋白含量", "mch"),
    "平均血红蛋白浓度": ("平均血红蛋白浓度", "mchc"), "血小板": ("血小板", "血小板计数", "plt"), "淋巴细胞百分数": ("淋巴细胞百分数", "lymph"),
    "嗜酸性粒细胞百分数": ("嗜酸性粒细胞百分数", "eo%", "eo"), "嗜碱性粒细胞百分数": ("嗜碱性粒细胞百分数", "baso%", "baso"),
    "卵巢储备评估得分": ("卵巢储备评估得分", "卵巢储备评分", "卵巢储备评估分数"),
    "推测卵巢储备减退年龄": ("推测卵巢储备减退年龄", "预测卵巢储备减退年龄", "卵巢储备减退年龄"),
    "推测围绝经期开始年龄": ("推测围绝经期开始年龄", "预测围绝经期开始年龄", "围绝经期开始年龄"),
    "卵巢功能等级": ("卵巢功能等级", "卵巢功能分级", "卵巢功能级别"),
}
SKIP_LABELS = {"序号", "项目", "检验项目", "检测项目", "结果", "检测结果", "参考范围", "单位", "指标"}
PROFILE_FIELDS = {"姓名", "年龄", "性别", "病历号", "报告日期"}

# 这些默认值来自当前项目同一检验体系的既有报告。性激素等会受月经周期、
# 年龄等影响的项目不预设范围，需由使用者在页面设置后才参与异常判断。
MEDICAL_NORMAL_RANGE_FIELDS = tuple(key for key in ALIASES if key not in PROFILE_FIELDS)
DEFAULT_MEDICAL_NORMAL_RANGES: dict[str, dict[str, float | None]] = {
    "转铁蛋白": {"min": 212, "max": 360},
    "25羟维生素d": {"min": 20, "max": None},
    "生长激素": {"min": None, "max": 8},
    "游离三碘甲状腺原氨酸": {"min": 2.3, "max": 4.2},
    "游离甲状腺素": {"min": 0.89, "max": 1.80},
    "促甲状腺素": {"min": 0.55, "max": 4.78},
    "雄烯二酮": {"min": 1.08, "max": 11.73},
    "抗缪勒管激素": {"min": 0.97, "max": 12.53},
    "垂体泌乳素": {"min": 3.34, "max": 26.72},
    "白细胞": {"min": 3.5, "max": 9.5},
    "红细胞": {"min": 3.8, "max": 5.1},
    "血红蛋白": {"min": 115, "max": 150},
    "红细胞压积": {"min": 0.35, "max": 0.45},
    "平均红细胞体积": {"min": 82, "max": 100},
    "平均血红蛋白含量": {"min": 27, "max": 34},
    "平均血红蛋白浓度": {"min": 316, "max": 354},
    "血小板": {"min": 125, "max": 350},
    "淋巴细胞百分数": {"min": 20, "max": 50},
    "嗜酸性粒细胞百分数": {"min": 0.4, "max": 8.0},
    "嗜碱性粒细胞百分数": {"min": 0, "max": 1},
}
METADATA_SOURCE_WORDS = (
    "采样", "送检", "签收", "接收", "登记", "审核", "审签", "打印", "时间", "日期", "科别", "科室",
    "病区", "病床", "床号", "标本", "样本", "诊断", "备注", "检验者", "审核者", "检测岗位", "申请医生",
)


def is_metadata_source(source: object, target: str) -> bool:
    """过滤医院报告页眉/页脚的时间、人员和样本信息，避免误写成检验结果。"""
    source_text = short_label(source)
    source_norm = normalized(source_text)
    if target == "报告日期" and source_norm in {normalized("报告日期"), normalized("检验日期")}:
        return False
    return any(word in source_text for word in METADATA_SOURCE_WORDS)


def result_value_is_safe(target: str, value: str) -> bool:
    """检验项目只允许数值/区间或少量明确的分级结果，杜绝英文缩写、时间被写入。"""
    text = clean(value)
    if not text:
        return False
    if target == "姓名":
        return bool(re.fullmatch(r"[\u4e00-\u9fff·]{2,12}", text))
    if target == "年龄":
        return bool(re.fullmatch(r"\d{1,3}", text))
    if target == "性别":
        return text in {"男", "女"}
    if target == "病历号":
        return bool(re.fullmatch(r"[A-Za-z0-9-]{3,32}", text))
    if target == "报告日期":
        return bool(re.fullmatch(r"\d{4}-\d{1,2}-\d{1,2}", text))
    if re.fullmatch(r"\d{4}[-/.]\d{1,2}[-/.]\d{1,2}(?:\s+\d{1,2}:\d{2}(?::\d{2})?)?", text):
        return False
    if re.fullmatch(r"\d{1,2}:\d{2}(?::\d{2})?", text):
        return False
    if re.fullmatch(r"[A-Za-z][A-Za-z0-9_#%.-]{0,20}", text):
        return "等级" in target and bool(re.fullmatch(r"[A-E]", text, re.I))
    return bool(re.fullmatch(r"[<>≤≥]?\s*\d+(?:\.\d+)?(?:\s*[-~至]\s*\d+(?:\.\d+)?)?%?", text))


def clean(value: object) -> str:
    return "" if value is None else str(value).strip().replace("\u3000", " ")


def normalized(value: object) -> str:
    text = clean(value).lower().replace("\n", "")
    text = re.sub(r"[（(][^）)]*[）)]", "", text)
    return re.sub(r"[\s:：_\-*★↑↓【】\[\]/]", "", text)


OCR_MEDICAL_LABEL_CORRECTIONS = {
    "25轻维生素d": "25羟维生素d",
    "抗缬勒管激素": "抗缪勒管激素",
    "抗继勒管激素": "抗缪勒管激素",
    "抗穆勒管激素": "抗缪勒管激素",
    "雌烯二酮": "雄烯二酮",
}
MEDICAL_TRAILING_CODES = (
    "lymph", "baso", "mchc", "igf1", "ft3", "ft4", "tsh", "trf", "hgh", "fsh", "amh",
    "wbc", "rbc", "hgb", "hct", "mcv", "mch", "plt", "prl", "vd", "e2", "eo",
)


def medical_label_variants(value: object) -> set[str]:
    """Generate stable medical-label forms for glued codes and recurring OCR glyph errors."""
    text = normalized(short_label(value))
    if not text:
        return set()
    for wrong, correct in OCR_MEDICAL_LABEL_CORRECTIONS.items():
        text = text.replace(wrong, correct)
    variants = {text}
    for suffix in MEDICAL_TRAILING_CODES:
        if text.endswith(suffix) and len(text) > len(suffix):
            variants.add(text[:-len(suffix)])
    # These reports commonly glue one-letter hormone codes directly to the Chinese name.
    if text == "孕酮p":
        variants.add("孕酮")
    if text == "睾酮t":
        variants.add("睾酮")
    if text == "游离三碘甲状腺原氨酸t3":
        variants.add("游离三碘甲状腺原氨酸")
    return {item for item in variants if item}


def medical_alias_keys(value: object) -> set[str]:
    variants = medical_label_variants(value)
    keys: set[str] = set()
    for canonical, aliases in ALIASES.items():
        alias_variants: set[str] = set()
        for alias in aliases:
            alias_variants.update(medical_label_variants(alias))
        if variants.intersection(alias_variants):
            keys.add(normalized(canonical))
    return keys


def canonical_medical_range_key(value: object) -> str:
    """Return the configured canonical field name for a template/OCR label."""
    candidates = medical_alias_keys(value)
    for field in MEDICAL_NORMAL_RANGE_FIELDS:
        if normalized(field) in candidates:
            return field
    return ""


def range_number(value: object) -> float | None:
    """Parse one user-entered range boundary; blanks intentionally disable a side."""
    if value in (None, ""):
        return None
    if isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def normalized_medical_normal_ranges(raw_ranges: object, *, strict: bool = False) -> dict[str, dict[str, float | None]]:
    """Validate the persisted/API range shape and keep only supported medical fields."""
    if not isinstance(raw_ranges, dict):
        if strict:
            raise HTTPException(status_code=400, detail="正常范围设置格式错误")
        return {}
    result: dict[str, dict[str, float | None]] = {}
    for raw_key, raw_range in raw_ranges.items():
        key = canonical_medical_range_key(raw_key) or clean(raw_key)
        if key not in MEDICAL_NORMAL_RANGE_FIELDS:
            if strict:
                raise HTTPException(status_code=400, detail=f"不支持设置“{clean(raw_key)}”的正常范围")
            continue
        if not isinstance(raw_range, dict):
            if strict:
                raise HTTPException(status_code=400, detail=f"{key} 的正常范围格式错误")
            continue
        lower_raw, upper_raw = raw_range.get("min"), raw_range.get("max")
        lower, upper = range_number(lower_raw), range_number(upper_raw)
        if lower_raw not in (None, "") and lower is None:
            if strict:
                raise HTTPException(status_code=400, detail=f"{key} 的下限必须是有限数字")
            continue
        if upper_raw not in (None, "") and upper is None:
            if strict:
                raise HTTPException(status_code=400, detail=f"{key} 的上限必须是有限数字")
            continue
        if lower is None and upper is None:
            continue
        if lower is not None and upper is not None and lower > upper:
            if strict:
                raise HTTPException(status_code=400, detail=f"{key} 的下限不能大于上限")
            continue
        result[key] = {"min": lower, "max": upper}
    return result


def load_medical_normal_ranges() -> dict[str, dict[str, float | None]]:
    """Load user settings once saved; otherwise expose verified project defaults."""
    if not MEDICAL_NORMAL_RANGE_CONFIG.exists():
        return {key: dict(value) for key, value in DEFAULT_MEDICAL_NORMAL_RANGES.items()}
    try:
        payload = json.loads(MEDICAL_NORMAL_RANGE_CONFIG.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {key: dict(value) for key, value in DEFAULT_MEDICAL_NORMAL_RANGES.items()}
    ranges = normalized_medical_normal_ranges(payload.get("ranges") if isinstance(payload, dict) else payload)
    return ranges


def save_medical_normal_ranges(ranges: dict[str, dict[str, float | None]]) -> None:
    """Persist the full setting atomically so future batches do not need OCR reference values."""
    payload = {
        "version": 1,
        "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "ranges": ranges,
    }
    temporary = MEDICAL_NORMAL_RANGE_CONFIG.with_suffix(".tmp")
    temporary.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(MEDICAL_NORMAL_RANGE_CONFIG)


def normal_range_text(value: dict[str, float | None] | None) -> str:
    if not value:
        return ""
    lower, upper = value.get("min"), value.get("max")
    def formatted(number: float | None) -> str:
        return "" if number is None else f"{number:g}"
    if lower is not None and upper is not None:
        return f"{formatted(lower)}–{formatted(upper)}"
    if lower is not None:
        return f"≥ {formatted(lower)}"
    if upper is not None:
        return f"≤ {formatted(upper)}"
    return ""


def measurement_for_range(value: object) -> tuple[str, float] | None:
    """Read a single numeric result, including a leading comparison operator when present."""
    text = normalize_measurement_text(value)
    match = re.fullmatch(r"\s*(<=|>=|<|>|≤|≥)?\s*(\d+(?:\.\d+)?)\s*%?\s*", text)
    if not match:
        return None
    return (match.group(1) or "", float(match.group(2)))


def medical_range_assessment(header: object, value: object, ranges: dict[str, dict[str, float | None]]) -> dict[str, Any]:
    """Assess only values that are definitely outside the configured numeric interval."""
    key = canonical_medical_range_key(header)
    reference = ranges.get(key) if key else None
    display = normal_range_text(reference)
    assessment: dict[str, Any] = {"range_key": key, "normal_range": display, "is_abnormal": False, "abnormal_reason": ""}
    measurement = measurement_for_range(value)
    if not reference or not measurement:
        return assessment
    operator, measured = measurement
    lower, upper = reference.get("min"), reference.get("max")
    # A comparison result can be indeterminate. Mark it only when every possible
    # value represented by the result lies outside the configured range.
    below = lower is not None and (
        measured < lower if operator in {"", "<="} else measured <= lower if operat…24583 tokens truncated…).hex
    job_dir = JOB_ROOT / job_id
    input_dir = job_dir / "input" / "reports"
    input_dir.mkdir(parents=True, exist_ok=True)
    groups: dict[str, list[Path]] = defaultdict(list)
    for index, pdf in enumerate(files):
        relative = relative_paths[index] or pdf.filename or f"report_{index + 1}.pdf"
        safe_path = safe_relative(relative, pdf.filename or f"report_{index + 1}.pdf")
        source_path = input_dir / safe_path
        source_path.parent.mkdir(parents=True, exist_ok=True)
        with source_path.open("wb") as target:
            shutil.copyfileobj(pdf.file, target)
        parent = str(safe_path.parent).replace("\\", "/")
        groups[parent if parent != "." else source_path.stem].append(source_path)
    if len(body_composition_files) != len(body_composition_relative_paths):
        raise HTTPException(status_code=400, detail="体成分报告与目录信息不一致，请重新选择文件夹")
    body_composition_dir = job_dir / "input" / "body_composition"
    body_composition_paths: list[Path] = []
    for index, pdf in enumerate(body_composition_files):
        relative = body_composition_relative_paths[index] or pdf.filename or f"body_composition_{index + 1}.pdf"
        safe_path = safe_relative(relative, pdf.filename or f"body_composition_{index + 1}.pdf")
        source_path = body_composition_dir / safe_path
        source_path.parent.mkdir(parents=True, exist_ok=True)
        with source_path.open("wb") as target:
            shutil.copyfileobj(pdf.file, target)
        body_composition_paths.append(source_path)
    template_path = None
    if template and template.filename:
        template_path = job_dir / "input" / Path(template.filename).name
        with template_path.open("wb") as target:
            shutil.copyfileobj(template.file, target)
    JOBS[job_id] = {"job_id": job_id, "kind": "nutrition", "parse_mode": parse_mode, "processing_mode": processing_mode, "template_path": template_path, "output_filename": timestamped_filename("食物频率调查_自动汇总"), "status": "queued", "job_dir": job_dir, "total_files": len(files) + len(body_composition_paths), "completed_files": 0, "current_file": "", "current_person": "", "message": "文件已上传到本机，等待开始识别…", "people": []}
    background_tasks.add_task(process_nutrition_job, job_id, ocr_url, dict(groups), parse_mode, processing_mode, body_composition_paths)
    return {"success": True, "job_id": job_id, "total_files": len(files) + len(body_composition_paths), "people_count": len(groups), "body_composition_count": len(body_composition_paths)}


@app.post("/nutrition/jobs/{job_id}/body-composition", status_code=202)
def append_body_composition_reports(
    job_id: str,
    background_tasks: BackgroundTasks,
    ocr_url: Annotated[str, Form()],
    people_json: Annotated[str, Form()],
    files: Annotated[list[UploadFile], File()],
    relative_paths: Annotated[list[str], Form()],
):
    """Append body-composition OCR to a completed nutrition job without rerunning food OCR."""
    state = JOBS.get(job_id) or recover_nutrition_job(job_id)
    if not state or state.get("kind") != "nutrition" or state.get("status") != "completed":
        raise HTTPException(status_code=409, detail="营养问卷任务尚未完成，无法追加体成分报告")
    if state.get("body_composition_status") == "processing":
        raise HTTPException(status_code=409, detail="体成分报告正在识别，请等待当前任务完成")
    if not files or len(files) != len(relative_paths):
        raise HTTPException(status_code=400, detail="体成分报告与目录信息不一致，请重新选择 PDF")
    try:
        people = json.loads(people_json)
    except json.JSONDecodeError as error:
        raise HTTPException(status_code=400, detail="当前人员数据格式不正确") from error
    if not isinstance(people, list) or not all(isinstance(person, dict) and clean(person.get("id")) for person in people):
        raise HTTPException(status_code=400, detail="当前人员数据不完整，无法保留人工核对结果")

    append_id = uuid.uuid4().hex
    input_dir = Path(state["job_dir"]) / "input" / "body_composition_append" / append_id
    source_paths: list[Path] = []
    for index, pdf in enumerate(files):
        relative = relative_paths[index] or pdf.filename or f"body_composition_{index + 1}.pdf"
        safe_path = safe_relative(relative, pdf.filename or f"body_composition_{index + 1}.pdf")
        source_path = input_dir / safe_path
        source_path.parent.mkdir(parents=True, exist_ok=True)
        with source_path.open("wb") as target:
            shutil.copyfileobj(pdf.file, target)
        source_paths.append(source_path)

    state["people"] = people
    state.update(
        body_composition_status="processing",
        body_composition_total_files=len(source_paths),
        body_composition_completed_files=0,
        body_composition_pending_matches=[],
        body_composition_assignments={},
        message=f"已保留当前食物问卷核对结果，正在追加识别 {len(source_paths)} 份体成分报告…",
    )
    persist_nutrition_job(job_id)
    background_tasks.add_task(process_appended_body_composition_reports, job_id, ocr_url, source_paths, append_id)
    return {"success": True, "body_composition_count": len(source_paths)}


@app.post("/nutrition/jobs/{job_id}/body-composition/apply-matches")
def apply_body_composition_matches(job_id: str, payload: dict[str, Any]):
    """Apply user-confirmed body-report pairings to the existing reviewed people."""
    state = JOBS.get(job_id) or recover_nutrition_job(job_id)
    if not state or state.get("kind") != "nutrition" or state.get("status") != "completed":
        raise HTTPException(status_code=409, detail="营养问卷任务不可用")
    if state.get("body_composition_status") not in {"awaiting_match", "completed"}:
        raise HTTPException(status_code=409, detail="当前体成分配对不可编辑")
    people = payload.get("people")
    assignments = payload.get("assignments")
    matches = payload.get("body_composition_matches", payload.get("body_composition_pending_matches", state.get("body_composition_match_history") or state.get("body_composition_pending_matches") or []))
    if not isinstance(people, list) or not isinstance(assignments, dict) or not isinstance(matches, list):
        raise HTTPException(status_code=400, detail="配对确认数据不完整")
    people_by_id = {clean(person.get("id")): person for person in people if isinstance(person, dict) and clean(person.get("id"))}
    if not people_by_id:
        raise HTTPException(status_code=400, detail="当前人员数据不完整")

    previous_history = [record for record in state.get("body_composition_match_history") or [] if isinstance(record, dict)]
    previously_assigned_people = {
        clean(record.get("assigned_person_id"))
        for record in previous_history
        if clean(record.get("assigned_person_id"))
    }
    previous_source_files = {
        clean(record.get("file_url"))
        for record in previous_history
        if clean(record.get("file_url"))
    }
    for person_id in previously_assigned_people:
        person = people_by_id.get(person_id)
        if person is None:
            continue
        person.setdefault("general", {}).pop("去脂体重", None)
        person["body_composition_files"] = [
            file_url for file_url in person.get("body_composition_files", [])
            if clean(file_url) not in previous_source_files
        ]

    applied, corrected_names, remaining, applied_records = 0, 0, [], []
    for record in matches:
        if not isinstance(record, dict):
            continue
        record_id = clean(record.get("id"))
        person_id = clean(assignments.get(record_id))
        person = people_by_id.get(person_id)
        fat_free_mass = clean(record.get("corrected_fat_free_mass")) or clean(record.get("fat_free_mass"))
        if person is None or not fat_free_mass:
            remaining.append(record)
            continue
        corrected_name = re.sub(r"\s+", "", clean(record.get("corrected_name")) or clean(record.get("name")))
        if corrected_name and person.setdefault("general", {}).get("姓名") != corrected_name:
            person["general"]["姓名"] = corrected_name
            corrected_names += 1
        person.setdefault("general", {})["去脂体重"] = fat_free_mass
        source_files = person.setdefault("body_composition_files", [])
        if record.get("file_url") and record["file_url"] not in source_files:
            source_files.append(record["file_url"])
        record["assigned_person_id"] = person_id
        record["match_status"] = "applied"
        applied_records.append(record)
        applied += 1

    remaining_assignments = {clean(record.get("id")): clean(assignments.get(clean(record.get("id")))) for record in remaining if clean(record.get("id"))}
    for record in remaining:
        record["assigned_person_id"] = ""
        record["match_status"] = "pending"
    current_ids = {clean(record.get("id")) for record in matches if isinstance(record, dict)}
    retained_history = [record for record in previous_history if clean(record.get("id")) not in current_ids]
    history = retained_history + applied_records + remaining
    state["people"] = list(people_by_id.values())
    state.update(
        body_composition_status="awaiting_match" if remaining else "completed",
        body_composition_pending_matches=remaining,
        body_composition_assignments=remaining_assignments,
        body_composition_match_history=history,
        body_composition_unmatched=[clean(record.get("filename")) or "未命名体成分报告" for record in remaining],
        message=f"已追加 {applied} 项去脂体重，并同步修正 {corrected_names} 名已配对人员的姓名。" + (f" 仍有 {len(remaining)} 份保留在黄色补录表。" if remaining else ""),
    )
    persist_nutrition_job(job_id)
    return {"success": True, "people": state["people"], "applied": applied, "pending_matches": remaining, "history": history, "assignments": remaining_assignments, "message": state["message"]}


@app.post("/nutrition/jobs/{job_id}/save-review")
def save_nutrition_review(job_id: str, payload: dict[str, Any]):
    state = JOBS.get(job_id) or recover_nutrition_job(job_id)
    if not state or state.get("kind") != "nutrition" or state.get("status") != "completed":
        raise HTTPException(status_code=409, detail="营养问卷任务尚未完成，无法保存")
    people = payload.get("people")
    if not isinstance(people, list):
        raise HTTPException(status_code=400, detail="缺少核对数据")
    pending_matches = payload.get("body_composition_pending_matches")
    assignments = payload.get("body_composition_assignments")
    if pending_matches is not None and not isinstance(pending_matches, list):
        raise HTTPException(status_code=400, detail="体成分配对草稿格式不正确")
    if assignments is not None and not isinstance(assignments, dict):
        raise HTTPException(status_code=400, detail="体成分配对选择格式不正确")
    write_output = payload.get("write_output", True)
    if not isinstance(write_output, bool):
        raise HTTPException(status_code=400, detail="保存模式格式不正确")
    if write_output:
        write_nutrition_workbook(state, people)
    state["people"] = people
    if pending_matches is not None:
        state["body_composition_pending_matches"] = pending_matches
        history_by_id = {
            clean(record.get("id")): record
            for record in state.get("body_composition_match_history") or []
            if isinstance(record, dict) and clean(record.get("id"))
        }
        for record in pending_matches:
            if isinstance(record, dict) and clean(record.get("id")):
                history_by_id[clean(record.get("id"))] = record
        if history_by_id:
            state["body_composition_match_history"] = list(history_by_id.values())
    if assignments is not None:
        state["body_composition_assignments"] = {clean(key): clean(value) for key, value in assignments.items() if clean(key)}
    persist_nutrition_job(job_id)
    return {"success": True, "message": "食物频率调查汇总表已保存" if write_output else "当前核对信息已保存", "output_url": state["output_url"]}


@app.get("/nutrition/latest-job")
def latest_nutrition_job():
    candidates = sorted((item for item in JOB_ROOT.iterdir() if item.is_dir()), key=lambda item: item.stat().st_mtime, reverse=True)
    for job_dir in candidates:
        state = JOBS.get(job_dir.name) or recover_nutrition_job(job_dir.name)
        if state and state.get("kind") == "nutrition" and state.get("status") == "completed":
            restore_unmatched_body_composition_records(state)
            ensure_body_composition_match_history(state)
            snapshot = {key: value for key, value in state.items() if key not in {"job_dir", "output_relative", "template_path"}}
            return {"success": True, "job_id": job_dir.name, **snapshot}
    raise HTTPException(status_code=404, detail="没有可恢复的本地营养任务")


@app.post("/nutrition/preview")
def preview_nutrition_person(payload: dict[str, Any]):
    person = payload.get("person")
    user_overrides = payload.get("user_overrides")
    if not isinstance(person, dict):
        raise HTTPException(status_code=400, detail="缺少当前受访者的核对数据")
    if user_overrides is not None and not isinstance(user_overrides, dict):
        raise HTTPException(status_code=400, detail="人员补充信息格式不正确")
    try:
        service = NutritionReportService()
        adapted = service.adapt_ocr_person(person, user_overrides=user_overrides, report_date=clean(payload.get("report_date")) or None)
        # 预览返回与导出相同的机器初稿（分类评价、建议、能量/钙评价），
        # 前端据此提供人工复核，而不是把固定阈值直接当最终结论。
        preview = service.build_report_data(adapted.payload)
    except Exception as error:
        raise HTTPException(status_code=500, detail=f"营养摄入预览计算失败：{error}") from error
    return {
        "success": True, "preview": preview, "warnings": adapted.warnings,
        "usable_food_count": adapted.usable_food_count,
        "food_data_source": preview.get("food_data_source", {}),
    }


@app.post("/nutrition/jobs/{job_id}/people/{person_id}/report")
def generate_nutrition_person_report(job_id: str, person_id: str, payload: dict[str, Any]):
    """Generate one migrated nutrition report from the reviewed OCR person data.

    The report is deliberately generated on the local OCR service: the cloud server
    remains responsible only for OCR and no reviewed questionnaire data is uploaded
    again.
    """
    state = JOBS.get(job_id) or recover_nutrition_job(job_id)
    if not state or state.get("kind") != "nutrition" or state.get("status") != "completed":
        raise HTTPException(status_code=409, detail="营养问卷任务尚未完成，无法生成报告")

    person = payload.get("person")
    if not isinstance(person, dict):
        raise HTTPException(status_code=400, detail="缺少当前受访者的核对数据")
    if clean(person.get("id")) != person_id:
        raise HTTPException(status_code=400, detail="受访者标识与当前任务不一致")
    if person_id not in {clean(item.get("id")) for item in state.get("people", []) if isinstance(item, dict)}:
        raise HTTPException(status_code=404, detail="当前任务中不存在该受访者")

    report_format = clean(payload.get("format")).lower()
    if report_format not in {"pdf", "excel"}:
        raise HTTPException(status_code=400, detail="报告格式仅支持 pdf 或 excel")
    user_overrides = payload.get("user_overrides")
    if user_overrides is not None and not isinstance(user_overrides, dict):
        raise HTTPException(status_code=400, detail="人员补充信息格式不正确")
    manual_evaluations = payload.get("manual_evaluations")
    if manual_evaluations is not None and not isinstance(manual_evaluations, dict):
        raise HTTPException(status_code=400, detail="人工营养评价格式不正确")
    auto_evaluation = payload.get("auto_evaluation") is not False
    if not auto_evaluation and manual_evaluations is None:
        raise HTTPException(status_code=400, detail="请先提交人工复核后的营养评价")

    report_date = clean(payload.get("report_date")) or None
    extension = ".pdf" if report_format == "pdf" else ".xlsx"
    label = "营养健康评估报告" if report_format == "pdf" else "营养分析报告"
    filename_person = person
    override_name = clean((user_overrides or {}).get("name"))
    if override_name:
        filename_person = {
            **person,
            "general": {**(person.get("general") or {}), "姓名": override_name},
        }
    # Timestamp plus a short task-local token keeps simultaneous exports from
    # different browser tabs from overwriting one another.
    output_path = state["job_dir"] / "output" / "nutrition_reports" / timestamped_filename(
        f"{report_filename_stem(filename_person)}_{label}_{uuid.uuid4().hex[:8]}", extension
    )
    try:
        service = NutritionReportService()
        if report_format == "pdf":
            result = service.generate_pdf_from_ocr_person(
                person, output_path, user_overrides=user_overrides, report_date=report_date,
                auto_evaluation=auto_evaluation, manual_evaluations=manual_evaluations,
            )
            media_type = "application/pdf"
        else:
            result = service.generate_excel_from_ocr_person(
                person, output_path, user_overrides=user_overrides, report_date=report_date,
                auto_evaluation=auto_evaluation, manual_evaluations=manual_evaluations,
            )
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    except Exception as error:
        raise HTTPException(status_code=500, detail=f"营养报告生成失败：{error}") from error

    source = result.get("report_data", {}).get("food_data_source", {})
    reports = state.setdefault("nutrition_reports", {})
    reports.setdefault(person_id, []).append({
        "format": report_format,
        "path": output_path.relative_to(state["job_dir"]).as_posix(),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "usable_food_count": result["usable_food_count"],
        "warning_count": len(result["warnings"]),
        "food_data_source": source,
        "evaluation_mode": "automatic" if auto_evaluation else "manual_reviewed",
    })
    return FileResponse(
        output_path,
        media_type=media_type,
        filename=output_path.name,
        headers={
            "X-Nutrition-Data-Source": str(source.get("kind") or "unknown"),
            "X-Nutrition-Cloud-Available": str(source.get("cloud_available") is True).lower(),
            "X-Nutrition-Queried-At": str(source.get("queried_at") or ""),
            "X-Report-Filename": quote(output_path.name, safe=""),
        },
    )


@app.post("/nutrition/jobs/{job_id}/report")
def generate_nutrition_report(job_id: str, payload: dict[str, Any]):
    """Slash-safe report endpoint: the person id stays in JSON instead of the URL path."""
    person = payload.get("person")
    if not isinstance(person, dict):
        raise HTTPException(status_code=400, detail="缺少当前受访者的核对数据")
    person_id = clean(person.get("id"))
    if not person_id:
        raise HTTPException(status_code=400, detail="当前受访者缺少人员标识")
    return generate_nutrition_person_report(job_id, person_id, payload)


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8000)

