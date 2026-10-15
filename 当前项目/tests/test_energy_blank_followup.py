import json
from pathlib import Path
import tempfile
import unittest
import fitz
import openpyxl
from nutrition_report import NutritionReportService
from nutrition_report.adapter import build_payload_from_ocr_person, parse_quantity_grams
from nutrition_report.nutrition import NutritionService


class EnergyBlankFollowupTest(unittest.TestCase):
    def setUp(self):
        self.person={'id':'synthetic/energy-blank','general':{'姓名':'测试运动员','性别':'女','年龄':25,'体重':65,
            '项目种类':'女足','当前状态':'正常训练','运动量':'2'},'food_rows':[
            {'食物名称':'米饭','平均每次食用量':'100g','次数':'1','频率周期(请核对)':'每天'},
            {'食物名称':'鲜奶','平均每次食用量':'250g','次数':'2','频率周期(请核对)':'每天'},
            {'食物名称':'其他饮料','平均每次食用量':'','次数':'','频率周期(请核对)':''},
            {'食物名称':'奶酪','平均每次食用量':'','次数':'','频率周期(请核对)':''}]}
        self.service=NutritionReportService(prefer_cloud=False)

    def report(self, **overrides):
        return self.service.build_report_data(build_payload_from_ocr_person(self.person,user_overrides=overrides).payload)

    def test_existing_eer_recommendations_vary_with_activity(self):
        for activity,value in [('1',1700),('2',2100),('3',2450)]:
            row=self.report(activityLevel=activity)['nutrition']['energy']
            self.assertEqual(value,row['standard'])
            self.assertEqual('过少',row['evaluation'])
            self.assertIn('EER',row['standard_source'])
        row=self.report(activityLevel='')['nutrition']['energy']
        self.assertIsNone(row['standard'])
        self.assertEqual('-',row['evaluation'])
        self.assertEqual('请选择活动量',row['standard_note'])

    def test_blanks_and_omitted_foods_are_not_eaten(self):
        adapted=build_payload_from_ocr_person(self.person)
        self.assertEqual('not_eaten',adapted.payload.get('blank_food_policy'))
        blanks={r['food']:r for r in adapted.source_foods if r.get('status')=='not_eaten'}
        self.assertEqual(0,blanks['其他饮料']['daily_grams'])
        self.assertEqual(0,blanks['奶酪']['daily_grams'])
        report=self.report()
        self.assertTrue(report['food_standard_evaluation']['nutrient_record_complete'])
        self.assertEqual('不足',report['category_result']['坚果']['evaluation'])
        self.assertNotIn('记录不完整',json.dumps(report,ensure_ascii=False))
        self.assertNotIn('摄入小计',json.dumps(report,ensure_ascii=False))

    def test_workspace_copy_matches_completed_questionnaire_policy(self):
        source=(Path(__file__).resolve().parents[1]/'src/components/NutritionWorkspace.vue').read_text(encoding='utf8')
        self.assertFalse('缺失数据不会当作零摄入' in source)
        self.assertIn('食物摄入空白项按未食用处理',source)

    def test_blank_required_quantity_is_zero_not_an_incomplete_record(self):
        self.person['food_rows'][0]['平均每次食用量']=''
        report=self.report()
        self.assertEqual(0,report['category_result']['谷类']['value'])
        self.assertTrue(report['food_standard_evaluation']['nutrient_record_complete'])

    def test_nonblank_invalid_units_are_not_silently_replaced_with_zero(self):
        self.person['food_rows'][0]['平均每次食用量']='2碗'
        report=self.report()
        self.assertFalse(report['food_standard_evaluation']['nutrient_record_complete'])
        self.assertTrue(any(w['code']=='unsupported_quantity_unit' for w in report['conversion_warnings']))
        self.assertEqual('-',report['nutrition']['energy']['evaluation'])
        self.assertEqual(2100,report['nutrition']['energy']['standard'])
        self.assertNotIn('当前仅为已记录摄入小计',json.dumps(report,ensure_ascii=False))

    def test_trailing_decimal_before_explicit_unit_is_supported_not_guessed(self):
        self.assertEqual(50,parse_quantity_grams('50.g','米饭')[0])
        self.assertEqual(280,parse_quantity_grams('280.ml','果汁饮料')[0])
        self.assertIsNone(parse_quantity_grams('100/80.g','包子/饺子')[0])
        self.assertIsNone(parse_quantity_grams('5只g','虾')[0])

    def test_pdf_excel_labels_energy_and_previous_changes_match(self):
        report=self.report()
        with tempfile.TemporaryDirectory() as d:
            pdf=Path(d)/'report.pdf';xlsx=Path(d)/'report.xlsx'
            self.service.pdf_builder.build(report,pdf)
            self.service.excel_builder.build(report,xlsx)
            with fitz.open(pdf) as doc: text=''.join(p.get_text() for p in doc)
            compact=''.join(text.split())
            self.assertIn('鲜奶等价量（按蛋白质估计）',compact)
            self.assertIn('推荐能量2100',compact)
            self.assertNotIn('记录不完整',text)
            wb=openpyxl.load_workbook(xlsx,data_only=True)
            self.assertIn('按蛋白质估计',wb['分类评估']['D1'].value)
            categories={r[0]:r[2] for r in wb['分类评估'].iter_rows(min_row=2,values_only=True)}
            self.assertEqual('建议尽量不吃',categories['零食'])
            for category in ['饮料','酒类']:self.assertEqual('-',categories[category])
            self.assertEqual('不足',categories['谷类'])
            self.assertEqual('不足',categories['畜禽肉'])
            totals={r[0]:r for r in wb['营养总计'].iter_rows(min_row=2,values_only=True)}
            self.assertEqual(2100,totals['推荐能量'][1])
            self.assertEqual('中等强度',totals['活动量'][1])
            self.assertTrue(wb['建议']['B2'].alignment.wrap_text)
            self.assertGreaterEqual(wb['建议'].row_dimensions[2].height,30)
            self.assertTrue(all('记录不完整' not in str(c.value) for w in wb for row in w for c in row))
            wb.close()


if __name__=='__main__':unittest.main()
