import json
import re
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import fitz
import openpyxl
from nutrition_report import NutritionReportService
from nutrition_report.measurements import load_person_measurements, body_composition_rows, resting_metabolism_rows
from nutrition_report.builders.presentation import evaluation_rows


class MeasurementReportTest(unittest.TestCase):
    def test_approved_estimates_are_identity_bound_and_preserve_measurements(self):
        from copy import deepcopy
        from nutrition_report.measurements import apply_approved_estimates
        person={'id':'ocr/test','general':{'姓名':'测试甲'}}
        user={'projectType':'足球','gender':'female','weight':60,'fatFreeMass':''}
        record=dict(name='测试甲',person_id='ocr/test',project_type='足球',gender='female',approved=True,
                    values={'ffm':45,'REE':1300},method='同项目实测均值',sample_counts={'ffm':10,'REE':5})
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'estimates.json';p.write_text(json.dumps({'records':[record]}),encoding='utf-8')
            old=deepcopy(user)
            m,u=apply_approved_estimates(person,{},user,p)
            self.assertEqual(old,user)
            self.assertEqual(45,u['fatFreeMass']);self.assertTrue(u['fatFreeMassEstimated'])
            self.assertIn('ffm',m['body_composition']['estimated_fields'])
            self.assertIn('REE',m['resting_metabolism']['estimated_fields'])
            self.assertEqual(({},old),apply_approved_estimates({**person,'id':'other'},{},user,p))
            self.assertEqual(({},old),apply_approved_estimates({**person,'general':{'姓名':'测试乙'}},{},user,p))
            measured={'body_composition':{'values':{'ffm':47}},'resting_metabolism':{'values':{'REE':1450}}}
            self.assertEqual((measured,old),apply_approved_estimates(person,measured,user,p))
            m,u=apply_approved_estimates(person,{},dict(user,fatFreeMass=48,bmrMin=1200,bmrMax=1400),p)
            self.assertEqual({},m);self.assertEqual(48,u['fatFreeMass'])
            p.write_text(json.dumps({'records':[dict(record,approved=False)]}),encoding='utf-8')
            self.assertEqual(({},old),apply_approved_estimates(person,{},user,p))

    def test_estimated_ree_is_used_and_labeled_without_claiming_measurement(self):
        service=NutritionReportService(prefer_cloud=False)
        d=service.build_report_data({'user':{'name':'测试甲','age':22,'projectType':'足球','weight':60,
            'gender':'female','athleteType':'team','trainingPhase':'training','fatFreeMass':45,'fatFreeMassEstimated':True},
            'body_composition':{'values':{'ffm':45},'estimated_fields':{'ffm':{'method':'均值'}}},
            'resting_metabolism':{'values':{'REE':1300},'estimated_fields':{'REE':{'method':'均值'}}},
            'meals':[{'name':'米饭','amount':800}]})
        e=d['nutrition']['energy']
        self.assertEqual('同项目REE均值估算',e['resting_source'])
        self.assertNotIn('restingMeasuredRee',d['user'])
        self.assertEqual([1300,1300],e['resting_bounds'])
        self.assertNotEqual('missing',e['availability']['status'])
        self.assertIn('估算',body_composition_rows(d)[5][0])
        self.assertIn('估算',resting_metabolism_rows(d)[0][0])
        with tempfile.TemporaryDirectory() as tmp:
            for style in ['modern','legacy']:
                d['report_style']=style;p=Path(tmp)/(style+'.pdf');service.pdf_builder.build(d,p)
                with fitz.open(p) as doc:text=''.join(page.get_text() for page in doc)
                self.assertIn('45 kg（估算）',text)
                self.assertIn('1300.00 kcal/天（估算）',text)
            p=Path(tmp)/'report.xlsx';service.excel_builder.build(d,p)
            w=openpyxl.load_workbook(p);text=' '.join(str(c.value) for s in w for row in s for c in row);w.close()
            self.assertIn('45 kg（估算）',text);self.assertIn('静息能量消耗（REE）（估算）',text)

    def test_energy_equations_match_correlated_reference_endpoints(self):
        from nutrition_report.training_energy import training_expenditure
        from nutrition_report.energy_reference import energy_calculation_lines
        for sport in ('体操', '足球', '篮球'):
            training = training_expenditure({'projectType': sport, 'trainingPhase': 'training', 'weight': 60}, [1200, 1400])
            lines = energy_calculation_lines({'reference_type': 'bmr_plus_net', 'resting_bounds': [1200, 1400], 'training_expenditure': training})
            for i, line in enumerate(lines[1:3]):
                numbers = list(map(float, re.findall(r'\d+(?:\.\d+)?', line)))
                self.assertAlmostEqual(sum(numbers[:3]), numbers[3], delta=.011)
                self.assertAlmostEqual(numbers[3], [1200, 1400][i]+training['net_daily_kcal'][i], delta=.006)
        text = ' '.join(energy_calculation_lines({'reference_type': 'eer', 'standard': 2100}))
        self.assertIn('未另加运动消耗', text)
        self.assertNotIn('专项训练日均', text)

    def test_identity_requires_id_name_and_no_conflicting_subject(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp)/'registry.json'
            record = {'name': '测试甲', 'person_id': 'ocr/1', 'subject_id': '001',
                      'body_composition': {'values': {'ffm': 45.1}}}
            p.write_text(json.dumps({'records': [record]}), encoding='utf-8')
            person = {'id': 'ocr/1', 'general': {'姓名': '测试甲', '受试者编号': '001'}}
            self.assertEqual(record, load_person_measurements(person, p))
            for key, value in [('姓名', '测试乙'), ('受试者编号', '002')]:
                self.assertEqual({}, load_person_measurements({**person, 'general': {**person['general'], key: value}}, p))
            self.assertEqual({}, load_person_measurements({**person, 'id': 'ocr/2'}, p))
            p.write_text(json.dumps({'records': [record, record]}), encoding='utf-8')
            with self.assertRaises(ValueError):
                load_person_measurements(person, p)

    def test_missing_stays_blank_zero_stays_zero(self):
        self.assertTrue(all(r[1] is None for r in body_composition_rows({})))
        self.assertEqual([], resting_metabolism_rows({}))
        rows = resting_metabolism_rows({'resting_metabolism': {'values': {'PRO%': 0}}})
        self.assertEqual(0, rows[-1][1])
        self.assertIsNone(rows[0][1])

    def test_supplement_categories_preserve_manual_results_and_dedupe(self):
        data = {'category_result': {'奶类': {'value': 450}, '饮料': {'value': 20, 'evaluation': '手工原文'}},
                'food_standard_evaluation': {'rows': [{'key': 'milk', 'label': '奶类'}]}}
        rows = evaluation_rows(data)
        self.assertEqual(1, sum(r['key'] == 'milk' for r in rows))
        self.assertEqual('手工原文', next(r['evaluation'] for r in rows if r['label'] == '饮料'))
        self.assertIsNone(next(r['actual'] for r in rows if r['label'] == '运动营养食品'))

    def test_saved_registry_used_without_attachments_in_pdf_excel(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            record = {'name': '测试甲', 'person_id': 'ocr/1',
                      'body_composition': {'values': {'height': 164, 'date_of_birth': '2008-01-18', 'asmi': 6.21}},
                      'resting_metabolism': {'values': {'REE': 1708, 'FAT%': 66.1, 'CHO%': 33.9, 'PRO%': 0}}}
            p = root/'registry.json'
            p.write_text(json.dumps({'records': [record]}), encoding='utf-8')
            person = {'id': 'ocr/1', 'general': {'姓名': '测试甲', '年龄': 18, '性别': '女', '体重': 60}, 'food_rows': []}
            with patch.dict('os.environ', {'NUTRITION_MEASUREMENTS_PATH': str(p)}):
                service = NutritionReportService(prefer_cloud=False)
                payload = service.adapt_ocr_person(person).payload
            self.assertEqual(record['body_composition'], payload['body_composition'])
            data = service.build_report_data(payload)
            for style in ('modern', 'legacy'):
                data['report_style'] = style
                pdf = root/(style+'.pdf')
                service.pdf_builder.build(data, pdf)
                with fitz.open(pdf) as doc:
                    first = doc[0].get_text()
                    text = '\n'.join(p.get_text() for p in doc)
                    self.assertIn('体成分结果', first)
                    self.assertNotIn('食物摄入量评估报告', first)
                    self.assertIn('1708', text)
                    self.assertIn('PRO%', text)
                    self.assertNotIn('实测值', text)
                    self.assertIn('ASMI', first)
            xlsx = root/'report.xlsx'
            service.excel_builder.build(data, xlsx)
            w = openpyxl.load_workbook(xlsx, data_only=True)
            self.assertNotIn('分类评估', w.sheetnames)
            self.assertIn('体成分结果', w.sheetnames)
            self.assertTrue(any(r[0] == '蛋白质消耗占比（PRO%）' and r[1] == 0 for r in w['营养总计'].values))
            metabolic_rows = [r for r in w['营养总计'].values if r[0] and any(k in str(r[0]) for k in ('REE', 'FAT%', 'CHO%', 'PRO%'))]
            self.assertEqual(4, len(metabolic_rows))
            self.assertEqual([1708, 66.1, 33.9, 0], [r[1] for r in metabolic_rows])
            self.assertTrue(all(r[3] is None for r in metabolic_rows))
            w.close()


if __name__ == '__main__':
    unittest.main()
