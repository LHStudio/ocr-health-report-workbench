from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import local_service as service  # noqa: E402


class NutritionExportColumnsTest(unittest.TestCase):
    def test_detail_sheets_remove_raw_ocr_and_add_name_after_folder(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template_path, output_path = root / "legacy_template.xlsx", root / "output.xlsx"
            workbook = Workbook()
            general = workbook.active
            general.title = "人员汇总"
            general.append(service.NUTRITION_GENERAL_COLUMNS)
            food = workbook.create_sheet("食物频率明细")
            food.append(["人员文件夹", "食物编号", "食物名称", "平均每次食用量", "次数", "频率周期(请核对)", "是否不吃", "OCR原始行", "人工核对备注"])
            supplement = workbook.create_sheet("营养保健品")
            supplement.append(["人员文件夹", "保健品种类", "保健品名称", "平均每次服用量", "次数", "频率周期(请核对)", "是否不吃", "备注"])
            workbook.save(template_path)
            workbook.close()

            person = {
                "id": "ocr/测试",
                "general": {"姓名": "测试甲"},
                "food_rows": [{"食物编号": "1.1", "食物名称": "米饭", "平均每次食用量": "100g", "次数": "2", "频率周期(请核对)": "每天", "OCR原始行": "不应导出", "人工核对备注": "人工复核备注"}],
                "supplement_rows": [{"保健品种类": "钙", "保健品名称": "钙片", "平均每次服用量": "1片", "次数": "1", "频率周期(请核对)": "每天", "备注": "按医嘱"}],
            }
            service.write_nutrition_workbook({"template_path": template_path, "job_dir": root, "output_filename": "unused.xlsx"}, [person], output_path)

            workbook = load_workbook(output_path, read_only=True, data_only=True)
            try:
                for sheet_name in ("食物频率明细", "营养保健品"):
                    worksheet = workbook[sheet_name]
                    headers = [worksheet.cell(1, column).value for column in range(1, worksheet.max_column + 1)]
                    self.assertEqual(["人员文件夹", "姓名"], headers[:2])
                    self.assertNotIn("OCR原始行", headers)
                    self.assertEqual("ocr/测试", worksheet.cell(2, 1).value)
                    self.assertEqual("测试甲", worksheet.cell(2, 2).value)
                self.assertEqual("人工复核备注", workbook["食物频率明细"].cell(2, 9).value)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
