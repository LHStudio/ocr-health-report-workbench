"""BMR intervals replace EER only when explicitly provided, across all exports."""
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path
import fitz
from openpyxl import load_workbook
from nutrition_report.nutrition import NutritionService
from nutrition_report.adapter import build_payload_from_ocr_person
from nutrition_report.builders.pdf import NutritionPDFBuilder
from nutrition_report.builders.excel import NutritionExcelBuilder


class BmrRangeTest(unittest.TestCase):
    def report(self, value=1300, **user):
        profile = dict(name='测试人员', gender='female', age=23, weight=54.6,
                       activityLevel='2', athleteType='team', trainingPhase='training')
        profile.update(user)
        payload = dict(user=profile, meals=[dict(name='主食', amount=100)])
        before = deepcopy(payload)
        result = NutritionService(food_db={'主食': dict(category='谷类', energy=value,
            protein=80, carbohydrate=200, fat=50, calcium=600)}).build_report_data(payload)
        self.assertEqual(before, payload)
        return result

    def test_bmr_inclusive_boundaries_not_eer_tolerance(self):
        for value, expected in [(1186.99,'过少'),(1187,'适宜'),(1300,'适宜'),(1369,'适宜'),(1369.01,'过多')]:
            with self.subTest(value=value):
                data = self.report(value, bmrMin='1187', bmrMax='1369')
                e = data['nutrition']['energy']
                self.assertEqual(expected, e['evaluation'])
                self.assertEqual((1187,1369), (e.get('standard_min'),e.get('standard_max')))
                self.assertEqual('bmr',e.get('reference_type'))
                self.assertNotIn('女性EER',e['standard_source'])
                for row in data['food_standard_evaluation']['rows']:
                    if row['key'] in {'beans_energy','tubers_energy','grains_remainder'}:
                        self.assertEqual('personal_target',row['evaluation_basis'])
                self.assertNotIn('参考目标 -', ''.join(data['suggestions']))

    def test_absent_keeps_eer_and_bmr_needs_no_gender_age_or_activity(self):
        self.assertEqual(2100,self.report()['nutrition']['energy']['standard'])
        self.assertEqual('适宜',self.report(bmrMin=1187,bmrMax=1369,gender='',age='',activityLevel='')['nutrition']['energy']['evaluation'])

    def test_invalid_or_partial_range_is_rejected_not_silently_eer(self):
        for lo,hi in [('',1369),(1187,''),(0,1369),(-1,1369),(1400,1300),('nan',1369),(1187,'inf'),('abc',1369)]:
            with self.subTest(lo=lo,hi=hi), self.assertRaisesRegex(ValueError,'BMR'):
                self.report(bmrMin=lo,bmrMax=hi)

    def test_adapter_and_missing_profile_warning(self):
        p=dict(id='P',general={'姓名':'测试','BMR最低值':'1187','BMR最高值':'1369'},food_rows=[])
        a=build_payload_from_ocr_person(p)
        self.assertEqual('1187',a.payload['user'].get('bmrMin'))
        self.assertFalse(any(w['code']=='missing_energy_profile' for w in a.warnings))

    def test_pdf_last_section_and_unique_date_with_quality_on_off_both_styles(self):
        with tempfile.TemporaryDirectory() as tmp:
            for style in ['modern','legacy']:
                for quality in [False,True]:
                    d=self.report(bmrMin=1187,bmrMax=1369)
                    d.update(report_style=style,include_quality_page=quality,generated_at='2026-10-19 12:34',suggestions=['测试建议一','测试建议二'])
                    for with_athlete in [True,False]:
                        if not with_athlete:
                            d.pop('athlete_evaluation',None); d.pop('food_standard_evaluation',None)
                        f=Path(tmp)/f'{style}_{quality}_{with_athlete}.pdf'
                        NutritionPDFBuilder().build(d,f)
                        with fitz.open(f) as pdf:
                            texts=[p.get_text() for p in pdf]
                            self.assertIn('整体评估及建议',texts[-1])
                            self.assertIn('报告生成日期',texts[-1])
                            self.assertEqual(1,sum(t.count('报告生成日期') for t in texts))
                            energy=next(t for t in texts if '总能量摄入量' in t)
                            self.assertIn('BMR',energy); self.assertNotIn('女性EER',energy)

    def test_excel_reference_and_last_suggestions_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            f=Path(tmp)/'report.xlsx'
            NutritionExcelBuilder().build(self.report(bmrMin=1187,bmrMax=1369),f)
            w=load_workbook(f)
            self.assertEqual('建议',w.sheetnames[-1])
            self.assertTrue(any('报告生成日期' in str(c.value) for row in w['建议'] for c in row))
            content=' '.join(str(c.value) for row in w['营养素摄入'] for c in row) if '营养素摄入' in w else ' '.join(str(c.value) for s in w for row in s for c in row)
            self.assertIn('1187',content); self.assertIn('1369',content)
            self.assertIn('BMR',content); self.assertNotIn('女性EER',content)

    def test_summary_export_adds_bmr_to_old_templates(self):
        import local_service as api
        with tempfile.TemporaryDirectory() as tmp:
            f=Path(tmp)/'summary.xlsx'
            api.write_nutrition_workbook(dict(template_path=api.NUTRITION_TEMPLATE),
                [dict(id='fixture',general={'姓名':'测试','BMR最低值':1187,'BMR最高值':1369})], f)
            w=load_workbook(f)
            rows=list(w['人员汇总'].values); headers=list(rows[0])
            self.assertIn('BMR最低值',headers)
            self.assertEqual(1369,rows[1][headers.index('BMR最高值')])

    def test_single_batch_and_preview_share_bmr_from_general(self):
        import local_service as api
        import zipfile
        from unittest.mock import patch
        from nutrition_report import NutritionReportService
        person=dict(id='fixture',general={'姓名':'测试','BMR最低值':1187,'BMR最高值':1369},food_rows=[{'食物名称':'米饭','平均每次食用量':'100g','次数':'1','频率周期(请核对)':'每天'}])
        with tempfile.TemporaryDirectory() as tmp, patch.object(api,'NutritionReportService',side_effect=lambda:NutritionReportService(prefer_cloud=False)):
            job='bmr-fixture';api.JOBS[job]=dict(kind='nutrition',status='completed',job_dir=Path(tmp),people=[person])
            try:
                p=api.preview_nutrition_person(dict(person=person))['preview']
                self.assertEqual('bmr',p['nutrition']['energy']['reference_type'])
                for style in ['modern','legacy']:
                    r=api.generate_nutrition_report(job,dict(person=person,format='pdf',report_style=style))
                    with fitz.open(r.path) as d:
                        self.assertIn('报告生成日期',d[-1].get_text())
                        self.assertNotIn('女性EER',''.join(p.get_text() for p in d))
                    r=api.generate_nutrition_batch_reports(job,dict(format='pdf',report_style=style))
                    with zipfile.ZipFile(r.path) as z:
                        with fitz.open(stream=z.read(next(n for n in z.namelist() if n.endswith('.pdf'))),filetype='pdf') as d:
                            self.assertIn('报告生成日期',d[-1].get_text())
                            self.assertNotIn('女性EER',''.join(p.get_text() for p in d))
            finally:api.JOBS.pop(job,None)


if __name__=='__main__': unittest.main()
