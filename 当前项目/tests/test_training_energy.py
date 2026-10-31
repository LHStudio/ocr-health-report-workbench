import unittest
from copy import deepcopy
from nutrition_report.nutrition import NutritionService
from nutrition_report.adapter import build_payload_from_ocr_person


class TrainingEnergyTest(unittest.TestCase):
    def report(self, **updates):
        user=dict(projectType='足球',weight=60,bmrMin=1300,bmrMax=1500,
                  gender='female',age=22,activityLevel='3',athleteType='team',trainingPhase='training')
        user.update(updates)
        payload=dict(user=user,meals=[dict(name='米饭',amount=100)])
        before=deepcopy(payload)
        result=NutritionService().build_report_data(payload)
        self.assertEqual(before,payload)
        return result

    def test_football_gross_and_net_and_daily_reference(self):
        e=self.report()['nutrition']['energy']
        self.assertEqual('tdee',e['reference_type'])
        t=e['training_expenditure']
        self.assertEqual([3780,6048],t['gross_weekly_kcal'])
        self.assertEqual([3024,5292],t['net_weekly_kcal'])
        self.assertAlmostEqual((1300+432+630)/.91,e['standard_min'])
        self.assertAlmostEqual((1500+756+630)/.91,e['standard_max'])
        self.assertEqual([1300,1500],e['resting_bounds'])

    def test_football_two_and_half_hours(self):
        e=self.report(trainingPlan={'weeklyMinutes':900})['nutrition']['energy']['training_expenditure']
        self.assertEqual([4725,7560],e['gross_weekly_kcal'])
        self.assertEqual([3780,6615],e['net_weekly_kcal'])

    def test_basketball_example_no_double_count(self):
        t=self.report(projectType='篮球',trainingPlan={'weeklyMinutes':1350,'strengthMinutes':90,
                'sportMET':6,'strengthMET':4.5})['nutrition']['energy']['training_expenditure']
        self.assertEqual([1260,1260],t['sport_weekly_minutes'])
        self.assertEqual([8363.25,8363.25],t['gross_weekly_kcal'])
        self.assertEqual([6945.75,6945.75],t['net_weekly_kcal'])

    def test_basketball_schedule_envelope(self):
        t=self.report(projectType='篮球')['nutrition']['energy']['training_expenditure']
        self.assertEqual([1260,1440],t['total_weekly_minutes'])
        self.assertEqual([48,180],t['strength_weekly_minutes'])
        self.assertEqual([5008.5,10483.2],t['net_weekly_kcal'])

    def test_missing_bmr_no_eer_plus_exercise_and_unsupported_sports(self):
        for user in [dict(bmrMin='',bmrMax=''),dict(projectType='排球'),
                     dict(trainingPhase='competition'),dict(currentStatus='康复期'),dict(weight='')]:
            with self.subTest(user=user):
                e=self.report(**user)['nutrition']['energy']
                self.assertNotEqual('tdee',e['reference_type'])
                self.assertTrue(e['training_expenditure']['reason'])

    def test_gymnastics_user_example_and_net_daily_target(self):
        e = self.report(projectType='体操')['nutrition']['energy']
        t = e['training_expenditure']
        self.assertEqual('tdee', e['reference_type'])
        self.assertEqual([720,720], t['total_weekly_minutes'])
        self.assertEqual([0,0], t['strength_weekly_minutes'])
        self.assertEqual([2872.8,2872.8], t['gross_weekly_kcal'])
        self.assertAlmostEqual(478.8, t['gross_weekly_kcal'][0]/6)
        self.assertEqual([2116.8,2116.8], t['net_weekly_kcal'])
        self.assertAlmostEqual(302.4, t['net_daily_kcal'][0])
        self.assertAlmostEqual((1602.4+630)/.91,e['standard_min'])
        self.assertAlmostEqual((1802.4+630)/.91,e['standard_max'])

    def test_gymnastics_preserves_overrides_missing_bmr_and_training_gates(self):
        e = self.report(projectType='体操',trainingPlan={'weeklyMinutes':0})['nutrition']['energy']
        self.assertEqual([(1300+630)/.91,(1500+630)/.91], [e['standard_min'],e['standard_max']])
        e = self.report(projectType='体操',weight=120)['nutrition']['energy']
        self.assertEqual([4233.6,4233.6], e['training_expenditure']['net_weekly_kcal'])
        for updates in [dict(bmrMin='',bmrMax=''),dict(trainingPhase='competition'),
                        dict(currentStatus='康复期'),dict(weight='')]:
            e = self.report(projectType='体操',**updates)['nutrition']['energy']
            self.assertNotEqual('tdee',e['reference_type'])

    def test_zero_attendance_and_invalid_plan(self):
        e=self.report(trainingPlan={'weeklyMinutes':0})['nutrition']['energy']
        self.assertEqual([(1300+630)/.91,(1500+630)/.91],[e['standard_min'],e['standard_max']])
        for plan in [{'weeklyMinutes':-1},{'weeklyMinutes':float('nan')},
                     {'weeklyMinutes':60,'strengthMinutes':61}, {'sportMET':0}, {'foo':2}]:
            with self.subTest(plan=plan),self.assertRaises(ValueError):self.report(trainingPlan=plan)

    def test_weight_scaling_and_inclusive_boundary_and_manual(self):
        a=self.report(weight=120)['nutrition']['energy']
        self.assertEqual([6048,10584],a['training_expenditure']['net_weekly_kcal'])
        from nutrition_report.energy_reference import energy_reference_text
        self.assertEqual('2595.6-3171.43',energy_reference_text(self.report()['nutrition']['energy']))
        fat=next(r for r in self.report()['athlete_evaluation']['rows'] if r['key']=='fat')
        self.assertAlmostEqual(((1300+432+630)/.91)*.2/9,fat['target_min'])

    def test_adapter_preserves_personal_override(self):
        p=dict(id='test',general={'姓名':'测试','训练计划':{'weeklyMinutes':600}})
        self.assertEqual({'weeklyMinutes':600},build_payload_from_ocr_person(p).payload['user']['trainingPlan'])

    def test_personal_minutes_survive_workbook_roundtrip(self):
        import tempfile
        from pathlib import Path
        from openpyxl import load_workbook
        import local_service as api
        p=dict(id='test',general={'姓名':'测试','每周训练总分钟':600,'每周力量分钟':0})
        adapted=build_payload_from_ocr_person(p)
        self.assertEqual({'weeklyMinutes':600,'strengthMinutes':0},adapted.payload['user']['trainingPlan'])
        with tempfile.TemporaryDirectory() as tmp:
            out=Path(tmp)/'summary.xlsx'
            api.write_nutrition_workbook(dict(template_path=api.NUTRITION_TEMPLATE),[p],out)
            w=load_workbook(out)
            rows=list(w['人员汇总'].values);headers=list(rows[0]);w.close()
            self.assertIn('每周训练总分钟',headers)
            self.assertEqual(600,rows[1][headers.index('每周训练总分钟')])
            self.assertEqual(0,rows[1][headers.index('每周力量分钟')])

    def test_reference_in_pdf_both_styles_and_excel(self):
        import tempfile
        from pathlib import Path
        import fitz
        from openpyxl import load_workbook
        from nutrition_report.builders.pdf import NutritionPDFBuilder
        from nutrition_report.builders.excel import NutritionExcelBuilder
        with tempfile.TemporaryDirectory() as tmp:
            d=self.report()
            for style in ['modern','legacy']:
                d['report_style']=style
                p=Path(tmp)/(style+'.pdf')
                NutritionPDFBuilder().build(d,p)
                with fitz.open(p) as pdf:text=''.join(page.get_text() for page in pdf)
                self.assertIn('2595.6-3171.43',text)
                self.assertIn('推荐能量（TDEE）：2595.6-3171.43',text)
                self.assertIn('非运动性活动产热（NEAT）：630.00',text)
                self.assertIn('食物热效应（TEF）：233.60–285.43',text)
                for formula in ('下限：', '上限：', '÷0.91', 'TDEE×9%'):
                    self.assertNotIn(formula,text)
                self.assertNotIn('女性EER',text)
            p=Path(tmp)/'report.xlsx';NutritionExcelBuilder().build(d,p)
            w=load_workbook(p);text=' '.join(str(c.value) for s in w for row in s for c in row);w.close()
            self.assertIn('2595.6-3171.43',text)
            self.assertIn('食物热效应（TEF）：233.60–285.43',text)
            self.assertNotIn('推荐能量计算',text)
            self.assertNotIn('÷0.91',text)

    def test_manual_energy_preserved_and_inclusive_bounds(self):
        user=dict(projectType='足球',weight=60,bmrMin=1300,bmrMax=1500,trainingPhase='training',athleteType='team',age=22)
        low,high=(1300+432+630)/.91,(1500+756+630)/.91
        for energy, expected in [(low,'适宜'),(high,'适宜'),(low-1,'适宜'),(high+1,'适宜'),
                                 (low*.9,'适宜'),(high*1.1,'适宜'),(low*.9-1,'不足'),(high*1.1+1,'偏多')]:
            service=NutritionService(food_db={'fixture':dict(energy=energy,protein=1,carbohydrate=1,fat=1,calcium=1,category='谷类')})
            d=service.build_report_data(dict(user=user,meals=[dict(name='fixture',amount=100)]))
            self.assertEqual(expected,d['nutrition']['energy']['evaluation'])
            self.assertAlmostEqual(low,d['nutrition']['energy']['standard_min'])
            self.assertAlmostEqual(high,d['nutrition']['energy']['standard_max'])
        d=service.build_report_data(dict(user=user,meals=[dict(name='fixture',amount=100)],
            autoEvaluation=False,manualEvaluations={'energyEvaluation':'人工保留'}))
        self.assertEqual('人工保留',d['nutrition']['energy']['evaluation'])

    def test_rest_day_still_has_zero_bounds_for_fat(self):
        e=self.report(trainingPlan={'weeklyMinutes':0})['nutrition']['energy']
        self.assertEqual([0,0],e['training_expenditure']['net_weekly_kcal'])

    def test_project_aliases_match_supported_sport(self):
        for alias,sport in [('女足','足球'),('男足','足球'),('足球队','足球'),('女篮','篮球'),('男篮','篮球'),('篮球队','篮球')]:
            self.assertEqual(self.report(projectType=sport)['nutrition']['energy']['standard'],
                             self.report(projectType=alias)['nutrition']['energy']['standard'])

    def test_nontraining_status_or_explicit_blank_blocks_default(self):
        for status in ['比赛期','赛期','不训练','休息期','非训练期']:
            self.assertNotEqual('tdee',self.report(currentStatus=status)['nutrition']['energy']['reference_type'])
        self.assertNotEqual('tdee',self.report(trainingPhase='',currentStatus='正常训练')['nutrition']['energy']['reference_type'])

    def test_override_plan_wins_general_minutes_and_zero_is_retained(self):
        p=dict(id='test',general={'每周训练总分钟':600,'每周力量分钟':60,'训练计划':{'sportMET':6}})
        for value in [900,0]:
            a=build_payload_from_ocr_person(p,user_overrides={'trainingPlan':{'weeklyMinutes':value,'strengthMinutes':0}})
            self.assertEqual(value,a.payload['user']['trainingPlan']['weeklyMinutes'])
            self.assertEqual(0,a.payload['user']['trainingPlan']['strengthMinutes'])

    def test_generation_date_does_not_create_a_date_only_page(self):
        import tempfile
        from pathlib import Path
        import fitz
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Spacer
        from nutrition_report.builders.pdf import GenerationDateFooter, NutritionPDFBuilder
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'filled.pdf';doc=BaseDocTemplate(str(p),pagesize=A4)
            doc.addPageTemplates([PageTemplate(id='full',frames=[Frame(0,0,*A4,
                leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)])])
            font=NutritionPDFBuilder().font_name
            doc.build([Spacer(1,A4[1]),GenerationDateFooter('报告生成日期：2026-10-20',font)])
            with fitz.open(p) as pdf:
                self.assertEqual(1,len(pdf))
                self.assertIn('报告生成日期',pdf[0].get_text())


if __name__=='__main__':unittest.main()
