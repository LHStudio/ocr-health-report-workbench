import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook

import local_service
from local_service import imaging_section, parse_imaging_fields, process_imaging_job, save_imaging_review, write_imaging_workbook


def imaging_payload(date_text: str = "2026年07月13日") -> str:
    return json.dumps(
        {
            "parsing_res_list": [
                {
                    "block_id": 0,
                    "block_label": "paragraph_title",
                    "block_content": "妇科超声诊断报告单",
                },
                {
                    "block_id": 1,
                    "block_label": "table",
                    "block_content": (
                        "<table>"
                        "<tr><td>病人ID</td><td>PID-2026-001</td><td>姓名</td><td>张三</td><td>性别</td><td>女</td><td>年龄</td><td>23岁</td></tr>"
                        "<tr><td colspan='6'>超声所见：子宫前位，内膜厚0.8cm。</td></tr>"
                        "<tr><td colspan='6'>超声诊断：子宫附件未见明显异常</td></tr>"
                        "</table>"
                    ),
                },
                {
                    "block_id": 2,
                    "block_label": "text",
                    "block_content": f"诊断医生：___ 审核医生：___ 时间：{date_text}",
                },
            ]
        },
        ensure_ascii=False,
    )


class ImagingParserTests(unittest.TestCase):
    def test_extracts_requested_ultrasound_fields(self):
        fields, needs_review = parse_imaging_fields("", imaging_payload(), "扫描件.pdf")

        self.assertEqual(fields["病人ID"], "PID-2026-001")
        self.assertEqual(fields["姓名"], "张三")
        self.assertEqual(fields["性别"], "女")
        self.assertEqual(fields["年龄"], "23")
        self.assertEqual(fields["超声所见"], "子宫前位，内膜厚0.8cm。")
        self.assertEqual(fields["超声诊断"], "子宫附件未见明显异常")
        self.assertEqual(fields["检查时间"], "2026-07-13")
        self.assertEqual(needs_review, [])

    def test_extracts_patient_id_from_markdown_when_json_has_only_an_unrelated_title(self):
        json_with_title_only = json.dumps(
            {"parsing_res_list": [{"block_id": 0, "block_label": "title", "block_content": "妇科超声诊断报告单"}]},
            ensure_ascii=False,
        )

        fields, _ = parse_imaging_fields("病人ID：MARKDOWN-2026", json_with_title_only, "扫描件.pdf")

        self.assertEqual(fields["病人ID"], "MARKDOWN-2026")

    def test_keeps_incomplete_date_visible_and_marks_review(self):
        fields, needs_review = parse_imaging_fields("", imaging_payload("2026年07月0日"), "扫描件.pdf")

        self.assertEqual(fields["检查时间"], "2026年07月0日")
        self.assertIn("检查时间", needs_review)

    def test_keeps_ocr_confused_date_visible_and_marks_review(self):
        fields, needs_review = parse_imaging_fields("", imaging_payload("2026年07月c7日"), "张羽彤.pdf")

        self.assertEqual(fields["检查时间"], "2026年07月c7日")
        self.assertIn("检查时间", needs_review)

    def test_report_title_is_not_mistaken_for_diagnosis(self):
        payload = json.dumps({"parsing_res_list": [{"block_id": 0, "block_content": "妇科超声诊断报告单"}]}, ensure_ascii=False)
        fields, needs_review = parse_imaging_fields("", payload, "张三.pdf")

        self.assertEqual(fields["超声诊断"], "")
        self.assertIn("超声诊断", needs_review)

    def test_extracts_sections_from_body_crop_text(self):
        body_text = "超声所见：\n子宫前位，内膜厚0.8cm。\n超声诊断：\n盆腔积液"
        findings = imaging_section(body_text, ("超声所见",), ("超声诊断", "备注"))
        diagnosis = imaging_section(body_text, ("超声诊断",), ("备注", "时间"))

        self.assertEqual(findings, "子宫前位，内膜厚0.8cm。")
        self.assertEqual(diagnosis, "盆腔积液")

    def test_exports_reviewable_excel(self):
        fields, _ = parse_imaging_fields("", imaging_payload(), "扫描件.pdf")
        records = [{"id": "imaging-001", "filename": "扫描件.pdf", "fields": fields}]
        with tempfile.TemporaryDirectory() as directory:
            state = {"job_dir": Path(directory), "output_relative": Path("output/result.xlsx")}
            output_path = write_imaging_workbook(state, records)
            workbook = load_workbook(output_path, read_only=True)
            try:
                worksheet = workbook["影像识别结果"]
                self.assertEqual([cell.value for cell in worksheet[1]], ["病人ID", "文件名", "姓名", "性别", "年龄", "超声所见", "超声诊断", "检查时间", "需核对字段"])
                self.assertEqual(worksheet.cell(2, 1).value, "PID-2026-001")
                self.assertEqual(worksheet.cell(2, 3).value, "张三")
                self.assertEqual(worksheet.cell(2, 8).value, "2026-07-13")
                self.assertIsNone(worksheet.cell(2, 9).value)
            finally:
                workbook.close()

    def test_imaging_job_processes_and_saves_review(self):
        response = Mock(ok=True, status_code=200)
        response.json.return_value = {"success": True, "markdown": "", "json_text": imaging_payload()}
        with tempfile.TemporaryDirectory() as directory, patch.object(local_service, "JOB_ROOT", Path(directory)), patch.object(local_service, "request_cloud_ocr", return_value=response):
            local_service.JOBS.clear()
            job_id = "a" * 32
            job_dir = Path(directory) / job_id
            source = job_dir / "input" / "扫描件.pdf"
            source.parent.mkdir(parents=True)
            source.write_bytes(b"%PDF-1.4 synthetic")
            local_service.JOBS[job_id] = {
                "job_id": job_id,
                "kind": "imaging",
                "processing_mode": "accurate",
                "status": "queued",
                "job_dir": job_dir,
                "output_relative": Path("output/result.xlsx"),
                "total_files": 1,
                "completed_files": 0,
                "records": [],
            }
            process_imaging_job(job_id, "http://ocr.example", [(source, "扫描件.pdf")], "accurate")
            state = local_service.JOBS[job_id]
            self.assertEqual(state["status"], "completed")
            self.assertEqual(state["records"][0]["fields"]["姓名"], "张三")

            saved = save_imaging_review(job_id, {"records": state["records"]})
            self.assertTrue(saved["success"])
            self.assertRegex(saved["output_filename"], r"^影像识别结果_\d{8}_\d{6}_\d{6}\.xlsx$")
            self.assertEqual(state["output_relative"].name, saved["output_filename"])
            self.assertEqual(len(list((job_dir / "output").glob("*.xlsx"))), 1)

    def test_imaging_job_uses_body_crop_when_layout_ocr_omits_sections(self):
        response = Mock(ok=True, status_code=200)
        header_only = json.dumps(
            {
                "parsing_res_list": [
                    {"block_id": 0, "block_content": "妇科超声诊断报告单"},
                    {"block_id": 1, "block_content": "<table><tr><td>姓名</td><td>张三</td><td>性别</td><td>女</td><td>年龄</td><td>23岁</td></tr></table>"},
                    {"block_id": 2, "block_content": "时间：2026年07月13日"},
                ]
            },
            ensure_ascii=False,
        )
        response.json.return_value = {"success": True, "markdown": "", "json_text": header_only}
        body_text = "超声所见：\n子宫前位，内膜厚0.8cm。\n超声诊断：\n盆腔积液"
        with tempfile.TemporaryDirectory() as directory, patch.object(local_service, "request_cloud_ocr", return_value=response), patch.object(local_service, "request_imaging_body_ocr", return_value=body_text):
            job_id = "b" * 32
            job_dir = Path(directory) / job_id
            source = job_dir / "input" / "扫描件.pdf"
            source.parent.mkdir(parents=True)
            source.write_bytes(b"%PDF-1.4 synthetic")
            local_service.JOBS[job_id] = {
                "job_id": job_id,
                "kind": "imaging",
                "processing_mode": "accurate",
                "status": "queued",
                "job_dir": job_dir,
                "output_relative": Path("output/result.xlsx"),
                "total_files": 1,
                "completed_files": 0,
                "records": [],
            }

            process_imaging_job(job_id, "http://ocr.example", [(source, "扫描件.pdf")], "accurate")

            record = local_service.JOBS[job_id]["records"][0]
            self.assertEqual(record["fields"]["超声所见"], "子宫前位，内膜厚0.8cm。")
            self.assertEqual(record["fields"]["超声诊断"], "盆腔积液")
            self.assertEqual(record["needs_review"], ["病人ID"])
            self.assertIn("body_text", record["raw_files"])

    def test_imaging_job_uses_upper_right_ocr_when_layout_ocr_omits_patient_id(self):
        response = Mock(ok=True, status_code=200)
        response.json.return_value = {"success": True, "markdown": "", "json_text": imaging_payload().replace("<td>病人ID</td><td>PID-2026-001</td>", "")}
        with tempfile.TemporaryDirectory() as directory, patch.object(local_service, "request_cloud_ocr", return_value=response), patch.object(local_service, "request_imaging_patient_id_ocr", return_value="病人ID：RIGHT-7788"):
            job_id = "c" * 32
            job_dir = Path(directory) / job_id
            source = job_dir / "input" / "扫描件.pdf"
            source.parent.mkdir(parents=True)
            source.write_bytes(b"%PDF-1.4 synthetic")
            local_service.JOBS[job_id] = {
                "job_id": job_id,
                "kind": "imaging",
                "processing_mode": "accurate",
                "status": "queued",
                "job_dir": job_dir,
                "output_relative": Path("output/result.xlsx"),
                "total_files": 1,
                "completed_files": 0,
                "records": [],
            }

            process_imaging_job(job_id, "http://ocr.example", [(source, "扫描件.pdf")], "accurate")

            record = local_service.JOBS[job_id]["records"][0]
            self.assertEqual(record["fields"]["病人ID"], "RIGHT-7788")
            self.assertIn("patient_id_text", record["raw_files"])

    def test_imaging_job_keeps_ambiguous_patient_id_for_manual_review(self):
        response = Mock(ok=True, status_code=200)
        response.json.return_value = {"success": True, "markdown": "", "json_text": imaging_payload().replace("<td>病人ID</td><td>PID-2026-001</td>", "")}
        with tempfile.TemporaryDirectory() as directory, patch.object(local_service, "request_cloud_ocr", return_value=response), patch.object(local_service, "request_imaging_patient_id_ocr", return_value="病人ID：12 ?8"):
            job_id = "d" * 32
            job_dir = Path(directory) / job_id
            source = job_dir / "input" / "扫描件.pdf"
            source.parent.mkdir(parents=True)
            source.write_bytes(b"%PDF-1.4 synthetic")
            local_service.JOBS[job_id] = {
                "job_id": job_id,
                "kind": "imaging",
                "processing_mode": "accurate",
                "status": "queued",
                "job_dir": job_dir,
                "output_relative": Path("output/result.xlsx"),
                "total_files": 1,
                "completed_files": 0,
                "records": [],
            }

            process_imaging_job(job_id, "http://ocr.example", [(source, "扫描件.pdf")], "accurate")

            record = local_service.JOBS[job_id]["records"][0]
            self.assertEqual(record["fields"]["病人ID"], "")
            self.assertIn("病人ID", record["needs_review"])
            self.assertIn("patient_id_text", record["raw_files"])

    def test_fast_imaging_job_uses_only_layout_ocr_when_fields_are_missing(self):
        response = Mock(ok=True, status_code=200)
        response.json.return_value = {
            "success": True,
            "markdown": "",
            "json_text": json.dumps({"parsing_res_list": [{"block_id": 0, "block_content": "妇科超声诊断报告单"}]}, ensure_ascii=False),
        }
        with tempfile.TemporaryDirectory() as directory, patch.object(local_service, "request_cloud_ocr", return_value=response), patch.object(local_service, "request_imaging_patient_id_ocr", side_effect=AssertionError("快速模式不应补充病人ID OCR")), patch.object(local_service, "request_imaging_body_ocr", side_effect=AssertionError("快速模式不应补充正文 OCR")):
            job_id = "e" * 32
            job_dir = Path(directory) / job_id
            source = job_dir / "input" / "扫描件.pdf"
            source.parent.mkdir(parents=True)
            source.write_bytes(b"%PDF-1.4 synthetic")
            local_service.JOBS[job_id] = {
                "job_id": job_id,
                "kind": "imaging",
                "processing_mode": "fast",
                "status": "queued",
                "job_dir": job_dir,
                "output_relative": Path("output/result.xlsx"),
                "total_files": 1,
                "completed_files": 0,
                "records": [],
            }

            process_imaging_job(job_id, "http://ocr.example", [(source, "扫描件.pdf")], "fast")

            record = local_service.JOBS[job_id]["records"][0]
            self.assertEqual(record["fields"]["病人ID"], "")
            self.assertIn("病人ID", record["needs_review"])
            self.assertNotIn("patient_id_text", record["raw_files"])
            self.assertNotIn("body_text", record["raw_files"])
            self.assertEqual(record["warnings"], [])


if __name__ == "__main__":
    unittest.main()
