import hashlib,importlib.util,sys,tempfile,unittest
import openpyxl
from pathlib import Path
from datetime import datetime
ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('push',ROOT/'scripts/push_clinical_reports.py')
p=importlib.util.module_from_spec(spec); spec.loader.exec_module(p)
class PushTests(unittest.TestCase):
 def test_dates(self):
  for value in ['2026/7/6','2026.7.6','2026年7月6日',' ２０２６年 ０７月 ０６日 ', '2026-07-06 08:09:10','2026-07-06T08:09:10+08:00']:
   self.assertEqual(p.report_date(value),'2026-07-06')
  self.assertEqual(p.report_date(datetime(2026,7,6)),'2026-07-06')
  self.assertEqual(p.report_date(46301),'2026-10-06')
  for value in ['',None,'2026-02-30','2026-13-0','26-7-6','2026年07月0日','2026年07月','2026.7',True,0,0.5,'2026-07-06 rubbish','2026-07-06 99:99:99']:
   with self.assertRaises(ValueError):p.report_date(value)
 def test_original_headers_project_only_six_push_fields(self):
  data={'病人ID':'original-id','文件名':'synthetic.pdf','姓名':'合成测试','性别':'女','年龄':21,'超声所见':'保留原文\n原样','超声诊断':'合成结论','检查时间':'2026.7.6','需核对字段':'保留'}
  for headers in [p.IMAGING_COLUMNS,['文件名']+p.IMAGING_COLUMNS,list(data),list(reversed(data))]:
   with tempfile.TemporaryDirectory() as d:
    path=Path(d)/'input.xlsx';wb=openpyxl.Workbook();wb.active.append(headers);wb.active.append([data[k] for k in headers]);wb.save(path);wb.close()
    before=hashlib.sha256(path.read_bytes()).hexdigest()
    items,rejected=p.prepare(path,'imaging')
    self.assertEqual(rejected,[]);self.assertEqual(len(items),1)
    self.assertEqual(items[0]['payload'],{k:('2026-07-06' if k=='检查时间' else data[k]) for k in p.IMAGING_COLUMNS})
    self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)
 def test_bad_headers_and_unknown_days(self):
  with tempfile.TemporaryDirectory() as d:
   path=Path(d)/'input.xlsx'
   for headers in [p.IMAGING_COLUMNS+['姓名'],p.IMAGING_COLUMNS[:-1],p.IMAGING_COLUMNS+['未知字段']]:
    wb=openpyxl.Workbook();wb.active.append(headers);wb.active.append(['synthetic']*len(headers));wb.save(path);wb.close()
    with self.assertRaises(ValueError):p.prepare(path,'imaging')
   wb=openpyxl.Workbook();wb.active.append(['文件名']+p.IMAGING_COLUMNS)
   for value in ['2026.7','2026年07月0日',None]:wb.active.append(['synthetic.pdf','合成测试','女',21,'原文','结论',value])
   wb.save(path);wb.close();items,rejected=p.prepare(path,'imaging')
   self.assertEqual(items,[]);self.assertEqual(len(rejected),3)
   self.assertTrue(all(r['name']=='合成测试' for r in rejected))
 def test_endpoints(self):
  self.assertEqual(p.endpoint_url('https://example.test/api','lab'),'https://example.test/api/open/lab')
  for url in ['http://evil.test','https://example.test/other','https://user:pass@example.test','https://example.test/?x=1']:
   with self.assertRaises(ValueError):p.endpoint_url(url,'imaging')
 def test_formula_literal(self):
  self.assertEqual(p.cell_value('="002316891900"','病历号'),'002316891900')
  with self.assertRaises(ValueError):p.cell_value('=1+1','TRF')
 def test_receipt(self):
  item={'name':'合成测试','analysisDate':'2026-07-06'}
  body={'code':0,'data':dict(item,athleteId='a',analysisId='b',created=True)}
  self.assertTrue(p.valid_receipt(body,item))
  body['data']['name']='错误人员'
  self.assertFalse(p.valid_receipt(body,item))
 def test_checkpoint(self):
  item={'name':'合成测试','analysisDate':'2026-07-06','payload':{'姓名':'合成测试','检查时间':'2026-07-06','超声诊断':'合成'},'source':'test.xlsx','source_sha256':'fake','row':2}
  response={'code':0,'data':{'name':item['name'],'analysisDate':item['analysisDate'],'athleteId':'a','analysisId':'b','created':True}}
  calls=[]
  def send(*args):calls.append(args);return 200,response
  with tempfile.TemporaryDirectory() as d:
   self.assertEqual(p.execute([item],'https://example.test/api/open/imaging','secret',Path(d),send=send)['success'],1)
   self.assertEqual(p.execute([item],'https://example.test/api/open/imaging','secret',Path(d),send=send)['success'],1)
   self.assertEqual(len(calls),1)
  with tempfile.TemporaryDirectory() as d:
   def uncertain(*args):return 502,{'code':1}
   self.assertEqual(p.execute([item],'https://example.test/api/open/imaging','secret',Path(d),send=uncertain)['uncertain'],1)
   self.assertEqual(p.execute([item],'https://example.test/api/open/imaging','secret',Path(d),send=send)['uncertain'],1)
   self.assertEqual(len(calls),1)
if __name__=='__main__':unittest.main()
