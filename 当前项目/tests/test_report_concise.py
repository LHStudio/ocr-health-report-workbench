import re
import tempfile
import unittest
from pathlib import Path

import fitz
from nutrition_report import NutritionReportService
from nutrition_report.builders.presentation import evaluation_rows, evaluation_warnings


class ConciseReportTest(unittest.TestCase):
    def setUp(self):
        self.service = NutritionReportService(prefer_cloud=False)
        self.payload = dict(user=dict(name='验证样例', gender='female', age=22, weight=65,
            activityLevel='2', athleteType='team', trainingPhase='training'),
            meals=[dict(name='米饭', amount=100), dict(name='牛肉', amount=20),
                   dict(name='鲜奶', amount=100), dict(name='蛋类', amount=20)])

    def test_outward_wording_is_concise_but_internal_basis_is_retained(self):
        data = self.service.build_report_data(self.payload)
        suggestions = '\n'.join(data['suggestions'])
        self.assertTrue(data['suggestions'][0].startswith('主食总量目前约 100 g/天；'))
        for text in ['总量满足后再比较各类占比', '暂不比较结构', '加权蛋白质密度折算', '蛋白质优选参考']:
            self.assertNotIn(text, suggestions)
        rows = {r['key']: r for r in evaluation_rows(data)}
        for key in ['fat', 'vegetables', 'beans_energy', 'tubers_energy', 'grains_remainder',
                    'seafood', 'quality_protein', 'milk', 'eggs', 'livestock_remainder']:
            self.assertEqual('', rows[key]['note'], key)
        self.assertNotIn('加权蛋白质密度折算', rows['meat']['note'])
        raw = {r['key']: r for r in data['food_standard_evaluation']['rows']}
        self.assertIn('加权蛋白质密度折算', raw['meat']['note'])
        self.assertEqual('', rows['meat']['note'])
        self.assertIn('谷类蛋白质', raw['meat']['note'])

    def test_pdf_suggestion_numbers_and_units_never_wrap_apart(self):
        data = self.service.build_report_data(self.payload)
        # Exercise both the screenshot and awkward line endings independently of wording cleanup.
        units = ['2100 kcal/天', '85.3 g/天', '97.5–130 g/天', '873.2 mg/天']
        data['suggestions'] = [f"{'参考目标' * n} {value}；继续增加摄入。"
                               for n in range(1, 15) for value in units]
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'report.pdf'
            self.service.pdf_builder.build(data, path)
            with fitz.open(path) as doc:
                # Section boundaries, not removed running headers, identify advice pages.
                start = next(i for i, p in enumerate(doc) if '整体评估及建议' in p.get_text())
                end = next(i for i, p in enumerate(doc) if '能量及钙摄入量评估报告' in p.get_text())
                text = '\n'.join(doc[i].get_text() for i in range(start, end))
                self.assertFalse(re.search(r'\d\s*\n\s*(?:kcal|mg|g)/天', text), text)
                for value in units:
                    self.assertEqual(14, len(re.findall(re.escape(value).replace(r'\ ', r'[ \u00a0]*'), text)), value)

    def test_meat_sentence_and_protein_advice_are_the_requested_short_version(self):
        data = self.service.build_report_data(self.payload)
        self.assertEqual('肉类（含水产品）：目前 20 g/天，参考目标约 448.4 g/天，总蛋白质不足，可优先补充。', data['suggestions'][11])
        self.assertNotIn('达标后再考虑结构占比', '\n'.join(data['suggestions']))
        self.assertEqual('', next(r for r in evaluation_rows(data) if r['key'] == 'meat')['note'])
        raw = next(r for r in data['food_standard_evaluation']['rows'] if r['key'] == 'meat')
        self.assertAlmostEqual(448.442211, raw['target_grams'], places=5)
        self.assertAlmostEqual(428.442211, raw['shortfall_grams'], places=5)

    def test_chinese_protein_advice_fills_the_line_before_wrapping(self):
        data = self.service.build_report_data(self.payload)
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / 'report.pdf'; self.service.pdf_builder.build(data, pdf)
            with fitz.open(pdf) as doc:
                page = next(p for p in doc if '整体评估及建议' in p.get_text())
                lines = [l for b in page.get_text('dict')['blocks'] for l in b.get('lines', [])]
                first = next(l for l in lines if '先补足蛋白质' in ''.join(s['text'] for s in l['spans']))
                text = ''.join(s['text'] for s in first['spans'])
                self.assertIn('85.3 g/天', text)
                self.assertGreater(first['bbox'][2], page.rect.width - 116)
                for line in lines:
                    text = ''.join(s['text'] for s in line['spans']).strip()
                    self.assertFalse(text.startswith(('）', '，', '。', '；', '、')), text)

    def test_remove_two_warnings_and_include_generation_time_in_pdf_and_excel(self):
        import openpyxl
        from datetime import datetime, timezone, timedelta
        before = datetime.now(timezone(timedelta(hours=8))).replace(microsecond=0)
        data = self.service.build_report_data(self.payload)
        self.assertRegex(data.get('generated_at', ''), r'^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$')
        generated = datetime.strptime(data['generated_at'], '%Y-%m-%d %H:%M:%S').replace(tzinfo=before.tzinfo)
        self.assertGreaterEqual(generated, before)
        self.assertLessEqual(generated, datetime.now(before.tzinfo))
        self.assertEqual(data['generated_date'], data['generated_at'][:10])
        outward = '\n'.join(evaluation_warnings(data))
        for phrase in ['奶粉、奶酪折算鲜奶等价量', '蔬菜最低400g/天']:
            self.assertNotIn(phrase, outward)
            self.assertIn(phrase, '\n'.join(data['food_standard_evaluation']['warnings']))
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / 'report.pdf'; xlsx = Path(tmp) / 'report.xlsx'
            self.service.pdf_builder.build(data, pdf); self.service.excel_builder.build(data, xlsx)
            with fitz.open(pdf) as doc:
                page = next(p for p in doc if '运动员专项膳食评价' in p.get_text())
                self.assertIn('报告生成日期：' + data['generated_at'], page.get_text())
                self.assertNotIn('奶粉、奶酪折算鲜奶等价量', page.get_text())
                self.assertNotIn('蔬菜最低400g/天', page.get_text())
                footer = page.search_for(data['generated_at'])[0]
                self.assertGreater(footer.x0, page.rect.width / 2)
                self.assertGreater(footer.y0, page.rect.height * .8)
            wb = openpyxl.load_workbook(xlsx, data_only=True)
            self.assertEqual(data['generated_at'], wb['运动员专项评价'].cell(wb['运动员专项评价'].max_row, 2).value)
            wb.close()


if __name__ == '__main__':
    unittest.main()
