import copy
import importlib.util
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from nutrition_report.builders.excel import NutritionExcelBuilder


def fixture():
    return {
        'user': {'name': '测试运动员', 'weight': 60}, 'date': '2026-06-12',
        'category_result': {
            '奶类': {'value': 100, 'evaluation': '不足', 'equivalent_value': 400},
            '谷类': {'value': 0, 'evaluation': '不足', 'reference_equivalent': {'representative': '米饭', 'value': 120}},
        },
        'athlete_evaluation': {'rows': [dict(key='protein', label='蛋白质', actual=0, unit='g', target='10', gap=-10, evaluation='不足', note='内部口径保留在计算数据')], 'warnings': []},
        'nutrition': {'energy': {'value': 100, 'evaluation': '不足', 'standard': 200}, 'protein': 0, 'calcium': {'value': 20, 'evaluation': '不足'}},
        'suggestions': ['人工建议不改'], 'generated_at': '2026-10-22 10:08:05',
    }


class ExcelExportColumnTests(unittest.TestCase):
    def test_export_has_only_requested_columns_and_does_not_mutate_input(self):
        data = fixture()
        original = copy.deepcopy(data)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'report.xlsx'
            NutritionExcelBuilder().build(data, path)
            w = load_workbook(path)
            self.assertEqual(w['分类评估'].max_column, 3)
            self.assertEqual(list(w['分类评估'].values), [('分类', '摄入原重(g)', '评价'), ('奶类', 100, '不足'), ('谷类', 0, '不足')])
            sport = w['运动员专项评价']
            self.assertEqual(sport.max_column, 6)
            self.assertEqual(tuple(next(row for row in sport.values if row[0] == '项目')), ('项目', '实际摄入', '单位', '参考目标', '差值', '最终评价'))
            self.assertEqual(list(sport.values)[-1], ('蛋白质', 0, 'g', '10', -10, '不足'))
            self.assertTrue(all(r.max_col <= 6 for r in sport.merged_cells.ranges))
            self.assertFalse(set(w['分类评估'].column_dimensions) & {'D', 'E', 'F'})
            self.assertNotIn('G', sport.column_dimensions)
            self.assertEqual(w['建议']['B2'].value, '人工建议不改')
            w.close()
        self.assertEqual(data, original)

    def test_reduced_export_is_accepted_by_push_reader(self):
        spec = importlib.util.spec_from_file_location('push', ROOT/'scripts/push_nutrition_reports.py')
        push = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(push)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp)/'report.xlsx'
            NutritionExcelBuilder().build(fixture(), path)
            p = push.read_report(path)
        self.assertEqual(set(p['分类评估'][0]), {'分类', '摄入原重(g)', '评价'})
        self.assertNotIn('口径', p['运动员专项评价'][0])
        self.assertEqual(p['运动员专项评价'][0]['实际摄入'], 0)


if __name__ == '__main__':
    unittest.main()
