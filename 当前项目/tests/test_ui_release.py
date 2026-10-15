from pathlib import Path
import unittest,tempfile,shutil
from unittest.mock import patch
from io import BytesIO
import fitz
from openpyxl import load_workbook
from fastapi import BackgroundTasks, UploadFile
import local_service as service
from nutrition_report.builders.pdf import NutritionPDFBuilder
from tests.test_nutrition_logic_followup import LogicTest

class UIReleaseTest(unittest.TestCase):
    # Missing packaged default must fail on a clean install, not on real business data.
    def test_clean_install_template_and_first_job(self):
        with tempfile.TemporaryDirectory() as t:
            d=Path(t);jobs=d/'jobs';jobs.mkdir();cache=d/'cache.xlsx';meta=d/'meta.json'
            with patch.multiple(service,JOB_ROOT=jobs,LAST_MEDICAL_TEMPLATE=cache,LAST_MEDICAL_TEMPLATE_META=meta):
                status=service.medical_template_status()
                self.assertTrue(status['exists'])
                self.assertEqual('检验结果_病例录入表.xlsx',status['filename'])
                wb=load_workbook(status['path']);ws=wb.active
                self.assertEqual((502,36),(ws.max_row,ws.max_column))
                self.assertTrue(all(c.value is None for row in ws.iter_rows(min_row=3,min_col=2) for c in row));wb.close()
                tasks=BackgroundTasks()
                result=service.process_reports(tasks,'https://fixture.invalid/parse-file',[UploadFile(filename='合成.pdf',file=BytesIO(b'%PDF-fixture'))],['合成.pdf'],None,'markdown','fast')
                try:
                    actual=service.JOBS[result['job_id']]['template_path']
                    self.assertEqual(Path(status['path']).read_bytes(),actual.read_bytes())
                    self.assertEqual(1,len(tasks.tasks))
                finally: service.JOBS.pop(result['job_id'],None)

    def test_custom_template_not_overwritten(self):
        with tempfile.TemporaryDirectory() as t:
            d=Path(t);jobs=d/'jobs';jobs.mkdir();cache=d/'cache.xlsx';meta=d/'meta.json'
            with patch.multiple(service,JOB_ROOT=jobs,LAST_MEDICAL_TEMPLATE=cache,LAST_MEDICAL_TEMPLATE_META=meta):
                source=Path(service.ROOT)/'templates/食物频率调查_汇总模板.xlsx'
                service.save_as_last_medical_template(source,'自定义.xlsx')
                status=service.medical_template_status()
                self.assertEqual('自定义.xlsx',status['filename']);self.assertEqual(source.read_bytes(),cache.read_bytes())

    def test_macro_text_has_space_and_matching_square_colors(self):
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'report.pdf';NutritionPDFBuilder().build(LogicTest().report(),p)
            with fitz.open(p) as doc:
                page=next(p for p in doc if '能量及钙摄入量评估报告' in p.get_text())
                spans=[s for b in page.get_text('dict')['blocks'] for l in b.get('lines',[]) for s in l['spans']]
                target=[s for s in spans if s['text'].startswith(('蛋白质 ','脂肪 ','碳水 '))]
                self.assertEqual(3,len(target))
                # Legend is now a full-width panel; text must stay in its left half.
                self.assertGreater(min(s['bbox'][0] for s in target),60)
                self.assertLess(max(s['bbox'][2] for s in target),page.rect.width/2+35)
                palette={(0.459,0.420,0.694),(0.886,0.549,0.271),(0.298,0.608,0.690)}
                squares=[d for d in page.get_drawings() if 7<=d['rect'].width<=10 and 7<=d['rect'].height<=10 and d['fill']]
                actual={tuple(round(x,3) for x in d['fill']) for d in squares}
                self.assertTrue(palette.issubset(actual),(palette,actual))
                for span in target:
                    self.assertTrue(any(abs(d['rect'].y0-span['bbox'][1])<12 and 0<span['bbox'][0]-d['rect'].x1<20 for d in squares))

if __name__=='__main__': unittest.main()
