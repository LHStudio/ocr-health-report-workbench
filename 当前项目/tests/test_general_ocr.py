import base64
import json
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock, patch

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook

import local_service
from local_service import general_artifact_bytes, general_json_bytes, health, normalize_general_output_format, process_general_job


def workbook_bytes() -> bytes:
    workbook = Workbook()
    workbook.active.append(["项目", "结果"])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def cloud_payload() -> dict:
    return {
        "success": True,
        "excel": base64.b64encode(workbook_bytes()).decode("ascii"),
        "markdown": "# 通用识别结果\n\n| 项目 | 结果 |\n|---|---|\n| 示例 | 正常 |",
        "json_text": json.dumps({"parsing_res_list": [{"block_label": "text", "block_content": "示例内容"}]}, ensure_ascii=False),
    }


class GeneralOcrTests(unittest.TestCase):
    def test_local_health_advertises_new_workspaces(self):
        payload = health()

        self.assertEqual(payload["status"], "ok")
        self.assertIn("imaging", payload["features"])
        self.assertIn("general", payload["features"])

    def test_output_format_validation(self):
        self.assertEqual(normalize_general_output_format("Excel"), "excel")
        self.assertEqual(normalize_general_output_format("markdown"), "markdown")
        self.assertEqual(normalize_general_output_format("json"), "json")
        with self.assertRaises(HTTPException):
            normalize_general_output_format("txt")

    def test_selected_formats_return_the_expected_artifact(self):
        payload = cloud_payload()
        excel = general_artifact_bytes(payload, "excel", "文档.pdf")
        markdown = general_artifact_bytes(payload, "markdown", "文档.pdf")
        structured = general_artifact_bytes(payload, "json", "文档.pdf")

        workbook = load_workbook(BytesIO(excel), read_only=True)
        self.assertEqual(workbook.active.cell(1, 1).value, "项目")
        workbook.close()
        self.assertIn("# 通用识别结果", markdown.decode("utf-8"))
        self.assertEqual(json.loads(structured)["parsing_res_list"][0]["block_content"], "示例内容")

    def test_concatenated_cloud_json_becomes_valid_json(self):
        data = {"json_text": '{"page": 1}\n\n{"page": 2}'}
        payload = json.loads(general_json_bytes(data, "两页.pdf"))

        self.assertEqual(payload["filename"], "两页.pdf")
        self.assertEqual(payload["documents"], [{"page": 1}, {"page": 2}])

    def test_json_fallback_is_converted_to_readable_markdown(self):
        json_text = json.dumps(
            {
                "parsing_res_list": [
                    {"block_id": 0, "block_label": "paragraph_title", "block_content": "文档标题"},
                    {"block_id": 1, "block_label": "text", "block_content": "正文内容"},
                ]
            },
            ensure_ascii=False,
        )
        markdown = general_artifact_bytes({"markdown": json_text, "json_text": json_text}, "markdown", "文档.pdf").decode("utf-8")

        self.assertIn("## 文档标题", markdown)
        self.assertIn("正文内容", markdown)
        self.assertNotIn('"parsing_res_list"', markdown)

    def test_multiple_documents_are_packaged_in_the_selected_format(self):
        response = Mock(ok=True, status_code=200)
        response.json.return_value = cloud_payload()
        with tempfile.TemporaryDirectory() as directory, patch.object(local_service, "request_cloud_ocr", return_value=response):
            job_id = "c" * 32
            job_dir = Path(directory) / job_id
            inputs = []
            for index, name in enumerate(("文档一.pdf", "文档二.png"), start=1):
                source = job_dir / "input" / f"{index:03d}_{name}"
                source.parent.mkdir(parents=True, exist_ok=True)
                source.write_bytes(b"synthetic")
                inputs.append((source, name))
            local_service.JOBS[job_id] = {
                "job_id": job_id,
                "kind": "general",
                "processing_mode": "accurate",
                "status": "queued",
                "job_dir": job_dir,
                "total_files": 2,
                "completed_files": 0,
                "records": [],
            }

            process_general_job(job_id, "http://ocr.example", inputs, "markdown")

            state = local_service.JOBS[job_id]
            self.assertEqual(state["status"], "completed")
            self.assertEqual(len(state["records"]), 2)
            self.assertTrue(state["output_relative"].name.endswith(".zip"))
            with zipfile.ZipFile(job_dir / state["output_relative"]) as archive:
                names = archive.namelist()
                self.assertEqual(len(names), 2)
                self.assertTrue(all(name.endswith(".md") for name in names))

    def test_single_document_returns_its_selected_file_directly(self):
        response = Mock(ok=True, status_code=200)
        response.json.return_value = cloud_payload()
        with tempfile.TemporaryDirectory() as directory, patch.object(local_service, "request_cloud_ocr", return_value=response):
            job_id = "d" * 32
            job_dir = Path(directory) / job_id
            source = job_dir / "input" / "文档.pdf"
            source.parent.mkdir(parents=True)
            source.write_bytes(b"synthetic")
            local_service.JOBS[job_id] = {
                "job_id": job_id,
                "kind": "general",
                "processing_mode": "accurate",
                "status": "queued",
                "job_dir": job_dir,
                "total_files": 1,
                "completed_files": 0,
                "records": [],
            }

            process_general_job(job_id, "http://ocr.example", [(source, "文档.pdf")], "json")

            state = local_service.JOBS[job_id]
            self.assertTrue(state["output_relative"].name.endswith(".json"))
            self.assertIn("示例内容", state["records"][0]["preview"])


if __name__ == "__main__":
    unittest.main()
