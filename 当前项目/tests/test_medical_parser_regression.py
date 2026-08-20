from __future__ import annotations

import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import local_service as service  # noqa: E402


JOB_DIR = PROJECT_ROOT / "data" / "jobs" / "32d9c7374eaa41bc987956ab1a19510f"
TEMPLATE = PROJECT_ROOT / "检验结果_病例录入表.xlsx"
OLD_OUTPUT = JOB_DIR / "output" / "检验结果_病例录入表_已自动填写_20260719_124933.xlsx"
ROUND1_OUTPUT = JOB_DIR / "output" / "检验结果_病例录入表_全字段增强第1轮_20260720.xlsx"
FRESH_ROUND2_OUTPUT = (
    PROJECT_ROOT
    / "data"
    / "jobs"
    / "ad7fcc9ec2724005af8de9077b144070"
    / "output"
    / "检验结果_病例录入表_全字段增强第2轮_20260720.xlsx"
)

RECOVERED_GAPS = {
    "K3": "3.84", "L3": "1.52", "M3": "3.74", "N3": "2.23", "O3": "2.4",
    "S3": "98", "T3": "36.5-40.3", "U3": "45.1-51.3",
    "K4": "3.29", "L4": "1.14", "M4": "3.1", "N4": "24.48", "O4": "1.67",
    "S4": "98", "T4": "36.5-40.3", "U4": "45.1-51.3",
    "K5": "3.36", "L5": "1.43", "M5": "3.56", "N5": "22.97", "O5": "1.81",
    "S5": "98", "T5": "34.5-38.3", "U5": "43.1-49.3", "AA5": "4", "AI5": "5",
    "K6": "3.02", "L6": "1.33", "M6": "2.07", "N6": "1.38", "O6": "1.38",
    "Q6": "3.35", "S6": "95", "T6": "32.3-35.1", "U6": "40.9-46.2",
    "H7": "17.8", "L7": "1.13", "M7": "2.48", "N7": "16.43", "O7": "2.13",
    "S7": "98", "T7": "35.5-39.3", "U7": "44.1-50.3",
    "H8": "16.8", "K8": "3.35", "L8": "1.21", "M8": "4.06", "N8": "1.43",
    "O8": "1.62", "R8": "8.46", "S8": "98", "T8": "35.5-39.3", "U8": "44.1-50.3",
    "X8": "35.19",
}


def parsed_people(parse_mode: str = "markdown") -> dict[str, dict[str, str]]:
    people: dict[str, dict[str, str]] = {}
    excel_root = JOB_DIR / "ocr_excel"
    markdown_root = JOB_DIR / "ocr_markdown"
    for excel_path in sorted(excel_root.rglob("*.xlsx")):
        relative = excel_path.relative_to(excel_root)
        person = relative.parent.as_posix()
        markdown_path = markdown_root / relative.with_suffix(".md")
        markdown = markdown_path.read_text(encoding="utf-8")
        fields = people.setdefault(person, {})
        for key, value in service.parse_medical_page_fields(excel_path.read_bytes(), markdown, parse_mode).items():
            if service.normalized(key) == service.normalized("姓名"):
                fields["姓名"] = value
            else:
                fields.setdefault(key, value)
    return people


class MedicalParserRegressionTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        if not JOB_DIR.is_dir() or not TEMPLATE.is_file() or not OLD_OUTPUT.is_file():
            raise unittest.SkipTest("The checked-in 6-person OCR regression fixture is unavailable")

    def test_all_180_result_cells_match_saved_ocr_evidence(self) -> None:
        self.assertEqual(54, len(RECOVERED_GAPS))
        template_book = load_workbook(TEMPLATE, read_only=True, data_only=True)
        template_sheet = template_book.active
        header_row = service.detect_header_row(template_sheet)
        headers = [
            template_sheet.cell(header_row, column).value for column in range(1, template_sheet.max_column + 1)
        ]
        template_book.close()

        old_book = load_workbook(OLD_OUTPUT, read_only=True, data_only=True)
        old_sheet = old_book.active
        people = parsed_people()
        self.assertEqual(6, len(people))

        for row, person in enumerate(sorted(people), start=3):
            mapped = {
                field["header"]: field["value"]
                for field in service.create_review_fields(headers, people[person])
            }
            for column in range(7, 37):  # G:AJ, 30 laboratory/result columns.
                coordinate = old_sheet.cell(row, column).coordinate
                header = service.clean(old_sheet.cell(header_row, column).value)
                old_value = service.clean(old_sheet.cell(row, column).value)
                expected = old_value or RECOVERED_GAPS.get(coordinate, "")
                self.assertTrue(expected, f"Missing independent expectation for {coordinate} {header}")
                self.assertEqual(expected, service.clean(mapped.get(header)), f"Mismatch at {coordinate} {header}")

        old_book.close()

    def test_round1_generated_workbook_contains_all_expected_values(self) -> None:
        if not ROUND1_OUTPUT.is_file():
            self.skipTest("Round-1 generated workbook is unavailable")
        old_book = load_workbook(OLD_OUTPUT, read_only=True, data_only=True)
        new_book = load_workbook(ROUND1_OUTPUT, read_only=True, data_only=True)
        old_sheet = old_book.active
        new_sheet = new_book.active

        for row in range(3, 9):
            for column in range(2, 37):  # B:AJ, basic information plus all 30 result columns.
                coordinate = old_sheet.cell(row, column).coordinate
                old_value = service.clean(old_sheet.cell(row, column).value)
                expected = old_value or RECOVERED_GAPS.get(coordinate, "")
                self.assertTrue(expected, f"Missing independent expectation for {coordinate}")
                self.assertEqual(expected, service.clean(new_sheet.cell(row, column).value), coordinate)

        old_book.close()
        new_book.close()

    def test_excel_only_mode_also_recovers_every_template_field(self) -> None:
        template_book = load_workbook(TEMPLATE, read_only=True, data_only=True)
        template_sheet = template_book.active
        header_row = service.detect_header_row(template_sheet)
        headers = [
            template_sheet.cell(header_row, column).value for column in range(1, template_sheet.max_column + 1)
        ]
        template_book.close()
        for person, fields in parsed_people("excel").items():
            review = service.create_review_fields(headers, fields)
            missing = [field["header"] for field in review if not field["value"]]
            self.assertFalse(missing, f"Excel-only gaps for {person}: {missing}")

    def test_ocr_range_separator_noise_is_normalized(self) -> None:
        self.assertEqual("44.1-50.3", service.normalize_measurement_text("44.1--50.3"))
        self.assertEqual("44.1-50.3", service.normalize_measurement_text("44.1—50.3"))
        self.assertEqual("44.1-50.3", service.result_only("推测围绝经期开始年龄", "44.1－50.3"))
        self.assertTrue(service.looks_like_measurement("44.1--50.3"))

    def test_fresh_server_round2_matches_the_independent_round1(self) -> None:
        if not ROUND1_OUTPUT.is_file() or not FRESH_ROUND2_OUTPUT.is_file():
            self.skipTest("Independent round outputs are unavailable")
        round1_book = load_workbook(ROUND1_OUTPUT, read_only=True, data_only=True)
        round2_book = load_workbook(FRESH_ROUND2_OUTPUT, read_only=True, data_only=True)
        round1_sheet = round1_book.active
        round2_sheet = round2_book.active
        for row in range(3, 9):
            for column in range(1, 37):
                coordinate = round1_sheet.cell(row, column).coordinate
                self.assertEqual(
                    service.clean(round1_sheet.cell(row, column).value),
                    service.clean(round2_sheet.cell(row, column).value),
                    coordinate,
                )
        round1_book.close()
        round2_book.close()


if __name__ == "__main__":
    unittest.main()
