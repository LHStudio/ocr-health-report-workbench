import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from fastapi import HTTPException
from openpyxl import load_workbook
import local_service
from nutrition_report import NutritionReportService


class NutritionBatchTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.job = 'batch-fixture'
        self.people = [dict(id=f'folder/{i}', general={'姓名': '同名', '体重': 65},
            food_rows=[{'食物名称': '坚果', '平均每次食用量': '10g', '次数': '1', '频率周期(请核对)': '每天'}])
            for i in range(3)]
        local_service.JOBS[self.job] = dict(kind='nutrition', status='completed',
            job_dir=Path(self.temp.name), people=self.people)
        # Only the external food service is replaced; real builders generate and reopen files.
        self.mock = patch.object(local_service, 'NutritionReportService',
            side_effect=lambda: NutritionReportService(prefer_cloud=False))
        self.mock.start()

    def tearDown(self):
        self.mock.stop()
        local_service.JOBS.pop(self.job, None)
        self.temp.cleanup()

    def export(self, payload):
        self.assertTrue(hasattr(local_service, 'generate_nutrition_batch_reports'), 'missing batch export')
        return local_service.generate_nutrition_batch_reports(self.job, payload)

    def test_all_people_both_formats_duplicate_names_and_manifest(self):
        response = self.export({'format': 'both'})
        self.assertEqual('application/zip', response.media_type)
        with zipfile.ZipFile(response.path) as z:
            self.assertIsNone(z.testzip())
            manifest = json.loads(z.read('manifest.json'))
            self.assertEqual(3, manifest['requested_people'])
            self.assertEqual(6, manifest['success_count'])
            self.assertEqual(0, manifest['failure_count'])
            files = [n for n in z.namelist() if n.endswith(('.pdf', '.xlsx'))]
            self.assertEqual(6, len(set(files)))
            for name in files:
                if name.endswith('.pdf'):
                    self.assertTrue(z.read(name).startswith(b'%PDF'))
                else:
                    from io import BytesIO
                    wb = load_workbook(BytesIO(z.read(name)))
                    self.assertIn('新增膳食标准', wb.sheetnames)

    def test_bad_person_does_not_silently_omit_or_abort_others(self):
        self.people[1]['food_rows'] = []
        response = self.export({'format': 'excel'})
        with zipfile.ZipFile(response.path) as z:
            m = json.loads(z.read('manifest.json'))
            self.assertEqual(2, m['success_count'])
            self.assertEqual(1, m['failure_count'])
            self.assertEqual('folder/1', next(r for r in m['results'] if r['status'] == 'failed')['person_id'])

    def test_input_scope_and_format_validation(self):
        for payload in ({'format': 'bad'}, {'format': 'pdf', 'people': self.people[:1]},
                        {'format': 'pdf', 'people': self.people + [self.people[0]]}):
            with self.assertRaises(HTTPException) as e:
                self.export(payload)
            self.assertEqual(400, e.exception.status_code)

    def test_overrides_and_manual_review_apply_to_correct_id(self):
        response = self.export({'format': 'excel', 'profiles': {'folder/0': {'name': '改名'}},
            'reviews': {'folder/0': {'standardEvaluations': {'nuts': '人工核查'}}}})
        with zipfile.ZipFile(response.path) as z:
            m = json.loads(z.read('manifest.json'))
            self.assertEqual('manual_reviewed', m['results'][0]['evaluation_mode'])
            self.assertEqual('automatic', m['results'][1]['evaluation_mode'])
            from io import BytesIO
            wb = load_workbook(BytesIO(z.read(m['results'][0]['file'])))
            values = [c.value for row in wb['新增膳食标准'] for c in row]
            self.assertIn('人工核查', values)
