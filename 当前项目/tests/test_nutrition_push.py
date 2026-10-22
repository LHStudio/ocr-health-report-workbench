import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock
from openpyxl import Workbook

SCRIPT = Path(__file__).resolve().parents[1] / 'scripts/push_nutrition_reports.py'
if SCRIPT.exists():
    spec = importlib.util.spec_from_file_location('nutrition_push', SCRIPT)
    push = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(push)
else:
    push = None


class PushTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(push, 'Independent nutrition push entry point is not implemented')
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.file = self.root / 'report.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = '个人信息'
        for row in [('项目', '内容'), ('姓名', '测试运动员'), ('日期', '2025.12.03'), ('体重', '57.9 kg')]:
            ws.append(row)
        for title, rows in {
            '食物明细': [('食物名称', '分类', '摄入量(g)', '能量(kcal)'), ('米饭', '谷类', 170, 197.2)],
            '分类评估': [('分类', '摄入原重(g)', '评价'), ('谷类', 170, '适中')],
            '营养总计': [('项目', '摄入量', '单位', '评价'), ('总能量', 197.2, 'kcal/day', '过少'), ('推荐能量', '1700-2200', 'kcal/天', None)],
            '运动员专项评价': [('运动员专项膳食评价',), ('说明',), ('项目', '实际摄入', '单位', '参考目标', '差值', '最终评价', '口径'), ('蛋白质', 0, 'g', '10', -10, '不足', None)],
            '建议': [('序号', '建议内容'), (1, '保持原建议'), ('报告生成日期', '2026-10-22 10:08:19')],
        }.items():
            s = wb.create_sheet(title)
            for row in rows:
                s.append(row)
        wb.save(self.file)

    def test_workbook_mapping_preserves_values_and_generated_date(self):
        before = self.file.read_bytes()
        p = push.read_report(self.file)
        self.assertEqual(p['analysisDate'], '2025-12-03')
        self.assertEqual(p['食物明细'][0]['摄入量 (g)'], 170)
        self.assertEqual(p['营养总计'][1]['摄入量'], '1700-2200')
        self.assertEqual(p['运动员专项评价'][0]['实际摄入'], 0)
        self.assertEqual(p['建议'][-1]['建议内容'], '2026-10-22 10:08:19')
        self.assertEqual(p['OCR换算说明'], [])
        self.assertEqual(before, self.file.read_bytes())

    def test_missing_date_rejected(self):
        from openpyxl import load_workbook
        w = load_workbook(self.file)
        w['个人信息']['B3'] = None
        w.save(self.file)
        with self.assertRaisesRegex(ValueError, '日期'):
            push.read_report(self.file)

    def test_invalid_birthday_rejected(self):
        with self.assertRaises(ValueError):
            push.read_report(self.file, birthday='2010-02-30')

    def test_chinese_report_date_normalized_without_changing_display(self):
        from openpyxl import load_workbook
        w = load_workbook(self.file)
        w['个人信息']['B3'] = '2026年6月4日'
        w.save(self.file)
        p = push.read_report(self.file)
        self.assertEqual(p['analysisDate'], '2026-06-04')
        self.assertEqual(p['个人信息']['日期'], '2026-06-04')

    def test_duplicate_person_date_rejected(self):
        other = self.root / 'duplicate.xlsx'
        other.write_bytes(self.file.read_bytes())
        with self.assertRaisesRegex(ValueError, '重复'):
            push.prepare([self.file, other], {})

    def test_formula_rejected_not_silently_dropped(self):
        from openpyxl import load_workbook
        w = load_workbook(self.file)
        w['食物明细']['B2'] = '=100+70'
        w.save(self.file)
        with self.assertRaisesRegex(ValueError, '公式'):
            push.read_report(self.file)

    def test_limit_rejected(self):
        from openpyxl import load_workbook
        w = load_workbook(self.file)
        for _ in range(500):
            w['食物明细'].append(['米饭', '谷类', 1, 1])
        w.save(self.file)
        with self.assertRaisesRegex(ValueError, '500'):
            push.read_report(self.file)

    def response(self, code=0):
        return {'code': code, 'message': 'success' if code == 0 else '没有找到运动员',
                'data': {'athleteId': '123', 'analysisId': '456', 'name': '测试运动员', 'analysisDate': '2025-12-03', 'created': True}}

    def test_success_resume_no_second_post(self):
        send = Mock(return_value=(200, self.response()))
        items = push.prepare([self.file], {})
        for _ in range(2):
            result = push.execute(items, 'https://example.test/open/nutrition', 'secret', self.root / 'audit', send=send)
            self.assertEqual(result['success'], 1)
        self.assertEqual(send.call_count, 1)
        self.assertNotIn('secret', (self.root / 'audit/checkpoint.json').read_text(encoding='utf-8'))

    def test_business_failure_not_success(self):
        r = push.execute(push.prepare([self.file], {}), 'https://example.test/open/nutrition', 'secret', self.root / 'audit', send=Mock(return_value=(200, self.response(1))))
        self.assertEqual(r['failed'], 1)
        self.assertEqual(r['success'], 0)

    def test_timeout_uncertain_not_retried(self):
        send = Mock(side_effect=TimeoutError('timeout'))
        for _ in range(2):
            r = push.execute(push.prepare([self.file], {}), 'https://example.test/open/nutrition', 'secret', self.root / 'audit', send=send)
            self.assertEqual(r['uncertain'], 1)
        self.assertEqual(send.call_count, 1)

    def test_wrong_response_identity_is_uncertain(self):
        response = self.response()
        response['data']['name'] = '另一个人'
        r = push.execute(push.prepare([self.file], {}), 'https://example.test/open/nutrition', 'secret', self.root / 'audit', send=Mock(return_value=(200, response)))
        self.assertEqual(r['uncertain'], 1)

    def test_endpoint_rejects_credentials_query_and_insecure_remote(self):
        for url in ['http://example.test', 'https://user:pass@example.test', 'https://example.test?token=x', 'https://example.test/wrong']:
            with self.assertRaises(ValueError):
                push.endpoint_url(url)
        self.assertEqual(push.endpoint_url('https://example.test/'), 'https://example.test/open/nutrition')

    def test_post_does_not_follow_redirects(self):
        from unittest.mock import patch
        response = Mock(status_code=302)
        response.json.return_value = {'message': 'redirect'}
        with patch.object(push.requests, 'post', return_value=response) as call:
            push.post('https://example.test/open/nutrition', 'secret', {'name': '测试运动员'})
        self.assertFalse(call.call_args.kwargs['allow_redirects'])
        self.assertEqual(call.call_args.kwargs['headers']['X-Open-Token'], 'secret')

    def test_documented_reverse_proxy_api_prefix(self):
        for url in ['https://example.test/api', 'https://example.test/api/open/nutrition']:
            self.assertEqual(push.endpoint_url(url), 'https://example.test/api/open/nutrition')

    def test_auth_failure_stops_remaining(self):
        items = push.prepare([self.file], {})
        import copy
        second = copy.deepcopy(items[0])
        second['source'] = 'other.xlsx'
        second['payload']['name'] = '测试运动员二'
        send = Mock(return_value=(401, {'code': 1, 'message': 'unauthorized'}))
        r = push.execute(items + [second], 'https://example.test/open/nutrition', 'secret', self.root / 'audit', send=send)
        self.assertEqual((r['failed'], r['pending'], send.call_count), (1, 1, 1))

    def test_remote_unknown_status_not_success(self):
        for status, body in [(302, {}), (500, {}), (200, {'invalid_json': True}), (200, {'code': False})]:
            r = push.execute(push.prepare([self.file], {}), 'https://example.test/open/nutrition', 'secret', self.root / ('audit' + str(status) + str(len(str(body)))), send=Mock(return_value=(status, body)))
            self.assertEqual(r['uncertain'], 1)

    def test_ambiguous_dates_not_guessed(self):
        for text in ['26.6.12', 'b.5', '', '2026-02-30']:
            with self.assertRaises(ValueError):
                push.normalize_date(text, '报告日期')

    def test_changed_endpoint_cannot_reuse_checkpoint(self):
        items = push.prepare([self.file], {})
        push.execute(items, 'https://example.test/open/nutrition', 'secret', self.root / 'audit', send=Mock(return_value=(200, self.response())))
        with self.assertRaisesRegex(ValueError, '批次'):
            push.execute(items, 'https://other.test/open/nutrition', 'secret', self.root / 'audit', send=Mock())

    def test_changed_batch_cannot_reuse_checkpoint(self):
        items = push.prepare([self.file], {})
        push.execute(items, 'https://example.test/open/nutrition', 'secret', self.root / 'audit', send=Mock(return_value=(200, self.response())))
        items[0]['payload']['个人信息']['体重'] = '60 kg'
        with self.assertRaisesRegex(ValueError, '批次'):
            push.execute(items, 'https://example.test/open/nutrition', 'secret', self.root / 'audit', send=Mock())


if __name__ == '__main__':
    unittest.main()
