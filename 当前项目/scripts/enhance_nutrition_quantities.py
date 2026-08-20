from __future__ import annotations

"""Batch OCR evidence for suspicious food-quantity cells.

The tool is intentionally non-mutating: it never writes to the source workbook
and never declares an OCR value accepted.  It aligns the questionnaire pages,
extracts the original gray quantity cell, sends gray and threshold variants in a
single classic-OCR request, and saves candidates for human review.
"""

import argparse
import hashlib
import json
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Sequence

import cv2
import numpy as np
import requests
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from scripts import enhance_nutrition_cells as cells  # noqa: E402


PARSER_VERSION = "nutrition-food-quantity-classic-v1"
SCHEMA_VERSION = 1
VARIANTS = ("grayscale", "threshold")
RISK_PHRASES = (
    "食用量疑似粘连",
    "食用量疑似异常",
    "食用量没有有效数字",
    "食用量未识别到单位",
)

QUANTITY_CELL_WIDTH = 360
QUANTITY_CELL_HEIGHT = 110
QUANTITY_CELL_GAP = 28
QUANTITY_BAND_GAP = 38
QUANTITY_BAND_WIDTH = len(VARIANTS) * QUANTITY_CELL_WIDTH + (len(VARIANTS) - 1) * QUANTITY_CELL_GAP
MIN_CANDIDATE_SCORE = 0.45

GLOBAL_COMPOSITE_FILENAME = "food_quantity_batch.png"
GLOBAL_BATCH_FILENAME = "food_quantity_batch.json"
GLOBAL_ROW_MONTAGE_FILENAME = "food_quantity_row_montage.png"
PERSON_RESULT_FILENAME = "food_quantity_classic_ocr.json"

_DIGIT_TRANSLATION = str.maketrans("０１２３４５６７８９", "0123456789")
_UNIT_PATTERN = r"(?:kg|mg|ml|g|L|l|克|千克|毫升|升|勺|个|只|粒|片|瓶|杯)"


def now_text() -> str:
    return datetime.now().isoformat(timespec="seconds")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _cell_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _food_sheet_rows(workbook_path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = next((sheet for sheet in workbook.worksheets if "食物" in sheet.title), None)
        if worksheet is None:
            raise RuntimeError("汇总表中没有名称包含“食物”的工作表")
        header_row = cells._header_row(worksheet)
        headers = [_cell_text(worksheet.cell(header_row, column).value) for column in range(1, worksheet.max_column + 1)]
        required = {"人员文件夹", "食物编号", "食物名称", "平均每次食用量", "人工核对备注"}
        if not required.issubset(headers):
            raise RuntimeError(f"食物明细缺少列：{'、'.join(sorted(required - set(headers)))}")
        result: list[dict[str, str]] = []
        for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            row = {
                header: _cell_text(values[index] if index < len(values) else None)
                for index, header in enumerate(headers)
                if header
            }
            if row.get("人员文件夹") and row.get("食物编号"):
                result.append(row)
        return result
    finally:
        workbook.close()


def parse_target_key(value: str) -> tuple[str, str]:
    raw = value.strip().replace("\\", "/")
    separator = ":" if ":" in raw else "#" if "#" in raw else ""
    if not separator:
        raise ValueError(f"显式目标必须为 person:code，例如 ocr/12:2.1；收到 {value!r}")
    person, code = (item.strip() for item in raw.rsplit(separator, 1))
    if person and "/" not in person:
        person = f"ocr/{person}"
    if not person or code not in cells.TABLE_BY_CODE:
        raise ValueError(f"无效显式目标：{value!r}")
    return person, code


def collect_quantity_targets(
    workbook_path: Path,
    *,
    scope: str = "flagged",
    explicit_targets: Sequence[str] | None = None,
) -> tuple[list[dict[str, str]], list[str]]:
    if scope not in {"flagged", "missing", "all"}:
        raise ValueError("scope must be flagged, missing, or all")
    rows = _food_sheet_rows(workbook_path)
    people = sorted({row["人员文件夹"].replace("\\", "/") for row in rows}, key=cells.person_sort_key)
    explicit = {parse_target_key(value) for value in explicit_targets or []}
    selected: list[dict[str, str]] = []
    found: set[tuple[str, str]] = set()
    for row in rows:
        person = row["人员文件夹"].replace("\\", "/")
        code = row["食物编号"]
        if code not in cells.TABLE_BY_CODE:
            continue
        key = (person, code)
        note = row.get("人工核对备注", "")
        if explicit:
            include = key in explicit
        elif scope == "flagged":
            include = any(phrase in note for phrase in RISK_PHRASES)
        elif scope == "missing":
            include = not row.get("平均每次食用量", "")
        else:
            include = True
        if not include or key in found:
            continue
        found.add(key)
        selected.append(
            {
                "person": person,
                "code": code,
                "name": row.get("食物名称", ""),
                "current_quantity": row.get("平均每次食用量", ""),
                "workbook_count": row.get("次数", ""),
                "workbook_period": row.get("频率周期(请核对)", ""),
                "workbook_raw": row.get("OCR原始行", ""),
                "workbook_review_note": note,
            }
        )
    if explicit:
        missing = explicit - found
        if missing:
            display = "、".join(f"{person}:{code}" for person, code in sorted(missing))
            raise RuntimeError(f"汇总表中找不到显式目标：{display}")
    code_order = {code: index for index, code in enumerate(cells.FOOD_CODES)}
    selected.sort(key=lambda item: (cells.person_sort_key(item["person"]), code_order[item["code"]]))
    return selected, people


def quantity_cell_bounds(spec: cells.TableSpec, row_index: int, inset: int = cells.CELL_INSET) -> tuple[int, int, int, int]:
    if not 0 <= row_index < len(spec.codes):
        raise IndexError(f"row_index {row_index} outside page {spec.page_number}")
    x0 = cells.ROW_CELL_X_BOUNDARIES[1] + inset
    x1 = cells.ROW_CELL_X_BOUNDARIES[2] - inset
    y0 = spec.header_height + row_index * spec.row_height + inset
    y1 = spec.header_height + (row_index + 1) * spec.row_height - inset
    return x0, y0, x1, y1


def extract_quantity_cell(current_table: np.ndarray, spec: cells.TableSpec, row_index: int) -> np.ndarray:
    x0, y0, x1, y1 = quantity_cell_bounds(spec, row_index)
    return current_table[y0:y1, x0:x1].copy()


def autocontrast_gray(image: np.ndarray) -> np.ndarray:
    if image.size == 0:
        return image.copy()
    foreground = image[image < 250]
    if foreground.size == 0:
        return np.full_like(image, 255)
    low = float(np.percentile(foreground, 3))
    normalized = np.clip((image.astype(np.float32) - low) * (255.0 / max(255.0 - low, 1.0)), 0, 255).astype(np.uint8)
    normalized[image >= 252] = 255
    return normalized


def make_quantity_variants(image: np.ndarray) -> dict[str, np.ndarray]:
    normalized = autocontrast_gray(image)
    thresholded = cv2.threshold(normalized, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
    return {"grayscale": normalized, "threshold": thresholded}


def _fit_variant(image: np.ndarray) -> np.ndarray:
    target = np.full((QUANTITY_CELL_HEIGHT, QUANTITY_CELL_WIDTH), 255, dtype=np.uint8)
    if image.size == 0:
        return target
    resized = cv2.resize(image, (QUANTITY_CELL_WIDTH - 16, QUANTITY_CELL_HEIGHT - 16), interpolation=cv2.INTER_CUBIC)
    target[8:-8, 8:-8] = resized
    return target


def build_quantity_strip(variants: dict[str, np.ndarray]) -> tuple[np.ndarray, list[dict[str, Any]]]:
    strip = np.full((QUANTITY_CELL_HEIGHT, QUANTITY_BAND_WIDTH), 255, dtype=np.uint8)
    mappings: list[dict[str, Any]] = []
    for index, variant in enumerate(VARIANTS):
        x0 = index * (QUANTITY_CELL_WIDTH + QUANTITY_CELL_GAP)
        x1 = x0 + QUANTITY_CELL_WIDTH
        strip[:, x0:x1] = _fit_variant(variants[variant])
        mappings.append({"variant": variant, "x_start": x0, "x_end": x1, "y_start": 0, "y_end": QUANTITY_CELL_HEIGHT})
    return strip, mappings


def build_batch_composite(prepared: Sequence[dict[str, Any]]) -> tuple[np.ndarray, list[dict[str, Any]]]:
    if not prepared:
        return np.full((1, QUANTITY_BAND_WIDTH), 255, dtype=np.uint8), []
    height = len(prepared) * QUANTITY_CELL_HEIGHT + (len(prepared) - 1) * QUANTITY_BAND_GAP
    composite = np.full((height, QUANTITY_BAND_WIDTH), 255, dtype=np.uint8)
    bands: list[dict[str, Any]] = []
    y = 0
    for row in prepared:
        composite[y : y + QUANTITY_CELL_HEIGHT] = row["ocr_strip"]
        mappings = []
        for mapping in row["variant_mappings"]:
            item = dict(mapping)
            item["y_start"] = y
            item["y_end"] = y + QUANTITY_CELL_HEIGHT
            mappings.append(item)
        bands.append(
            {
                "id": row["id"],
                "person": row["person"],
                "code": row["code"],
                "x_start": 0,
                "x_end": QUANTITY_BAND_WIDTH,
                "y_start": y,
                "y_end": y + QUANTITY_CELL_HEIGHT,
                "variants": mappings,
            }
        )
        y += QUANTITY_CELL_HEIGHT + QUANTITY_BAND_GAP
    return composite, bands


def build_row_montage(prepared: Sequence[dict[str, Any]]) -> tuple[np.ndarray, list[dict[str, Any]]]:
    if not prepared:
        return np.full((1, cells.TABLE_WIDTH * 2), 255, dtype=np.uint8), []
    header_height = 28
    row_height = cells.TABLE_ROW_HEIGHT * 2
    gap = 10
    band_height = header_height + row_height
    width = cells.TABLE_WIDTH * 2
    height = len(prepared) * band_height + (len(prepared) - 1) * gap
    montage = np.full((height, width), 255, dtype=np.uint8)
    bands: list[dict[str, Any]] = []
    y = 0
    for row in prepared:
        label = f"{row['person']}  #{row['code']}"
        cv2.putText(montage, label, (10, y + 20), cv2.FONT_HERSHEY_SIMPLEX, 0.55, 40, 1, cv2.LINE_AA)
        enlarged = cv2.resize(row["raw_row"], (width, row_height), interpolation=cv2.INTER_CUBIC)
        montage[y + header_height : y + band_height] = enlarged
        bands.append({"person": row["person"], "code": row["code"], "y_start": y, "y_end": y + band_height})
        y += band_height + gap
    return montage, bands


def _overlap(a0: float, a1: float, b0: float, b1: float) -> float:
    return max(0.0, min(a1, b1) - max(a0, b0))


def map_ocr_items(items: Sequence[object], bands: Sequence[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    mapped: list[dict[str, Any]] = []
    ambiguous: list[dict[str, Any]] = []
    unmapped: list[dict[str, Any]] = []
    for source_index, raw_item in enumerate(items):
        if not isinstance(raw_item, dict):
            unmapped.append({"source_index": source_index, "reason": "item_not_object", "item": raw_item})
            continue
        box = cells.normalized_box(raw_item.get("box"))
        if not box:
            unmapped.append({"source_index": source_index, "reason": "box_missing", "item": raw_item})
            continue
        x0, y0, x1, y1 = box
        center_x, center_y = (x0 + x1) / 2, (y0 + y1) / 2
        height = max(1.0, y1 - y0)
        band_scores = sorted(
            ((_overlap(y0, y1, float(band["y_start"]), float(band["y_end"])), band) for band in bands),
            key=lambda pair: pair[0],
            reverse=True,
        )
        if not band_scores or band_scores[0][0] / height < 0.50:
            unmapped.append({"source_index": source_index, "reason": "band_overlap_low", "item": raw_item, "box": box})
            continue
        best_y, band = band_scores[0]
        if not (float(band["y_start"]) <= center_y <= float(band["y_end"])):
            unmapped.append({"source_index": source_index, "reason": "band_center_outside", "item": raw_item, "box": box})
            continue
        if len(band_scores) > 1 and band_scores[1][0] >= best_y * 0.80 and band_scores[1][0] > 0:
            ambiguous.append({"source_index": source_index, "reason": "band_ambiguous", "item": raw_item, "box": box})
            continue
        width = max(1.0, x1 - x0)
        variant_scores = sorted(
            ((_overlap(x0, x1, float(item["x_start"]), float(item["x_end"])), item) for item in band["variants"]),
            key=lambda pair: pair[0],
            reverse=True,
        )
        best_x, variant = variant_scores[0]
        if best_x / width < 0.50 or not (float(variant["x_start"]) <= center_x <= float(variant["x_end"])):
            unmapped.append({"source_index": source_index, "reason": "variant_overlap_low", "item": raw_item, "box": box})
            continue
        if len(variant_scores) > 1 and variant_scores[1][0] >= best_x * 0.80 and variant_scores[1][0] > 0:
            ambiguous.append({"source_index": source_index, "reason": "variant_ambiguous", "item": raw_item, "box": box})
            continue
        mapped.append(
            {
                "source_index": source_index,
                "person": band["person"],
                "code": band["code"],
                "variant": variant["variant"],
                "text": str(raw_item.get("text") or "").strip(),
                "score": round(float(raw_item.get("score") or 0.0), 6),
                "box": box,
                "vertical_overlap_ratio": round(best_y / height, 4),
                "horizontal_overlap_ratio": round(best_x / width, 4),
            }
        )
    return {"mapped": mapped, "ambiguous": ambiguous, "unmapped": unmapped}


def normalize_quantity_candidate(value: object) -> str:
    text = str(value or "").translate(_DIGIT_TRANSLATION)
    text = text.replace("ＭＬ", "ml").replace("ｍｌ", "ml").replace("ML", "ml")
    text = text.replace("～", "-").replace("~", "-").replace("—", "-").replace("–", "-").replace("至", "-")
    text = re.sub(r"\s+", "", text).strip("。；;，,:：")
    match = re.fullmatch(rf"(\d+(?:\.\d+)?(?:-\d+(?:\.\d+)?)?)({_UNIT_PATTERN})", text, re.I)
    if not match:
        return ""
    number, unit = match.groups()
    parts = [float(item) for item in number.split("-")]
    if len(parts) == 2 and parts[0] > parts[1]:
        return ""
    number_text = "-".join(str(int(item)) if item.is_integer() else f"{item:g}" for item in parts)
    unit_text = unit.lower() if unit.lower() in {"kg", "mg", "ml", "g", "l"} else unit
    return f"{number_text}{unit_text}"


def quantity_candidates(items: Sequence[dict[str, Any]], minimum_score: float = MIN_CANDIDATE_SCORE) -> list[str]:
    reliable = [item for item in items if float(item.get("score") or 0.0) >= minimum_score and str(item.get("text") or "").strip()]
    candidates: list[str] = []
    sources = [item["text"] for item in reliable]
    if sources:
        sources.append(" ".join(item["text"] for item in sorted(reliable, key=lambda item: item.get("box", [0])[0] if item.get("box") else 0)))
    for source in sources:
        candidate = normalize_quantity_candidate(source)
        if candidate and candidate not in candidates:
            candidates.append(candidate)
    return candidates


def review_quantity(items_by_variant: dict[str, list[dict[str, Any]]]) -> dict[str, Any]:
    candidates = {variant: quantity_candidates(items_by_variant.get(variant, [])) for variant in VARIANTS}
    sets = [set(values) for values in candidates.values() if values]
    union = set().union(*sets) if sets else set()
    if len(sets) == len(VARIANTS) and len(union) == 1 and all(len(values) == 1 for values in sets):
        status = "consistent_candidate"
        proposed = next(iter(union))
        reasons = ["two_variants_agree"]
    elif not union:
        status = "no_candidate"
        proposed = ""
        reasons = ["no_structured_quantity_candidate"]
    elif len(union) == 1:
        status = "single_variant_candidate"
        proposed = next(iter(union))
        reasons = ["only_one_variant_has_candidate"]
    else:
        status = "conflict"
        proposed = ""
        reasons = ["quantity_variants_conflict"]
    return {
        "status": status,
        "proposed_quantity": proposed,
        "candidates_by_variant": candidates,
        "reason_codes": reasons,
        "manual_review_required": True,
    }


def _crop_paths(job_dir: Path, person: str, code: str) -> tuple[Path, Path]:
    safe = code.replace(".", "_")
    root = job_dir / "ocr_enhanced" / Path(person) / "food_quantities"
    return root / f"quantity_{safe}.png", root / f"quantity_{safe}_row.png"


def prepare_targets(
    job_dir: Path,
    targets: Sequence[dict[str, str]],
    references: dict[int, np.ndarray],
) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, str]]] = {}
    for target in targets:
        grouped.setdefault(target["person"], []).append(target)
    prepared: list[dict[str, Any]] = []
    for person in sorted(grouped, key=cells.person_sort_key):
        sources = cells.source_pages(job_dir, person)
        page_cache: dict[int, tuple[np.ndarray, dict[str, Any], Path]] = {}
        for spec in {cells.TABLE_BY_CODE[target["code"]] for target in grouped[person]}:
            source_pdf = sources[spec.source_page_index]
            reference = references[spec.page_number]
            rendered = cells.render_pdf_page_gray(source_pdf, reference.shape)
            aligned, alignment = cells.align_orb(rendered, reference)
            page_cache[spec.page_number] = (cells.warp_table(aligned, spec), alignment, source_pdf)
        for target in grouped[person]:
            spec = cells.TABLE_BY_CODE[target["code"]]
            row_index = spec.row_index(target["code"])
            current_table, alignment, source_pdf = page_cache[spec.page_number]
            quantity = extract_quantity_cell(current_table, spec, row_index)
            variants = make_quantity_variants(quantity)
            strip, mappings = build_quantity_strip(variants)
            crop_path, row_path = _crop_paths(job_dir, person, target["code"])
            cells.write_png(crop_path, strip)
            raw_y0 = spec.header_height + row_index * spec.row_height
            raw_row = current_table[raw_y0 : raw_y0 + spec.row_height, :].copy()
            cells.write_png(row_path, cv2.resize(raw_row, (cells.TABLE_WIDTH * 2, cells.TABLE_ROW_HEIGHT * 2), interpolation=cv2.INTER_CUBIC))
            confidence = float(alignment.get("confidence") or 0.0)
            row: dict[str, Any] = {
                "id": f"{person}|{target['code']}",
                **target,
                "source_pdf": str(source_pdf.resolve()),
                "source_pdf_sha256": sha256_file(source_pdf),
                "source_page_index": spec.source_page_index + 1,
                "table_page": spec.page_number,
                "table_row_index": row_index,
                "alignment_confidence": round(confidence, 6),
                "alignment": alignment,
                "crop": str(crop_path.resolve()),
                "crop_sha256": sha256_file(crop_path),
                "raw_row_crop": str(row_path.resolve()),
                "raw_row_crop_sha256": sha256_file(row_path),
                "variants": list(VARIANTS),
                "items_by_variant": {variant: [] for variant in VARIANTS},
                "text_by_variant": {variant: "" for variant in VARIANTS},
                "status": "ocr_pending" if confidence >= cells.MIN_ALIGNMENT_CONFIDENCE else "alignment_low",
                "proposed_quantity": "",
                "candidates_by_variant": {variant: [] for variant in VARIANTS},
                "reason_codes": ["classic_ocr_pending"] if confidence >= cells.MIN_ALIGNMENT_CONFIDENCE else ["alignment_confidence_below_threshold"],
                "manual_review_required": True,
                "ocr_strip": strip,
                "variant_mappings": mappings,
                "raw_row": raw_row,
            }
            prepared.append(row)
    return prepared


def apply_response(rows: Sequence[dict[str, Any]], response: dict[str, Any], bands: Sequence[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    raw_items = response.get("items") if isinstance(response.get("items"), list) else []
    diagnostics = map_ocr_items(raw_items, bands)
    grouped: dict[str, list[dict[str, Any]]] = {}
    for item in diagnostics["mapped"]:
        grouped.setdefault(f"{item['person']}|{item['code']}", []).append(item)
    for row in rows:
        if row["status"] != "ocr_pending":
            continue
        items_by_variant = {variant: [] for variant in VARIANTS}
        for item in grouped.get(row["id"], []):
            items_by_variant[item["variant"]].append(item)
        row["items_by_variant"] = items_by_variant
        row["text_by_variant"] = {
            variant: "\n".join(item["text"] for item in variant_items if item.get("text"))
            for variant, variant_items in items_by_variant.items()
        }
        row.update(review_quantity(items_by_variant))
    return diagnostics


def _public_row(row: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in row.items() if key not in {"ocr_strip", "variant_mappings", "raw_row"}}


def _load_json(path: Path) -> dict[str, Any]:
    if not path.is_file():
        return {}
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError):
        return {}
    return value if isinstance(value, dict) else {}


def _write_results(job_dir: Path, workbook_path: Path, rows: Sequence[dict[str, Any]], payload: dict[str, Any]) -> None:
    public_rows = [_public_row(row) for row in rows]
    payload["rows"] = public_rows
    batch_path = job_dir / "ocr_enhanced" / GLOBAL_BATCH_FILENAME
    cells.write_json_atomic(batch_path, payload)
    by_person: dict[str, list[dict[str, Any]]] = {}
    for row in public_rows:
        by_person.setdefault(row["person"], []).append(row)
    for person, person_rows in by_person.items():
        path = job_dir / "ocr_enhanced" / Path(person) / PERSON_RESULT_FILENAME
        cells.write_json_atomic(
            path,
            {
                "schema_version": SCHEMA_VERSION,
                "parser_version": PARSER_VERSION,
                "generated_at": payload["generated_at"],
                "person": person,
                "workbook": str(workbook_path.resolve()),
                "workbook_sha256": payload["workbook_sha256"],
                "batch": str(batch_path.resolve()),
                "rows": person_rows,
            },
        )


def enhance_quantities(
    job_dir: Path,
    workbook_path: Path,
    cloud_url: str | None,
    *,
    scope: str = "flagged",
    explicit_targets: Sequence[str] | None = None,
    force: bool = False,
    prepare_only: bool = False,
    timeout: int = 900,
) -> dict[str, Any]:
    job_dir = job_dir.resolve()
    workbook_path = workbook_path.resolve()
    if not job_dir.is_dir():
        raise FileNotFoundError(f"任务目录不存在：{job_dir}")
    if not workbook_path.is_file():
        raise FileNotFoundError(f"汇总表不存在：{workbook_path}")
    targets, people = collect_quantity_targets(workbook_path, scope=scope, explicit_targets=explicit_targets)
    page1 = cells.read_gray_image(cells.PAGE1_REFERENCE)
    references = {1: page1}
    page2_metadata: dict[str, Any] = {}
    if any(cells.TABLE_BY_CODE[target["code"]].page_number == 2 for target in targets):
        page2, page2_metadata = cells.build_page2_reference(job_dir, people, page1.shape, force=force)
        references[2] = page2
    rows = prepare_targets(job_dir, targets, references)
    ocr_ready = [row for row in rows if row["status"] == "ocr_pending"]
    composite, bands = build_batch_composite(ocr_ready)
    montage, montage_bands = build_row_montage(rows)
    output_root = job_dir / "ocr_enhanced"
    composite_path = output_root / GLOBAL_COMPOSITE_FILENAME
    montage_path = output_root / GLOBAL_ROW_MONTAGE_FILENAME
    cells.write_png(composite_path, composite)
    cells.write_png(montage_path, montage)
    composite_bytes = cv2.imencode(".png", composite, [cv2.IMWRITE_PNG_COMPRESSION, 6])[1].tobytes()
    composite_hash = hashlib.sha256(composite_bytes).hexdigest()
    workbook_hash = sha256_file(workbook_path)
    batch_path = output_root / GLOBAL_BATCH_FILENAME
    cached = _load_json(batch_path)
    cached_response = cached.get("cloud_response") if isinstance(cached.get("cloud_response"), dict) else {}
    cloud_response: dict[str, Any] = {}
    cloud_error = ""
    cache_reused = False
    if (
        not force
        and cached.get("parser_version") == PARSER_VERSION
        and cached.get("workbook_sha256") == workbook_hash
        and cached.get("composite_sha256") == composite_hash
        and cached_response.get("success")
    ):
        cloud_response = cached_response
        cache_reused = True
    elif ocr_ready and not prepare_only:
        if not cloud_url:
            raise ValueError("非 --prepare-only 模式必须提供 --cloud-url")
        try:
            response = requests.post(
                cells.classic_api_url(cloud_url),
                files={"file": (composite_path.name, composite_bytes, "image/png")},
                timeout=(30, timeout),
            )
            try:
                cloud_response = response.json()
            except ValueError as error:
                raise RuntimeError(f"经典 OCR 返回非 JSON：{response.text[:200]}") from error
            if not response.ok or not cloud_response.get("success"):
                raise RuntimeError(cloud_response.get("detail") or f"HTTP {response.status_code}")
        except Exception as error:
            cloud_error = str(error)
    diagnostics: dict[str, list[dict[str, Any]]] = {"mapped": [], "ambiguous": [], "unmapped": []}
    if cloud_response.get("success"):
        diagnostics = apply_response(rows, cloud_response, bands)
    elif cloud_error:
        for row in rows:
            if row["status"] == "ocr_pending":
                row["status"] = "ocr_error"
                row["reason_codes"] = ["classic_ocr_request_failed"]
                row["error"] = cloud_error
    status_counts: dict[str, int] = {}
    for row in rows:
        status_counts[row["status"]] = status_counts.get(row["status"], 0) + 1
    payload: dict[str, Any] = {
        "schema_version": SCHEMA_VERSION,
        "parser_version": PARSER_VERSION,
        "generated_at": now_text(),
        "job_dir": str(job_dir),
        "workbook": str(workbook_path),
        "workbook_sha256": workbook_hash,
        "scope": scope,
        "explicit_targets": list(explicit_targets or []),
        "risk_phrases": list(RISK_PHRASES),
        "prepare_only": prepare_only,
        "force": force,
        "variants": list(VARIANTS),
        "minimum_candidate_score": MIN_CANDIDATE_SCORE,
        "page1_reference": {"path": str(cells.PAGE1_REFERENCE.resolve()), "sha256": sha256_file(cells.PAGE1_REFERENCE)},
        "page2_reference": page2_metadata,
        "composite_image": str(composite_path.resolve()),
        "composite_sha256": composite_hash,
        "row_montage": str(montage_path.resolve()),
        "row_montage_sha256": sha256_file(montage_path),
        "row_montage_bands": montage_bands,
        "bands": bands,
        "cloud_url": cells.classic_api_url(cloud_url) if cloud_url else "",
        "cloud_response_reused": cache_reused,
        "cloud_response": cloud_response,
        "cloud_error": cloud_error,
        "mapped_items": diagnostics["mapped"],
        "ambiguous_items": diagnostics["ambiguous"],
        "unmapped_items": diagnostics["unmapped"],
        "target_count": len(rows),
        "ocr_target_count": len(ocr_ready),
        "status_counts": status_counts,
    }
    _write_results(job_dir, workbook_path, rows, payload)
    return {
        "parser_version": PARSER_VERSION,
        "job_dir": str(job_dir),
        "workbook": str(workbook_path),
        "scope": scope,
        "batch": str(batch_path.resolve()),
        "composite": str(composite_path.resolve()),
        "row_montage": str(montage_path.resolve()),
        "target_count": len(rows),
        "ocr_target_count": len(ocr_ready),
        "status_counts": status_counts,
        "cloud_response_reused": cache_reused,
        "cloud_error": cloud_error,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="增强食物平均每次食用量的 OCR 证据；只保存核对结果，不修改 Excel。")
    parser.add_argument("--job-dir", type=Path, required=True)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--cloud-url", help="云端基地址、/parse-file 或 /classic-ocr 地址。")
    parser.add_argument("--scope", choices=("flagged", "missing", "all"), default="flagged")
    parser.add_argument("--targets", nargs="*", help="显式 person:code 列表；提供后覆盖 --scope，例如 ocr/12:2.1。")
    parser.add_argument("--timeout", type=int, default=900)
    parser.add_argument("--force", action="store_true", help="重建参考并忽略缓存 OCR 响应。")
    parser.add_argument("--prepare-only", action="store_true", help="只生成双路裁剪、批次图和原始行拼图，不请求云端。")
    args = parser.parse_args()
    result = enhance_quantities(
        args.job_dir,
        args.workbook,
        args.cloud_url,
        scope=args.scope,
        explicit_targets=args.targets,
        force=args.force,
        prepare_only=args.prepare_only,
        timeout=args.timeout,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if result.get("cloud_error"):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
