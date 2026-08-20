from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import local_service as service  # noqa: E402


# Preserve already reviewed identity fields and meal-location choices when a saved
# output is supplied.  Fixed numeric questions are deliberately excluded so the
# focused classic OCR artifacts can still improve them on every rebuild.
SEED_GENERAL_FIELDS = (
    "受试者编号",
    "姓名",
    "调查日期",
    "访视号",
    "早餐地点",
    "午餐地点",
    "晚餐地点",
)
CLASSIC_GENERAL_FILENAME = "general_classic_ocr.json"


def load_seed_general(workbook_path: Path | None) -> dict[str, dict[str, str]]:
    if workbook_path is None or not workbook_path.is_file():
        return {}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = next((sheet for sheet in workbook.worksheets if "人员" in sheet.title), workbook.worksheets[0])
    header_row = service.detect_header_row(worksheet)
    headers = [service.clean(worksheet.cell(header_row, column).value) for column in range(1, worksheet.max_column + 1)]
    person_column = next((index for index, header in enumerate(headers) if header == "人员文件夹"), None)
    result: dict[str, dict[str, str]] = {}
    if person_column is not None:
        for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
            person_id = service.clean(values[person_column] if person_column < len(values) else "")
            if not person_id:
                continue
            row = {header: service.clean(values[index] if index < len(values) else "") for index, header in enumerate(headers) if header}
            result[person_id] = {field: row.get(field, "") for field in SEED_GENERAL_FIELDS if row.get(field)}
    workbook.close()
    return result


def merge_page(person: dict[str, Any], parsed: dict[str, Any], name_candidates: list[tuple[str, str]]) -> None:
    for key, value in parsed["general"].items():
        if key == "姓名":
            name = service.valid_person_name(value)
            if name:
                name_candidates.append((name, "saved_ocr"))
        elif key == "人工备注":
            person["general"][key] = service.append_nutrition_note(person["general"].get(key), value)
        elif value and not person["general"].get(key):
            person["general"][key] = value
    person["food_rows"].extend(parsed["food_rows"])
    person["supplement_rows"].extend(parsed["supplement_rows"])


def apply_classic_general(
    person: dict[str, Any],
    payload: dict[str, Any],
    name_candidates: list[tuple[str, str]],
) -> None:
    """Apply validated first-page/header and numeric crop fields after legacy OCR artifacts."""
    structured = payload.get("structured") or {}
    general_fields = structured.get("general_fields") or {}
    field_reviews = structured.get("field_reviews") or {}
    person["classic_general_review"] = field_reviews

    for field, review in field_reviews.items():
        if not isinstance(review, dict):
            continue
        status = service.clean(review.get("status"))
        if status not in {"conflict", "low_confidence", "invalid"}:
            continue
        raw_candidates = list(review.get("candidates") or []) + list(review.get("invalid_candidates") or [])
        candidate_values = list(
            dict.fromkeys(
                service.clean(candidate.get("value"))
                for candidate in raw_candidates
                if isinstance(candidate, dict) and service.clean(candidate.get("value"))
            )
        )
        detail = "、".join(candidate_values) or "无可用候选"
        label = {"conflict": "候选冲突", "low_confidence": "低置信度", "invalid": "超出合理范围"}[status]
        person["general"]["人工备注"] = service.append_nutrition_note(
            person["general"].get("人工备注"), f"首面经典 OCR {label}：{field}={detail}"
        )

    for field, raw_value in general_fields.items():
        if field not in service.NUTRITION_GENERAL_COLUMNS:
            continue
        value = service.clean(raw_value)
        if not value:
            continue
        if field == "姓名":
            name = service.valid_person_name(value)
            if name:
                # The classic identity band is built exclusively from the first PDF.
                name_candidates.append((name, "首份裁剪增强"))
            continue
        if field in {"人员文件夹", "人工备注"}:
            continue
        if field in SEED_GENERAL_FIELDS and service.clean(person["general"].get(field)):
            continue
        previous = service.clean(person["general"].get(field))
        if previous and previous != value:
            person["general"]["人工备注"] = service.append_nutrition_note(
                person["general"].get("人工备注"),
                f"整页/首面经典 OCR 冲突：{field}={previous}/{value}",
            )
            # 身份页眉可能把相邻手写字符粘连或漏掉一位。已有结构化身份值
            # 与经典裁剪不一致时保留原值并提示核对；固定数值题则采用更聚焦的裁剪值。
            if field in {"受试者编号", "调查日期", "访视号"}:
                continue
        person["general"][field] = value


def load_food_visual_audits(
    paths: Path | list[Path] | tuple[Path, ...] | None,
) -> dict[str, dict[str, dict[str, Any]]]:
    """Index explicit visual-audit files in order; later fields override earlier ones."""
    if paths is None:
        return {}
    audit_paths = [paths] if isinstance(paths, Path) else list(paths)
    indexed: dict[str, dict[str, dict[str, Any]]] = {}
    for path in audit_paths:
        if not path.is_file():
            continue
        payload = json.loads(path.read_text(encoding="utf-8"))
        rows = payload.get("rows") or payload.get("items") or payload.get("records") or []
        for row in rows:
            if not isinstance(row, dict) or row.get("apply") is False:
                continue
            person_id = service.clean(row.get("person"))
            code = service.clean(row.get("code"))
            if person_id and code:
                previous = indexed.setdefault(person_id, {}).get(code, {})
                indexed[person_id][code] = {**previous, **row}
    return indexed


# Backwards-compatible alias for callers that pass one audit path.
load_food_visual_audit = load_food_visual_audits


def load_food_cell_audits(paths: list[Path] | None) -> dict[str, dict[str, list[dict[str, Any]]]]:
    indexed: dict[str, dict[str, list[dict[str, Any]]]] = {}
    for path in paths or []:
        if not path.is_file():
            continue
        payload = json.loads(path.read_text(encoding="utf-8"))
        rows = payload.get("rows") or []
        for row in rows:
            if not isinstance(row, dict):
                continue
            person_id = service.clean(row.get("person"))
            code = service.clean(row.get("code"))
            if person_id and code:
                indexed.setdefault(person_id, {}).setdefault(code, []).append(row)
    return indexed


def strip_resolved_period_notes(value: object) -> str:
    resolved_markers = (
        "频率列疑似错位",
        "多个频率列同时有正数",
        "唯一正数与零/不吃标记并存",
        "仅识别到 0",
        "视觉复核：",
        "视觉候选=",
        "单元格 OCR：",
        "次数疑似异常",
    )
    parts = [part.strip() for part in service.clean(value).split("；") if part.strip()]
    return "；".join(part for part in parts if not any(marker in part for marker in resolved_markers))


def strip_resolved_quantity_notes(value: object) -> str:
    resolved_markers = (
        "食用量疑似粘连",
        "食用量疑似异常",
        "食用量没有有效数字",
        "食用量未识别到单位",
        "食用量仅识别到符号",
        "平均每次量仍缺少",
    )
    parts = [part.strip() for part in service.clean(value).split("；") if part.strip()]
    return "；".join(part for part in parts if not any(marker in part for marker in resolved_markers))


def raw_cell_count_candidates(value: object, period: str) -> list[str]:
    text = service.clean(value).translate(str.maketrans("一二三四五六七八九", "123456789"))
    text = text.replace("～", "-").replace("~", "-").replace("—", "-").replace("至", "-")
    maximum = {"每天": 10.0, "每周": 21.0, "每月": 31.0, "每年": 365.0}.get(period)
    result: list[str] = []
    for token in re.findall(r"\d+(?:\.\d+)?(?:\s*-\s*\d+(?:\.\d+)?)?", text):
        numbers = [float(item) for item in re.sub(r"\s+", "", token).split("-")]
        if not numbers or all(number == 0 for number in numbers):
            continue
        if len(numbers) == 2 and numbers[0] > numbers[1]:
            continue
        if maximum is not None and any(number > maximum for number in numbers):
            continue
        normalized = "-".join(str(int(number)) if number.is_integer() else f"{number:g}" for number in numbers)
        if normalized not in result:
            result.append(normalized)
    return result


def apply_food_visual_audit(person: dict[str, Any], indexed: dict[str, dict[str, dict[str, Any]]]) -> None:
    """Apply explicit final conclusions: blank, no, or a high-confidence checked value.

    Ordinary handwriting suggestions still wait for an independent cell OCR.  A
    ``confirmed_value`` is reserved for a separately retained human/visual audit
    with confidence >= 0.9, so it cannot arise from an OCR candidate by accident.
    """
    records = indexed.get(person["id"], {})
    if not records:
        return
    rows_by_code = {service.clean(row.get("食物编号")): row for row in person.get("food_rows") or []}
    review: list[dict[str, Any]] = []
    for code, record in records.items():
        row = rows_by_code.get(code)
        if row is None:
            continue
        visual_class = service.clean(record.get("visual_class"))
        if visual_class:
            conclusion = service.clean(record.get("visual_conclusion"))
            if visual_class == "confirmed_period_blank":
                confidence = float(record.get("confidence") or 0.0)
                if confidence < 0.9:
                    raise ValueError(
                        f"Unsafe confirmed_period_blank audit for {person['id']} food {code}: "
                        f"confidence={confidence!r}"
                    )
                row["次数"] = ""
                row["频率周期(请核对)"] = ""
                row["是否不吃"] = ""
                row["人工核对备注"] = strip_resolved_period_notes(row.get("人工核对备注"))
                status = "confirmed_source_period_blank"
            elif visual_class == "true_blank":
                row["次数"] = ""
                row["频率周期(请核对)"] = ""
                row["是否不吃"] = ""
                row["人工核对备注"] = ""
                status = "confirmed_blank"
            elif visual_class == "explicit_no":
                row["次数"] = "0"
                row["频率周期(请核对)"] = "不吃"
                row["是否不吃"] = "是"
                suggested_quantity = service.clean(record.get("suggested_quantity"))
                if suggested_quantity:
                    row["平均每次食用量"] = suggested_quantity
                row["人工核对备注"] = ""
                status = "accepted_no"
            elif visual_class == "confirmed_value":
                period = service.clean(record.get("suggested_period"))
                count = service.nutrition_count_value(record.get("suggested_count"))
                confidence = float(record.get("confidence") or 0.0)
                if period not in service.NUTRITION_PERIODS or not count or confidence < 0.9:
                    raise ValueError(
                        f"Unsafe confirmed_value audit for {person['id']} food {code}: "
                        f"period={period!r}, count={count!r}, confidence={confidence!r}"
                    )
                row["频率周期(请核对)"] = period
                row["次数"] = count
                row["是否不吃"] = "是" if period == "不吃" else ""
                suggested_quantity = service.clean(record.get("suggested_quantity"))
                if suggested_quantity and suggested_quantity != "待补":
                    row["平均每次食用量"] = suggested_quantity
                    row["人工核对备注"] = strip_resolved_quantity_notes(row.get("人工核对备注"))
                row["人工核对备注"] = strip_resolved_period_notes(row.get("人工核对备注"))
                status = "accepted_visual_confirmation"
            else:
                suggestion = "/".join(
                    value
                    for value in (
                        service.clean(record.get("suggested_period")),
                        service.clean(record.get("suggested_count")),
                    )
                    if value
                )
                detail = conclusion or "存在手写或多列冲突"
                if suggestion:
                    detail = f"{detail}；视觉候选={suggestion}"
                row["人工核对备注"] = service.append_nutrition_note(
                    row.get("人工核对备注"), f"视觉复核：{detail}"
                )
                status = "awaiting_cell_ocr" if suggestion else "manual_review"
            review.append(
                {
                    "code": code,
                    "status": status,
                    "visual_class": visual_class,
                    "suggested_period": service.clean(record.get("suggested_period")),
                    "suggested_count": service.clean(record.get("suggested_count")),
                    "confidence": record.get("confidence"),
                    "conclusion": conclusion,
                }
            )

        quantity_class = service.clean(record.get("quantity_class"))
        if quantity_class:
            quantity = service.clean(record.get("suggested_quantity"))
            confidence = float(record.get("quantity_confidence", record.get("confidence")) or 0.0)
            if quantity_class == "confirmed_quantity":
                if not quantity or confidence < 0.9:
                    raise ValueError(
                        f"Unsafe confirmed_quantity audit for {person['id']} food {code}: "
                        f"quantity={quantity!r}, confidence={confidence!r}"
                    )
                row["平均每次食用量"] = quantity
                row["人工核对备注"] = strip_resolved_quantity_notes(row.get("人工核对备注"))
                quantity_status = "accepted_visual_quantity"
            elif quantity_class == "confirmed_quantity_blank":
                if confidence < 0.9:
                    raise ValueError(
                        f"Unsafe confirmed_quantity_blank audit for {person['id']} food {code}: "
                        f"confidence={confidence!r}"
                    )
                row["平均每次食用量"] = ""
                row["人工核对备注"] = strip_resolved_quantity_notes(row.get("人工核对备注"))
                quantity_status = "confirmed_source_blank"
            else:
                raise ValueError(f"Unknown quantity_class for {person['id']} food {code}: {quantity_class!r}")
            review.append(
                {
                    "code": code,
                    "status": quantity_status,
                    "quantity_class": quantity_class,
                    "suggested_quantity": quantity,
                    "confidence": confidence,
                    "conclusion": service.clean(record.get("quantity_conclusion")),
                }
            )
    person["food_visual_audit"] = review


def apply_food_cell_audit(
    person: dict[str, Any],
    visual_index: dict[str, dict[str, dict[str, Any]]],
    cell_index: dict[str, dict[str, list[dict[str, Any]] | dict[str, Any]]],
) -> None:
    """Accept a value only when visual and independent cell OCR evidence agree."""
    visual_records = visual_index.get(person["id"], {})
    cell_records = cell_index.get(person["id"], {})
    if not visual_records or not cell_records:
        return
    rows_by_code = {service.clean(row.get("食物编号")): row for row in person.get("food_rows") or []}
    audit_rows: list[dict[str, Any]] = []
    for code, visual in visual_records.items():
        raw_cells = cell_records.get(code)
        row = rows_by_code.get(code)
        if raw_cells is None or row is None or service.clean(visual.get("visual_class")) != "handwriting_or_conflict":
            continue
        cells = raw_cells if isinstance(raw_cells, list) else [raw_cells]

        visual_period = service.clean(visual.get("suggested_period"))
        visual_count = service.nutrition_count_value(visual.get("suggested_count"))
        visual_confidence = float(visual.get("confidence") or 0.0)
        evidence: list[dict[str, Any]] = []
        supporting: dict[str, Any] | None = None
        contradictory = False
        matching_variants = 0
        for cell in cells:
            cell_status = service.clean(cell.get("status"))
            cell_period = service.clean(cell.get("selected_period"))
            cell_count = service.nutrition_count_value(cell.get("count"))
            reliable = cell.get("reliable_counts") if isinstance(cell.get("reliable_counts"), dict) else {}
            reliable_values = [service.nutrition_count_value(value) for value in reliable.get(visual_period, [])]
            reliable_values = [value for value in reliable_values if value]
            exact_selected = (
                cell_status == "accepted"
                and cell_period == visual_period
                and cell_count == visual_count
                and bool(visual_period and visual_count)
            )
            exact_candidate = bool(visual_period and visual_count and visual_count in reliable_values)
            text_by_period = cell.get("text_by_period") if isinstance(cell.get("text_by_period"), dict) else {}
            raw_values = raw_cell_count_candidates(text_by_period.get(visual_period), visual_period)
            raw_match = bool(visual_count and visual_count in raw_values)
            if exact_candidate or raw_match:
                matching_variants += 1
            supports = (
                exact_selected
                or (visual_confidence >= 0.85 and exact_candidate)
                or (visual_confidence >= 0.90 and raw_match)
            )
            if supporting is None and supports:
                supporting = cell
            if cell_status == "accepted" and (cell_period, cell_count) != (visual_period, visual_count):
                contradictory = True
            if any(value != visual_count for value in reliable_values):
                contradictory = True
            classic_display = "、".join(
                f"{period}={service.clean(text_by_period.get(period))}"
                for period in service.NUTRITION_PERIODS
                if service.clean(text_by_period.get(period))
            )
            evidence.append(
                {
                    "variant": service.clean(cell.get("image_variant")) or "ink",
                    "status": cell_status,
                    "period": cell_period,
                    "count": cell_count,
                    "reliable_counts": reliable,
                    "raw_counts_in_visual_period": raw_values,
                    "classic_text": classic_display,
                }
            )
        if supporting is None and matching_variants >= 2 and visual_period and visual_count:
            supporting = cells[0]
        accepted = supporting is not None and not contradictory
        classic_display = " | ".join(
            f"{item['variant']}:{item['classic_text'] or item['status']}" for item in evidence
        )
        if accepted:
            row["频率周期(请核对)"] = visual_period
            row["次数"] = visual_count
            row["是否不吃"] = "是" if visual_period == "不吃" else ""
            suggested_quantity = service.clean(visual.get("suggested_quantity"))
            if suggested_quantity and suggested_quantity != "待补":
                row["平均每次食用量"] = suggested_quantity
            row["人工核对备注"] = strip_resolved_period_notes(row.get("人工核对备注"))
            quantity_confirmed_blank = service.clean(visual.get("quantity_class")) == "confirmed_quantity_blank"
            if (
                visual_period in service.NUTRITION_PERIODS[:4]
                and not quantity_confirmed_blank
                and not re.search(r"\d", service.clean(row.get("平均每次食用量")))
            ):
                row["人工核对备注"] = service.append_nutrition_note(
                    row.get("人工核对备注"), "单元格增强已确定频率，但食用量没有有效数字"
                )
            status = "accepted_agreement"
        else:
            detail = classic_display or "未提取到可靠数字"
            row["人工核对备注"] = service.append_nutrition_note(
                row.get("人工核对备注"), f"单元格 OCR：{detail}"
            )
            status = "manual_review" if not visual_period else "evidence_mismatch"
        audit_rows.append(
            {
                "code": code,
                "status": status,
                "visual_period": visual_period,
                "visual_count": visual_count,
                "visual_confidence": visual_confidence,
                "supporting_variant": service.clean(supporting.get("image_variant")) if supporting else "",
                "contradictory": contradictory,
                "matching_variants": matching_variants,
                "cell_evidence": evidence,
                "classic_text": classic_display,
            }
        )
    person["food_cell_audit"] = audit_rows


def rebuild(
    job_dir: Path,
    output_path: Path,
    parse_mode: str = "markdown",
    template_path: Path | None = None,
    seed_output: Path | None = None,
    food_visual_audits: list[Path] | None = None,
    food_cell_audits: list[Path] | None = None,
) -> dict[str, Any]:
    excel_root = job_dir / "ocr_excel"
    markdown_root = job_dir / "ocr_markdown"
    json_root = job_dir / "ocr_json"
    if not excel_root.is_dir():
        raise FileNotFoundError(f"Missing OCR Excel directory: {excel_root}")

    grouped: dict[str, list[Path]] = {}
    for excel_path in sorted(excel_root.rglob("*.xlsx")):
        relative = excel_path.relative_to(excel_root)
        grouped.setdefault(relative.parent.as_posix(), []).append(excel_path)
    if not grouped:
        raise RuntimeError(f"No OCR Excel files found under {excel_root}")

    seeds = load_seed_general(seed_output)
    visual_audit = load_food_visual_audits(food_visual_audits)
    cell_audit = load_food_cell_audits(food_cell_audits)
    people: list[dict[str, Any]] = []
    for person_id, excel_paths in sorted(grouped.items()):
        person = service.nutrition_blank_person(person_id)
        person["general"].update(seeds.get(person_id, {}))
        name_candidates: list[tuple[str, str]] = []
        source_files: list[str] = []
        for page_index, excel_path in enumerate(excel_paths, start=1):
            relative = excel_path.relative_to(excel_root)
            markdown_path = markdown_root / relative.with_suffix(".md")
            markdown_text = markdown_path.read_text(encoding="utf-8") if markdown_path.is_file() else ""
            parsed = service.parse_nutrition_page_fields(excel_path.read_bytes(), markdown_text, parse_mode)
            name_candidates.extend(
                service.first_page_name_candidates(service.extract_person_name_candidates(markdown_text), page_index)
            )

            json_path = json_root / relative.with_suffix(".json")
            if json_path.is_file():
                cloud_data = json.loads(json_path.read_text(encoding="utf-8"))
                name_candidates.extend(
                    service.first_page_name_candidates(service.cloud_name_candidates(cloud_data), page_index)
                )
                service.apply_cloud_questionnaire_result(cloud_data, person)
            merge_page(person, parsed, name_candidates)
            source_files.append(relative.as_posix())

        enhanced_dir = job_dir / "ocr_enhanced" / Path(person_id)
        for enhanced_markdown in sorted(enhanced_dir.glob("*_ocr.md")):
            enhanced_text = enhanced_markdown.read_text(encoding="utf-8")
            plain_text = BeautifulSoup(enhanced_text, "html.parser").get_text("\n", strip=True)
            enhanced_general = service.nutrition_general_from_text(plain_text)
            for name, _ in service.extract_person_name_candidates(enhanced_text):
                name_candidates.append((name, "裁剪增强"))
            for key, value in enhanced_general.items():
                if key == "姓名":
                    name = service.valid_person_name(value)
                    if name:
                        name_candidates.append((name, "裁剪增强"))
                elif key == "人工备注":
                    person["general"][key] = service.append_nutrition_note(person["general"].get(key), value)
                elif value and key not in SEED_GENERAL_FIELDS:
                    previous = service.clean(person["general"].get(key))
                    if previous and previous != value:
                        person["general"]["人工备注"] = service.append_nutrition_note(
                            person["general"].get("人工备注"), f"整页/裁剪 OCR 冲突：{key}={previous}/{value}（采用裁剪值）"
                        )
                    person["general"][key] = value
            source_files.append(enhanced_markdown.relative_to(job_dir).as_posix())

        classic_json = enhanced_dir / CLASSIC_GENERAL_FILENAME
        if classic_json.is_file():
            classic_payload = json.loads(classic_json.read_text(encoding="utf-8"))
            apply_classic_general(person, classic_payload, name_candidates)
            source_files.append(classic_json.relative_to(job_dir).as_posix())

        selected_name = service.choose_person_name(name_candidates)
        if selected_name:
            person["general"]["姓名"] = selected_name
        service.finalize_nutrition_person(person)
        apply_food_visual_audit(person, visual_audit)
        apply_food_cell_audit(person, visual_audit, cell_audit)
        person["source_files"] = source_files
        people.append(person)

    state = {"template_path": template_path, "job_dir": job_dir, "output_filename": output_path.name}
    service.write_nutrition_workbook(state, people, output_path=output_path)

    general_filled = {
        field: sum(bool(service.clean(person["general"].get(field))) for person in people)
        for field in service.NUTRITION_GENERAL_COLUMNS
        if field not in {"人员文件夹", "人工备注"}
    }
    food_rows = [row for person in people for row in person["food_rows"]]
    supplement_rows = [row for person in people for row in person["supplement_rows"]]
    result = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "job_dir": str(job_dir.resolve()),
        "output": str(output_path.resolve()),
        "parse_mode": parse_mode,
        "template": str(template_path.resolve()) if template_path else str(service.NUTRITION_TEMPLATE.resolve()),
        "seed_output": str(seed_output.resolve()) if seed_output else None,
        "food_visual_audits": [str(path.resolve()) for path in food_visual_audits or []],
        "food_cell_audits": [str(path.resolve()) for path in food_cell_audits or []],
        "people_count": len(people),
        "source_files": sum(len(person["source_files"]) for person in people),
        "food_rows": len(food_rows),
        "supplement_rows": len(supplement_rows),
        "food_periods": dict(Counter(service.clean(row.get("频率周期(请核对)")) for row in food_rows)),
        "supplement_periods": dict(Counter(service.clean(row.get("频率周期(请核对)")) for row in supplement_rows)),
        "food_review_rows": sum(service.nutrition_food_note_requires_review(row.get("人工核对备注")) for row in food_rows),
        "supplement_review_rows": sum(bool(service.clean(row.get("备注"))) for row in supplement_rows),
        "general_filled": general_filled,
        "people": [
            {
                "id": person["id"],
                "general": person["general"],
                "unrecognized_food_codes": [row["食物编号"] for row in person["food_rows"] if row["频率周期(请核对)"] == "未识别"],
                "food_review_codes": [
                    row["食物编号"]
                    for row in person["food_rows"]
                    if service.nutrition_food_note_requires_review(row.get("人工核对备注"))
                ],
                "supplement_review": [row["保健品种类"] for row in person["supplement_rows"] if service.clean(row.get("备注"))],
                "classic_general_review": person.get("classic_general_review", {}),
                "food_visual_audit": person.get("food_visual_audit", []),
                "food_cell_audit": person.get("food_cell_audit", []),
                "source_files": person["source_files"],
            }
            for person in people
        ],
    }
    audit_path = output_path.with_suffix(".audit.json")
    audit_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    result["audit"] = str(audit_path.resolve())
    return result


def find_latest_output(job_dir: Path) -> Path | None:
    candidates = sorted((job_dir / "output").glob("*.xlsx"), key=lambda path: path.stat().st_mtime, reverse=True)
    return candidates[0] if candidates else None


def main() -> None:
    parser = argparse.ArgumentParser(description="Rebuild a nutrition questionnaire summary from locally saved OCR artifacts.")
    parser.add_argument("--job-dir", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--template", type=Path)
    parser.add_argument("--seed-output", type=Path, help="Preserve structured meal-location values from an earlier output.")
    parser.add_argument(
        "--food-visual-audit",
        type=Path,
        action="append",
        help="Apply a visual audit JSON; repeat to layer later confirmations over earlier audits.",
    )
    parser.add_argument(
        "--food-cell-audit",
        type=Path,
        action="append",
        help="Cross-check visual candidates with a classic cell OCR batch JSON; repeat for multiple image variants.",
    )
    parser.add_argument("--no-auto-seed", action="store_true")
    parser.add_argument("--parse-mode", choices=("markdown", "excel"), default="markdown")
    args = parser.parse_args()

    job_dir = args.job_dir.resolve()
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = (args.output or (job_dir / "output" / f"食物频率调查_本地增强重建_{timestamp}.xlsx")).resolve()
    seed_output = args.seed_output.resolve() if args.seed_output else (None if args.no_auto_seed else find_latest_output(job_dir))
    result = rebuild(
        job_dir=job_dir,
        output_path=output_path,
        parse_mode=args.parse_mode,
        template_path=args.template.resolve() if args.template else None,
        seed_output=seed_output,
        food_visual_audits=[path.resolve() for path in args.food_visual_audit or []],
        food_cell_audits=[path.resolve() for path in args.food_cell_audit or []],
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
