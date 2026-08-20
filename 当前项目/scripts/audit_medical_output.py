from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def audit(workbook_path: Path) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook.active
    header_row = max(
        range(1, min(30, worksheet.max_row) + 1),
        key=lambda row: sum(bool(clean(worksheet.cell(row, column).value)) for column in range(1, worksheet.max_column + 1)),
    )
    headers = [clean(worksheet.cell(header_row, column).value) for column in range(1, worksheet.max_column + 1)]
    name_column = next((index + 1 for index, header in enumerate(headers) if header == "姓名"), None)
    if name_column is None:
        workbook.close()
        raise RuntimeError("Could not find the 姓名 column")

    data_rows = [
        row
        for row in range(header_row + 1, worksheet.max_row + 1)
        if clean(worksheet.cell(row, name_column).value)
    ]
    columns: list[dict[str, Any]] = []
    for column, header in enumerate(headers, start=1):
        if not header or header == "序号":
            continue
        values = [clean(worksheet.cell(row, column).value) for row in data_rows]
        missing_rows = [row for row, value in zip(data_rows, values) if not value]
        columns.append(
            {
                "column": worksheet.cell(header_row, column).column_letter,
                "header": header,
                "filled": len(values) - len(missing_rows),
                "total": len(values),
                "missing_rows": missing_rows,
                "all_empty": bool(values) and len(missing_rows) == len(values),
            }
        )
    workbook.close()

    total_cells = sum(column["total"] for column in columns)
    filled_cells = sum(column["filled"] for column in columns)
    return {
        "workbook": str(workbook_path.resolve()),
        "sheet": worksheet.title,
        "header_row": header_row,
        "data_rows": data_rows,
        "people_count": len(data_rows),
        "columns_count": len(columns),
        "filled": filled_cells,
        "total": total_cells,
        "missing": total_cells - filled_cells,
        "completeness": round(filled_cells / total_cells, 6) if total_cells else 0,
        "all_empty_columns": [column["header"] for column in columns if column["all_empty"]],
        "partial_columns": [
            column["header"] for column in columns if column["missing_rows"] and not column["all_empty"]
        ],
        "columns": columns,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Report column-level completeness for a medical summary workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--json-out", type=Path)
    args = parser.parse_args()
    result = audit(args.workbook)
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(payload, encoding="utf-8")
    print(payload)


if __name__ == "__main__":
    main()
