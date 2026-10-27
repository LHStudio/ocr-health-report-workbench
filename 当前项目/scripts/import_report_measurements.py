"""Import only unique exact roster matches from InBody and measured REE workbooks."""
import argparse
from collections import Counter
from datetime import date, datetime
import hashlib
import io
import json
import math
from pathlib import Path
import re
import zipfile

import openpyxl

BODY_HEADERS = {
    'height': 'Height', 'date_of_birth': 'Date of Birth', 'age': 'Age',
    'weight': 'Weight', 'bfm': 'BFM (Body Fat Mass)', 'ffm': 'FFM (Fat Free Mass)',
    'bmi': 'BMI (Body Mass Index)', 'pbf': 'PBF (Percent Body Fat)',
    'vfl': 'VFL (Visceral Fat Level)', 'vfa': 'VFA (Visceral Fat Area)', 'asmi': 'ASMI (kg/m²)',
}


def digest(data):
    return hashlib.sha256(data).hexdigest()


def clean(value, key=''):
    if value is None or str(value).strip() in ('', '-', '--', 'N/A'):
        return None
    if key == 'date_of_birth':
        if isinstance(value, (date, datetime)):
            return value.strftime('%Y-%m-%d')
        parts = re.fullmatch(r'(\d{4})[./-](\d{1,2})[./-](\d{1,2})\.?', str(value).strip())
        if not parts:
            raise ValueError('无法读取出生日期')
        return date(*map(int, parts.groups())).isoformat()
    value = float(re.sub(r'^Level\s*', '', str(value).strip()))
    if not math.isfinite(value):
        raise ValueError('测量值不是有限数')
    return value


def extract(job, inbody, metabolism):
    people = job['people']
    counts = Counter(str(p.get('general', {}).get('姓名') or '').strip() for p in people)
    roster = {str(p['general'].get('姓名') or '').strip(): p for p in people
              if counts[str(p['general'].get('姓名') or '').strip()] == 1}
    records = {}
    stats = {'roster_people': len(people), 'body_imported': 0, 'metabolism_imported': 0,
             'body_discarded': 0, 'metabolism_discarded': 0}

    def record(name):
        p = roster[name]
        return records.setdefault(name, {'name': name, 'person_id': p['id'],
            'subject_id': str(p['general'].get('受试者编号') or '').strip()})

    content = Path(inbody).read_bytes()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    sheet = wb['InBody']
    headers = {re.sub(r'^\d+\.\s*', '', str(c.value)): c.column for c in sheet[1]}
    cols = {key: headers[label] for key, label in BODY_HEADERS.items()}
    name_col = headers['Name']
    source_counts = Counter(str(sheet.cell(r, name_col).value or '').strip()
                            for r in range(2, sheet.max_row + 1))
    for r in range(2, sheet.max_row + 1):
        name = str(sheet.cell(r, name_col).value or '').strip()
        if name not in roster or source_counts[name] != 1:
            stats['body_discarded'] += 1
            continue
        cells = {key: sheet.cell(r, col) for key, col in cols.items()}
        values = {key: clean(cell.value, key) for key, cell in cells.items()}
        record(name)['body_composition'] = {'values': values, 'source': {
            'filename': Path(inbody).name, 'sha256': digest(content), 'sheet': sheet.title,
            'cells': {key: cell.coordinate for key, cell in cells.items()},
            'raw_values': {key: str(cell.value) if cell.value is not None else None for key, cell in cells.items()},
            'test_date': str(sheet.cell(r, headers['Test Date / Time']).value)
            if 'Test Date / Time' in headers else None}}
        stats['body_imported'] += 1
    wb.close()
    with zipfile.ZipFile(metabolism) as archive:
        for info in archive.infolist():
            if not info.filename.lower().endswith('.xlsx'):
                continue
            content = archive.read(info)
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            source = wb['数据']
            # Internal name is authoritative; archive filenames may have prefixes.
            name = str(source['B2'].value or '').strip()
            if name not in roster:
                stats['metabolism_discarded'] += 1
                wb.close()
                continue
            if 'resting_metabolism' in record(name):
                raise ValueError('同名静息代谢结果重复，拒绝自动覆盖')
            sheet = wb['结果']
            values, cells = {}, {}
            measured_col = None
            for row in sheet:
                for cell in row:
                    if cell.value == '实测值':
                        measured_col = cell.column
                key = str(row[0].value or '').strip()
                if key in ('REE', 'FAT%', 'CHO%', 'PRO%'):
                    if measured_col is None or key in values:
                        raise ValueError('实测值列缺失或参数重复')
                    cell = sheet.cell(row[0].row, measured_col)
                    values[key] = clean(cell.value)
                    cells[key] = cell.coordinate
            if set(values) != {'REE', 'FAT%', 'CHO%', 'PRO%'}:
                raise ValueError('静息代谢参数不完整')
            try:
                member = info.filename.encode('cp437').decode('utf-8')
            except (UnicodeError, LookupError):
                member = info.filename
            record(name)['resting_metabolism'] = {'values': values, 'source': {
                'archive': Path(metabolism).name, 'member': member, 'sha256': digest(content),
                'sheet': '结果', 'column': '实测值', 'cells': cells,
                'test_date': str(source['E1'].value)}}
            stats['metabolism_imported'] += 1
            wb.close()
    return {'schema_version': 1, 'job_id': job['job_id'], 'records': list(records.values()), 'counts': stats}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for flag in ('job-state', 'inbody', 'metabolism', 'output'):
        parser.add_argument('--' + flag, required=True, type=Path)
    args = parser.parse_args()
    result = extract(json.loads(args.job_state.read_text(encoding='utf-8')), args.inbody, args.metabolism)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open('x', encoding='utf-8', newline='\n') as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    assert json.loads(args.output.read_text(encoding='utf-8')) == result
    print(json.dumps(result['counts'], ensure_ascii=False))


if __name__ == '__main__':
    main()
