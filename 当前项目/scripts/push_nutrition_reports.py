"""Push existing nutrition XLSX exports without importing report generators.

Dry-run is the default. Set SURVEY_OPEN_TOKEN and use --commit to transmit.
Checkpoint, payloads and receipts contain personal data; keep them outside Git.
"""
from __future__ import annotations

import argparse
from collections import Counter
from datetime import date, datetime
import hashlib
import json
import math
import os
from pathlib import Path
import re
import sys
from urllib.parse import urlsplit, urlunsplit

import openpyxl
import requests


HEADERS = {
    '个人信息': ['项目', '内容'],
    # The deployed UI reads these exact compact keys. PDF extraction inserts
    # spaces before units; accepting those as input must not change output keys.
    '食物明细': ['食物名称', '分类', '摄入量(g)', '能量(kcal)', '蛋白质(g)', '脂肪(g)', '碳水(g)', '钙(mg)', '数据来源'],
    '分类评估': ['分类', '摄入原重(g)', '评价', '奶类鲜奶等价量(g)', '评价等价食物', '评价等价量(g)'],
    '营养总计': ['项目', '摄入量', '单位', '评价'],
    'OCR换算说明': ['级别', '代码', '食物', '说明'],
    '运动员专项评价': ['项目', '实际摄入', '单位', '参考目标', '差值', '最终评价', '口径'],
    '建议': ['序号', '建议内容'],
}


def encoded(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False).encode('utf-8')


def digest(value):
    return hashlib.sha256(encoded(value)).hexdigest()


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + '.tmp')
    with temporary.open('w', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write('\n')
        stream.flush()
        os.fsync(stream.fileno())
    temporary.replace(path)


def normalize_date(value, label):
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-%d')
    text = str(value or '').strip()
    text = re.sub(r'^(\d{4})年(\d{1,2})月(\d{1,2})日$', r'\1-\2-\3', text)
    match = re.fullmatch(r'(\d{4})[-./](\d{1,2})[-./](\d{1,2})', text)
    if not match:
        raise ValueError(f'{label}缺失或格式无效：{text!r}')
    return date(*map(int, match.groups())).isoformat()


def endpoint_url(base):
    parts = urlsplit(base.strip())
    if (parts.scheme != 'https' or not parts.hostname or parts.username or parts.password
            or parts.query or parts.fragment or parts.path.rstrip('/') not in ('', '/open/nutrition', '/api', '/api/open/nutrition')):
        raise ValueError('平台地址必须为 HTTPS 根地址、/api 或其 /open/nutrition，不允许凭证、查询参数或其他路径')
    prefix = '/api' if parts.path.rstrip('/') in ('/api', '/api/open/nutrition') else ''
    return urlunsplit((parts.scheme, parts.netloc, prefix + '/open/nutrition', '', ''))


def _key(value):
    return re.sub(r'\s+', '', str(value or ''))


def _value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError('报告包含非有限数值')
    return value


def read_report(path, birthday=None):
    """Read literal cells; no recalculation, identity guessing or date invention."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    payload = {}
    try:
        for title, headers in HEADERS.items():
            if title not in wb.sheetnames:
                if title in ('个人信息', '营养总计'):
                    raise ValueError(f'报告缺少工作表：{title}')
                payload[title] = []
                continue
            rows = list(wb[title].iter_rows())
            for row in rows:
                if any(cell.data_type == 'f' for cell in row):
                    raise ValueError(f'{title}包含公式，请使用已导出的数值报告')
            values = [[_value(c.value) for c in row] for row in rows]
            required = headers[:2]
            header_index = next((i for i, row in enumerate(values)
                                 if all(_key(h) in [_key(v) for v in row] for h in required)), None)
            if header_index is None:
                raise ValueError(f'{title}缺少必要表头：{required}')
            mapping = {_key(h): h for h in headers}
            names = [mapping.get(_key(v)) for v in values[header_index]]
            nonempty_names = [n for n in names if n]
            if len(set(nonempty_names)) != len(nonempty_names):
                raise ValueError(f'{title}存在重复表头')
            # Unknown non-empty columns must not be silently lost.
            if any(v not in (None, '') and n is None for v, n in zip(values[header_index], names)):
                raise ValueError(f'{title}存在接口未定义的表头')
            data = [{n: v for n, v in zip(names, row) if n is not None}
                    for row in values[header_index + 1:] if any(v not in (None, '') for v in row)]
            if len(data) > 500:
                raise ValueError(f'{title}超过500行')
            if title == '个人信息':
                info = {}
                for row in data:
                    key = str(row['项目'] or '').strip()
                    if not key or key in info:
                        raise ValueError('个人信息存在空项目或重复项目')
                    info[key] = row['内容']
                payload[title] = info
            else:
                payload[title] = data
    finally:
        wb.close()
    info = payload['个人信息']
    name = str(info.get('姓名') or '').strip()
    if not name:
        raise ValueError('姓名不能为空')
    payload['name'] = name
    payload['analysisDate'] = normalize_date(info.get('日期'), '报告日期')
    info['日期'] = payload['analysisDate']
    existing = info.get('出生日期')
    if birthday and existing and normalize_date(birthday, '出生日期') != normalize_date(existing, '出生日期'):
        raise ValueError('出生日期与报告冲突')
    if birthday or existing:
        payload['birthday'] = normalize_date(birthday or existing, '出生日期')
    if not any(payload[k] for k in HEADERS if k not in ('个人信息', 'OCR换算说明')):
        raise ValueError('请提供营养报告内容')
    encoded(payload)
    return payload


def prepare(files, birthdays):
    items, identities = [], {}
    for path in sorted(map(Path, files)):
        absolute = str(path.resolve())
        payload = read_report(path, birthdays.get(absolute))
        key = (payload['name'], payload['analysisDate'])
        birthday = payload.get('birthday')
        previous = identities.setdefault(key, [])
        if previous and (not birthday or None in previous or birthday in previous):
            raise ValueError(f'同一人员日期重复且不能唯一消歧：{key}')
        previous.append(birthday)
        items.append({'source': absolute, 'source_sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'payload': payload})
    if not items:
        raise ValueError('没有找到营养报告 XLSX')
    unknown = set(birthdays) - {item['source'] for item in items}
    if unknown:
        raise ValueError('出生日期映射包含不属于本批次的文件')
    return items


def post(endpoint, token, payload):
    response = requests.post(endpoint, headers={'X-Open-Token': token, 'Content-Type': 'application/json'},
                             data=encoded(payload), timeout=(10, 45), allow_redirects=False)
    try:
        body = response.json()
    except ValueError:
        body = {'invalid_json': True, 'body_preview': response.text[:1000]}
    return response.status_code, body


def valid_receipt(body, payload):
    if not isinstance(body, dict) or type(body.get('code')) is not int or body['code'] != 0:
        return False
    data = body.get('data')
    return (isinstance(data, dict) and isinstance(data.get('athleteId'), str) and bool(data['athleteId'])
            and isinstance(data.get('analysisId'), str) and bool(data['analysisId'])
            and str(data.get('name', '')).strip() == payload['name']
            and data.get('analysisDate') == payload['analysisDate'] and type(data.get('created')) is bool)


def execute(items, endpoint, token, audit_dir, *, send=post, retry_failed=False):
    endpoint = endpoint_url(endpoint)
    if not token or '\n' in token or '\r' in token:
        raise ValueError('请设置有效的 SURVEY_OPEN_TOKEN')
    directory = Path(audit_dir)
    directory.mkdir(parents=True, exist_ok=True)
    # Exclusive local batch lock prevents concurrent duplicate transmissions.
    lock = directory / 'push.lock'
    try:
        handle = lock.open('x')
    except FileExistsError:
        raise ValueError('批次正在运行或上次中断，请人工核对 push.lock 后再继续') from None
    try:
        handle.write(str(os.getpid()))
        handle.close()
        checkpoint = directory / 'checkpoint.json'
        batch = digest({'endpoint': endpoint, 'items': items})
        state = json.loads(checkpoint.read_text(encoding='utf-8')) if checkpoint.exists() else {
            'batch': batch, 'endpoint': endpoint, 'records': {}}
        if state['batch'] != batch:
            raise ValueError('批次内容或目标地址已改变，请使用新的审计目录')
        for item in items:
            key = digest({'source': item['source'], 'payload': item['payload']})
            prior = state['records'].get(key)
            if prior and not (retry_failed and prior['status'] == 'failed'):
                continue
            record = {k: v for k, v in item.items() if k != 'payload'}
            record.update(name=item['payload']['name'], analysisDate=item['payload']['analysisDate'],
                          payload_sha256=digest(item['payload']), status='sending')
            state['records'][key] = record
            write_json(directory / 'payloads' / f'{key}.json', item['payload'])
            write_json(checkpoint, state)  # Crash after this point means unknown, never blind retry.
            try:
                status, response = send(endpoint, token, item['payload'])
                # Remove the configured secret if a server echoes it in an error.
                response = json.loads(json.dumps(response, ensure_ascii=False).replace(token, '[REDACTED]'))
                record.update(http_status=status, response=response)
                if status == 200 and valid_receipt(response, item['payload']):
                    record['status'] = 'success'
                elif status == 401 or (status == 200 and isinstance(response, dict) and type(response.get('code')) is int and response['code'] == 1):
                    record['status'] = 'failed'
                else:
                    record['status'] = 'uncertain'
            except Exception as exc:
                record.update(status='uncertain', error=str(exc).replace(token, '[REDACTED]'), error_type=type(exc).__name__)
            write_json(checkpoint, state)
            if record.get('http_status') == 401 or record['status'] == 'uncertain':
                break
        counts = Counter(r['status'] for r in state['records'].values())
        summary = {'total': len(items), 'success': counts['success'], 'failed': counts['failed'],
                   'uncertain': counts['uncertain'] + counts['sending'], 'pending': len(items) - len(state['records']),
                   'endpoint': endpoint, 'checkpoint': str(checkpoint.resolve())}
        write_json(directory / 'summary.json', summary)
        return summary
    finally:
        if not handle.closed:
            handle.close()
        lock.unlink()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True, type=Path, help='One exported XLSX or a directory of reports')
    parser.add_argument('--base-url', default=os.environ.get('SURVEY_BASE_URL', ''), help='Confirmed HTTPS platform address')
    parser.add_argument('--audit-dir', required=True, type=Path, help='Local private output directory, outside Git')
    parser.add_argument('--birthdays', type=Path, help='JSON mapping absolute workbook path to confirmed yyyy-MM-dd')
    parser.add_argument('--commit', action='store_true', help='Actually transmit; same athlete/date may be overwritten')
    parser.add_argument('--retry-failed', action='store_true', help='Retry definitive business failures only; never uncertain requests')
    args = parser.parse_args(argv)
    try:
        endpoint = endpoint_url(args.base_url)
        files = [args.input] if args.input.is_file() else list(args.input.rglob('*.xlsx'))
        files = [p for p in files if not p.name.startswith('~$')]
        birthdays = json.loads(args.birthdays.read_text(encoding='utf-8-sig')) if args.birthdays else {}
        if not isinstance(birthdays, dict):
            raise ValueError('出生日期映射必须为对象')
        accepted, rejected = [], []
        for path in sorted(files):
            try:
                read_report(path, birthdays.get(str(path.resolve())))
                accepted.append(path)
            except (ValueError, OSError) as exc:
                rejected.append({'source': str(path.resolve()), 'error': str(exc)})
        write_json(args.audit_dir / 'rejected.json', rejected)
        accepted_birthdays = {k: v for k, v in birthdays.items() if k in {str(p.resolve()) for p in accepted}}
        if set(birthdays) - {str(p.resolve()) for p in files}:
            raise ValueError('出生日期映射包含不属于本批次的文件')
        items = prepare(accepted, accepted_birthdays)
        if not args.commit:
            summary = {'mode': 'dry-run', 'total': len(items), 'rejected': rejected, 'endpoint': endpoint,
                       'reports': [{k: v for k, v in item.items() if k != 'payload'} | {
                           'name': item['payload']['name'], 'analysisDate': item['payload']['analysisDate'],
                           'rows': {k: len(v) for k, v in item['payload'].items() if isinstance(v, list)}} for item in items]}
            write_json(args.audit_dir / 'preflight.json', summary)
            print(json.dumps({'mode': 'dry-run', 'total': len(items), 'rejected': len(rejected), 'endpoint': endpoint}, ensure_ascii=False))
            return 2 if rejected else 0
        result = execute(items, endpoint, os.environ.get('SURVEY_OPEN_TOKEN', ''), args.audit_dir, retry_failed=args.retry_failed)
        result['rejected'] = len(rejected)
        write_json(args.audit_dir / 'summary.json', result)
        print(json.dumps(result, ensure_ascii=False))
        return 0 if result['success'] == result['total'] and not rejected else 2
    except (ValueError, OSError, KeyError) as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
