from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def table_bounds(worksheet) -> tuple[int, int, list[int]]:
    header_row = max(
        range(1, min(30, worksheet.max_row) + 1),
        key=lambda row: sum(bool(clean(worksheet.cell(row, column).value)) for column in range(1, worksheet.max_column + 1)),
    )
    max_column = max(
        column for column in range(1, worksheet.max_column + 1) if clean(worksheet.cell(header_row, column).value)
    )
    name_column = next(
        column for column in range(1, max_column + 1) if clean(worksheet.cell(header_row, column).value) == "姓名"
    )
    data_rows: list[int] = []
    blank_run = 0
    for row in range(header_row + 1, min(worksheet.max_row, header_row + 10000) + 1):
        if clean(worksheet.cell(row, name_column).value):
            data_rows.append(row)
            blank_run = 0
        elif data_rows:
            blank_run += 1
            if blank_run >= 20:
                break
    return header_row, max_column, data_rows


def main() -> None:
    parser = argparse.ArgumentParser(description="Compare populated cells between two medical summary workbooks.")
    parser.add_argument("baseline", type=Path)
    parser.add_argument("candidate", type=Path)
    parser.add_argument("--json-out", type=Path)
    args = parser.parse_args()

    baseline_book = load_workbook(args.baseline, read_only=True, data_only=True)
    candidate_book = load_workbook(args.candidate, read_only=True, data_only=True)
    baseline_sheet = baseline_book.active
    candidate_sheet = candidate_book.active
    baseline_header, baseline_max_column, baseline_rows = table_bounds(baseline_sheet)
    candidate_header, candidate_max_column, candidate_rows = table_bounds(candidate_sheet)
    if baseline_header != candidate_header:
        raise RuntimeError(f"Header-row mismatch: {baseline_header} != {candidate_header}")
    max_column = max(baseline_max_column, candidate_max_column)
    rows = sorted(set(baseline_rows).union(candidate_rows))
    differences = []
    compared = 0
    for row in rows:
        for column in range(1, max_column + 1):
            baseline_value = clean(baseline_sheet.cell(row, column).value)
            candidate_value = clean(candidate_sheet.cell(row, column).value)
            if not baseline_value and not candidate_value:
                continue
            compared += 1
            if baseline_value != candidate_value:
                differences.append(
                    {
                        "cell": baseline_sheet.cell(row, column).coordinate,
                        "baseline": baseline_value,
                        "candidate": candidate_value,
                    }
                )
    baseline_book.close()
    candidate_book.close()

    result = {
        "baseline": str(args.baseline.resolve()),
        "candidate": str(args.candidate.resolve()),
        "compared_nonempty_cells": compared,
        "differences": len(differences),
        "items": differences,
    }
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(payload, encoding="utf-8")
    print(payload)


if __name__ == "__main__":
    main()
