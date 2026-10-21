from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from pathlib import Path
import math
from copy import copy
from ..energy_reference import energy_reference_text
from .presentation import evaluation_rows, evaluation_warnings, evaluation_context, report_alerts


class NutritionExcelBuilder:

    def build(self, data, output_path):
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()

        self._create_user_sheet(wb, data)
        self._create_food_details_sheet(wb, data)
        self._create_category_sheet(wb, data)
        self._create_nutrition_sheet(wb, data)
        self._create_suggestions_sheet(wb, data)
        self._create_conversion_notes_sheet(wb, data)
        athlete = data.get("athlete_evaluation")
        if evaluation_rows(data):
            ws = wb.create_sheet("运动员专项评价")
            ws.append(["运动员专项膳食评价"])
            ws.append([evaluation_context(data)])
            ws.merge_cells('A2:G2')
            for note in evaluation_warnings(data):
                ws.append([note])
            ws.append(["项目", "实际摄入", "单位", "参考目标", "差值", "最终评价", "口径"])
            for row in evaluation_rows(data):
                ws.append([row[k] for k in ("label", "actual", "unit", "target", "gap", "evaluation")] + [row.get("note", "")])
            for column, width in zip("ABCDEFG", (38, 22, 12, 40, 32, 18, 50)):
                ws.column_dimensions[column].width = width
            for cells in ws:
                for cell in cells:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws['A2'].alignment = Alignment(wrap_text=False, vertical="top")

        # Highlight alert evaluations and the exact snack advice; normal values keep their style.
        for sheet in wb:
            for row in sheet:
                for cell in row:
                    if cell.value in ("不足", "过少", "过多", "不建议吃零食", "能量占比偏低", "能量占比偏高", "占比偏低", "占比偏高"):
                        font = copy(cell.font)
                        font.color = "DC2626"
                        cell.font = font
        ws = wb['建议']
        wb.move_sheet(ws, offset=len(wb.sheetnames)-1-wb.sheetnames.index('建议'))
        ws.append(['报告生成日期', data.get('generated_at') or data.get('generated_date', '')])
        if '运动员专项评价' in wb.sheetnames:
            for row in wb['运动员专项评价']:
                cell = row[3] if len(row) > 3 else None
                if cell is not None and cell.value == '不建议吃零食':
                    font = copy(cell.font)
                    font.color = '000000'
                    cell.font = font
        if '运动员专项评价' in wb.sheetnames:
            for cells in wb['运动员专项评价']:
                for cell in cells:
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = '0.00'
        wb.save(output_path)

    def _create_user_sheet(self, wb, data):
        ws = wb.active
        ws.title = "个人信息"

        header_fill = PatternFill(start_color="4F75B5", end_color="4F75B5", fill_type="solid")
        header_font = Font(name="微软雅黑", size=12, bold=True, color="FFFFFF")
        cell_font = Font(name="微软雅黑", size=11)
        cell_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        headers = ["项目", "内容"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = cell_alignment
            cell.border = thin_border

        user = data.get("user", {})
        date = data.get("date", "")
        source = data.get("food_data_source") or {}

        rows = [
            ("姓名", user.get("name", "")),
            ("年龄", f"{user.get('age', '')} 岁"),
            ("项目种类", user.get("projectType", "")),
            ("训练年限", user.get("trainingYears", "")),
            ("当前状态", user.get("currentStatus", "")),
            ("身高", f"{user.get('height', '')} cm"),
            ("体重", f"{user.get('weight', '')} kg"),
            ("去脂体重（瘦体重）", f"{user.get('fatFreeMass', '')} kg"),
            ("日期", date),
            ("营养数据来源", source.get("label", "未记录")),
            ("营养库查询时间", source.get("queried_at", "")),
            ("云端营养库地址", source.get("base_url", "")),
            ("云端状态", source.get("message", "")),
            ("回退原因", source.get("error", "")),
        ]

        for row_idx, (label, value) in enumerate(rows, 2):
            ws.cell(row=row_idx, column=1, value=label).font = cell_font
            ws.cell(row=row_idx, column=1).alignment = cell_alignment
            ws.cell(row=row_idx, column=1).border = thin_border
            ws.cell(row=row_idx, column=2, value=value).font = cell_font
            ws.cell(row=row_idx, column=2).alignment = cell_alignment
            ws.cell(row=row_idx, column=2).border = thin_border

        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 70

    def _create_food_details_sheet(self, wb, data):
        ws = wb.create_sheet("食物明细")

        header_fill = PatternFill(start_color="4F75B5", end_color="4F75B5", fill_type="solid")
        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        cell_font = Font(name="微软雅黑", size=10)
        cell_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        headers = ["食物名称", "分类", "摄入量(g)", "能量(kcal)", "蛋白质(g)", "脂肪(g)", "碳水(g)", "钙(mg)", "数据来源"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = cell_alignment
            cell.border = thin_border

        foods = data.get("foods", [])
        for row_idx, food in enumerate(foods, 2):
            nutrition = food.get("nutrition", {})
            row_data = [
                food.get("name", ""),
                food.get("category", ""),
                food.get("amount", ""),
                nutrition.get("energy", ""),
                nutrition.get("protein", ""),
                nutrition.get("fat", ""),
                nutrition.get("carbohydrate", ""),
                nutrition.get("calcium", ""),
                food.get("data_source", ""),
            ]
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = cell_font
                cell.alignment = cell_alignment
                cell.border = thin_border

        col_widths = [15, 10, 12, 12, 12, 12, 12, 12, 24]
        for idx, width in enumerate(col_widths, 1):
            ws.column_dimensions[chr(64 + idx)].width = width

    def _create_category_sheet(self, wb, data):
        ws = wb.create_sheet("分类评估")

        header_fill = PatternFill(start_color="4F75B5", end_color="4F75B5", fill_type="solid")
        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        cell_font = Font(name="微软雅黑", size=10)
        cell_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        headers = ["分类", "摄入原重(g)", "评价", "奶类鲜奶等价量(g)", "评价等价食物", "评价等价量(g)"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = cell_alignment
            cell.border = thin_border

        category_result = data.get("category_result", {})
        row_idx = 2
        for cat, info in category_result.items():
            ws.cell(row=row_idx, column=1, value=cat).font = cell_font
            ws.cell(row=row_idx, column=1).alignment = cell_alignment
            ws.cell(row=row_idx, column=1).border = thin_border

            ws.cell(row=row_idx, column=2, value=info.get("value", "")).font = cell_font
            ws.cell(row=row_idx, column=2).alignment = cell_alignment
            ws.cell(row=row_idx, column=2).border = thin_border

            ws.cell(row=row_idx, column=3, value=info.get("evaluation", "")).font = cell_font
            ws.cell(row=row_idx, column=3).alignment = cell_alignment
            ws.cell(row=row_idx, column=3).border = thin_border
            if cat == '奶类':
                ws.cell(row=row_idx, column=4, value=info.get('equivalent_value'))
            equivalent = info.get('reference_equivalent')
            if equivalent:
                ws.cell(row=row_idx, column=5, value=equivalent['representative'])
                ws.cell(row=row_idx, column=6, value=equivalent['value'])
            row_idx += 1

        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 15
        ws.column_dimensions["D"].width = 36
        ws.column_dimensions['E'].width = 20
        ws.column_dimensions['F'].width = 20
        ws['D1'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 34

    def _create_nutrition_sheet(self, wb, data):
        ws = wb.create_sheet("营养总计")

        header_fill = PatternFill(start_color="4F75B5", end_color="4F75B5", fill_type="solid")
        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        cell_font = Font(name="微软雅黑", size=10)
        cell_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        headers = ["项目", "摄入量", "单位", "评价"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = cell_alignment
            cell.border = thin_border

        nutrition = data.get("nutrition", {})
        athlete_rows = {row["key"]: row for row in (data.get("athlete_evaluation") or {}).get("rows", [])}
        rows = [
            ("总能量", nutrition.get("energy", {}).get("value", ""), "kcal/day", nutrition.get("energy", {}).get("evaluation", "")),
            ("蛋白质", nutrition.get("protein", ""), "g", athlete_rows.get("protein", {}).get("evaluation", "-")),
            ("脂肪", nutrition.get("fat", ""), "g", athlete_rows.get("fat", {}).get("evaluation", "-")),
            ("碳水化合物", nutrition.get("carbohydrate", ""), "g", athlete_rows.get("carbohydrate", {}).get("evaluation", "-")),
            ("钙", nutrition.get("calcium", {}).get("value", ""), "mg/day", nutrition.get("calcium", {}).get("evaluation", "")),
            ("推荐能量", energy_reference_text(nutrition.get("energy", {})) if nutrition.get("energy", {}).get("reference_type") in {'bmr', 'bmr_plus_net'} else nutrition.get("energy", {}).get("standard"), "kcal/天", '静息代谢＋日均运动净消耗' if nutrition.get("energy", {}).get("reference_type") == 'bmr_plus_net' else nutrition.get("energy", {}).get("standard_note", "")),
            ("活动量", nutrition.get("energy", {}).get("activity_label") or "未选择", "", ""),
        ]

        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.font = cell_font
                cell.alignment = cell_alignment
                cell.border = thin_border

        note_row = len(rows) + 2
        note = data.get("nutrition_disclaimer") or "以上营养成分仅根据记录饮食计算，不包含其他补充剂摄入。"
        ws.cell(row=note_row, column=1, value="说明").font = cell_font
        ws.cell(row=note_row, column=1).alignment = cell_alignment
        ws.cell(row=note_row, column=1).border = thin_border
        ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=4)
        note_cell = ws.cell(row=note_row, column=2, value=note)
        note_cell.font = cell_font
        note_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        note_cell.border = thin_border
        for column in range(3, 5):
            ws.cell(row=note_row, column=column).border = thin_border
        ws.row_dimensions[note_row].height = 30

        ws.column_dimensions["A"].width = 15
        ws.column_dimensions["B"].width = 42
        ws.column_dimensions["C"].width = 12
        ws.column_dimensions["D"].width = 24

    def _create_suggestions_sheet(self, wb, data):
        ws = wb.create_sheet("建议")

        header_fill = PatternFill(start_color="4F75B5", end_color="4F75B5", fill_type="solid")
        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        cell_font = Font(name="微软雅黑", size=10)
        cell_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin")
        )

        headers = ["序号", "建议内容"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = cell_alignment
            cell.border = thin_border

        suggestions = data.get("suggestions", [])
        for row_idx, suggestion in enumerate(suggestions, 2):
            ws.cell(row=row_idx, column=1, value=row_idx - 1).font = cell_font
            ws.cell(row=row_idx, column=1).alignment = cell_alignment
            ws.cell(row=row_idx, column=1).border = thin_border

            ws.cell(row=row_idx, column=2, value=suggestion).font = cell_font
            ws.cell(row=row_idx, column=2).alignment = cell_alignment
            ws.cell(row=row_idx, column=2).border = thin_border
            # 中文按双宽估算行数；避免Excel导出建议被固定行高截断。
            display_width = sum(2 if ord(c) > 127 else 1 for c in str(suggestion))
            ws.row_dimensions[row_idx].height = max(30, 16 * max(len(str(suggestion).splitlines()), math.ceil(display_width / 96)) + 8)

        if data.get('joint_plan_summary'):
            summary_row = len(suggestions)+3
            ws.cell(summary_row,2,data['joint_plan_summary']).font=cell_font
            ws.cell(summary_row,2).alignment=cell_alignment
            ws.row_dimensions[summary_row].height=72
        for message in report_alerts(data):
            row = ws.max_row + 1
            ws.cell(row, 1, '核查').font = cell_font
            ws.cell(row, 2, message).font = cell_font
            ws.cell(row, 2).alignment = cell_alignment
            ws.row_dimensions[row].height = max(30, 16 * math.ceil(sum(2 if ord(c)>127 else 1 for c in message)/96)+8)
        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 100

    def _create_conversion_notes_sheet(self, wb, data):
        """保留 OCR 频率换算和代表食物替代说明，避免估算依据在导出时丢失。"""
        ws = wb.create_sheet("OCR换算说明")
        header_fill = PatternFill(start_color="4F75B5", end_color="4F75B5", fill_type="solid")
        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        body_font = Font(name="微软雅黑", size=10)
        wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)

        headers = ["级别", "代码", "食物", "说明"]
        for column, value in enumerate(headers, 1):
            cell = ws.cell(row=1, column=column, value=value)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = wrap

        warnings = data.get("conversion_warnings") or []
        if not warnings:
            warnings = [{"severity": "info", "code": "none", "food": "", "message": "无 OCR 换算告警"}]
        for row_index, warning in enumerate(warnings, 2):
            if not isinstance(warning, dict):
                warning = {"severity": "warning", "code": "unknown", "food": "", "message": str(warning)}
            values = [
                warning.get("severity", "warning"),
                warning.get("code", ""),
                warning.get("food", ""),
                warning.get("message", ""),
            ]
            for column, value in enumerate(values, 1):
                cell = ws.cell(row=row_index, column=column, value=value)
                cell.font = body_font
                cell.alignment = wrap

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 24
        ws.column_dimensions["C"].width = 18
        ws.column_dimensions["D"].width = 90
