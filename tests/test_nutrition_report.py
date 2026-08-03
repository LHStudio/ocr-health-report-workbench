from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import fitz
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from nutrition_report import FOOD_REFERENCE_MAP, NutritionReportService, build_payload_from_ocr_person  # noqa: E402
from local_service import body_composition_match_candidates, extract_fat_free_mass, merge_body_composition_records  # noqa: E402


def sample_person() -> dict:
    return {
        "id": "ocr/测试人员",
        "general": {"姓名": "测试甲", "调查日期": "2026-07-21"},
        "food_rows": [
            {"食物名称": "米饭", "平均每次食用量": "100g", "次数": "2", "频率周期(请核对)": "每天", "是否不吃": ""},
            {"食物名称": "牛肉", "平均每次食用量": "70克", "次数": "3", "频率周期(请核对)": "每周", "是否不吃": ""},
            {"食物名称": "鲜奶", "平均每次食用量": "250ml", "次数": "1", "频率周期(请核对)": "每天", "是否不吃": ""},
            {"食物名称": "糕点", "平均每次食用量": "50-70g", "次数": "2", "频率周期(请核对)": "每月", "是否不吃": ""},
            {"食物名称": "奶粉", "平均每次食用量": "2勺", "次数": "1", "频率周期(请核对)": "每天", "是否不吃": ""},
            {"食物名称": "茶/茶饮料", "平均每次食用量": "300ml", "次数": "0", "频率周期(请核对)": "不吃", "是否不吃": "是"},
        ],
    }


class OCRNutritionAdapterTest(unittest.TestCase):
    def test_all_38_fixed_foods_have_a_real_database_reference(self) -> None:
        service = NutritionReportService(prefer_cloud=False)
        self.assertEqual(38, len(FOOD_REFERENCE_MAP))
        missing = {
            reference
            for reference, _ in FOOD_REFERENCE_MAP.values()
            if reference not in service.nutrition_service.food_db
        }
        self.assertEqual(set(), missing)

    def test_frequency_rows_are_converted_to_daily_grams_with_traceability(self) -> None:
        adapted = build_payload_from_ocr_person(
            sample_person(),
            user_overrides={"age": 22, "projectType": "女足", "trainingYears": "5年", "currentStatus": "正常训练", "gender": "female", "activityLevel": "2", "weight": 55, "fatFreeMass": "46.6"},
        )

        meals = {item["name"]: item["amount"] for item in adapted.payload["meals"]}
        self.assertEqual(200.0, meals["米饭"])
        self.assertEqual(30.0, meals["牛肉"])
        self.assertEqual(250.0, meals["鲜奶"])
        self.assertEqual(4.0, meals["蛋糕"])
        self.assertEqual(4, adapted.usable_food_count)
        self.assertEqual("测试甲", adapted.payload["user"]["name"])
        self.assertEqual("female", adapted.payload["user"]["gender"])
        self.assertEqual("女足", adapted.payload["user"]["projectType"])
        self.assertEqual("5年", adapted.payload["user"]["trainingYears"])
        self.assertEqual("正常训练", adapted.payload["user"]["currentStatus"])
        self.assertEqual("46.6", adapted.payload["user"]["fatFreeMass"])
        self.assertEqual("2026-07-21", adapted.payload["date"])

        codes = {warning["code"] for warning in adapted.warnings}
        self.assertIn("volume_density_assumed", codes)
        self.assertIn("range_quantity", codes)
        self.assertIn("proxy_food_reference", codes)
        self.assertIn("unsupported_quantity_unit", codes)
        not_eaten = [item for item in adapted.source_foods if item.get("status") == "not_eaten"]
        self.assertEqual(["茶/茶饮料"], [item["food"] for item in not_eaten])

    def test_unmapped_and_unresolved_rows_are_not_silently_counted(self) -> None:
        person = {
            "id": "ocr/2",
            "general": {},
            "food_rows": [
                {"食物名称": "陌生食物", "平均每次食用量": "100g", "次数": "1", "频率周期(请核对)": "每天"},
                {"食物名称": "米饭", "平均每次食用量": "100g", "次数": "1", "频率周期(请核对)": "未识别"},
            ],
        }
        adapted = build_payload_from_ocr_person(person)

        self.assertEqual([], adapted.payload["meals"])
        codes = {warning["code"] for warning in adapted.warnings}
        self.assertTrue({"unmapped_food", "unresolved_period", "no_usable_foods"}.issubset(codes))


class NutritionReportServiceTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.service = NutritionReportService(prefer_cloud=False)

    def test_original_user_meals_json_remains_supported(self) -> None:
        data = self.service.build_report_data({
            "user": {"name": "原格式"},
            "date": "2026-07-21",
            "meals": [{"name": "米饭", "amount": 100}],
            "autoEvaluation": True,
        })

        self.assertEqual(116.0, data["nutrition"]["energy"]["value"])
        self.assertEqual(2.6, data["nutrition"]["protein"])
        self.assertEqual([], data["unmatched_foods"])

    def test_single_person_pdf_and_excel_are_generated(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            pdf_path = temp / "测试甲_营养报告.pdf"
            excel_path = temp / "测试甲_营养报告.xlsx"
            overrides = {"age": 22, "projectType": "女足", "trainingYears": "5年", "currentStatus": "正常训练", "gender": "female", "activityLevel": "2", "weight": 55, "fatFreeMass": "46.6"}

            pdf_result = self.service.generate_pdf_from_ocr_person(
                sample_person(), pdf_path, user_overrides=overrides
            )
            excel_result = self.service.generate_excel_from_ocr_person(
                sample_person(), excel_path, user_overrides=overrides
            )

            self.assertTrue(pdf_path.is_file())
            self.assertEqual(b"%PDF", pdf_path.read_bytes()[:4])
            self.assertGreater(pdf_path.stat().st_size, 10_000)
            self.assertFalse(list(temp.glob("nutrition_pie_*.png")))
            with fitz.open(pdf_path) as pdf_document:
                page_texts = [page.get_text() for page in pdf_document]
                pdf_text = "\n".join(page_texts)
                image_counts = [len(page.get_images(full=True)) for page in pdf_document]
            self.assertEqual(3, len(page_texts))
            self.assertIn("姓名", page_texts[0])
            self.assertIn("测试甲", page_texts[0])
            self.assertIn("年龄", page_texts[0])
            self.assertIn("22 岁", page_texts[0])
            self.assertIn("项目种类", page_texts[0])
            self.assertIn("女足", page_texts[0])
            self.assertIn("训练年限", page_texts[0])
            self.assertIn("5年", page_texts[0])
            self.assertIn("当前状态", page_texts[0])
            self.assertIn("正常训练", page_texts[0])
            self.assertIn("去脂体重", page_texts[0])
            self.assertIn("46.6 kg", page_texts[0])
            self.assertTrue(all(category in page_texts[0] for category in ("谷类", "薯类", "蔬菜", "水果", "畜禽肉", "水产品", "蛋类")))
            self.assertTrue(all(category in page_texts[1] for category in ("豆类", "坚果", "奶类", "零食", "运动营养", "饮料", "酒类")))
            self.assertIn("能量及钙摄入量评估报告", page_texts[2])
            self.assertNotIn("食物图片", pdf_text)
            self.assertFalse(any(image_counts))
            self.assertNotIn("营养数据来源", page_texts[0])
            self.assertIn("本地食物营养库", pdf_text)
            self.assertIn("查询时间", pdf_text)
            self.assertIn("营养数据来源", page_texts[-1])
            self.assertIn("OCR 换算与数据质量说明", page_texts[-1])
            self.assertEqual(4, pdf_result["usable_food_count"])
            self.assertEqual(2100, pdf_result["report_data"]["nutrition"]["energy"]["standard"])

            workbook = load_workbook(excel_path, read_only=True, data_only=True)
            try:
                self.assertEqual(
                    ["个人信息", "食物明细", "分类评估", "营养总计", "建议", "OCR换算说明"],
                    workbook.sheetnames,
                )
                note_codes = {
                    workbook["OCR换算说明"].cell(row, 2).value
                    for row in range(2, workbook["OCR换算说明"].max_row + 1)
                }
                self.assertIn("proxy_food_reference", note_codes)
                personal = workbook["个人信息"]
                personal_values = {
                    personal.cell(row, 1).value: personal.cell(row, 2).value
                    for row in range(2, personal.max_row + 1)
                }
                self.assertEqual("本地食物营养库", personal_values["营养数据来源"])
                self.assertEqual("女足", personal_values["项目种类"])
                self.assertEqual("5年", personal_values["训练年限"])
                self.assertEqual("正常训练", personal_values["当前状态"])
                self.assertEqual("46.6 kg", personal_values["去脂体重（瘦体重）"])
                self.assertIn("api/trophic", personal_values["云端营养库地址"])
            finally:
                workbook.close()
            self.assertEqual(4, excel_result["usable_food_count"])

    def test_fat_free_mass_is_extracted_and_merged_only_for_a_unique_name(self) -> None:
        transcript = "InBody770\nID 260604-5（刘紫玉）\n人体成分分析\n去脂体重 (kg) 46.6\n体重 62.4"
        self.assertEqual("46.6", extract_fat_free_mass(transcript))

        people = [
            {"id": "A", "general": {"姓名": "刘紫玉"}},
            {"id": "B", "general": {"姓名": "陈晨"}},
        ]
        matched, unmatched = merge_body_composition_records(people, [
            {"name": "刘紫玉", "fat_free_mass": "46.6", "filename": "inbody.pdf", "file_url": "/local-files/demo/inbody.md"},
            {"name": "未知", "fat_free_mass": "40.1", "filename": "unknown.pdf", "file_url": ""},
        ])
        self.assertEqual(1, matched)
        self.assertEqual(["unknown.pdf"], unmatched)
        self.assertEqual("46.6", people[0]["general"]["去脂体重"])
        self.assertNotIn("去脂体重", people[1]["general"])

    def test_body_composition_name_pairing_suggests_similar_names_without_auto_match(self) -> None:
        people = [
            {"id": "A", "general": {"姓名": "刘紫玉"}},
            {"id": "B", "general": {"姓名": "刘子玉"}},
            {"id": "C", "general": {"姓名": "陈晨"}},
        ]
        exact = body_composition_match_candidates("刘紫玉", people)
        similar = body_composition_match_candidates("刘紫雨", people)
        self.assertEqual("A", exact[0]["person_id"])
        self.assertEqual(100, exact[0]["score"])
        self.assertTrue(any(item["person_id"] == "A" for item in similar))
        self.assertTrue(all(item["score"] < 100 for item in similar))

    def test_manual_review_replaces_machine_evaluation_and_exports_disclaimer(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp = Path(temp_dir)
            pdf_path = temp / "人工复核.pdf"
            excel_path = temp / "人工复核.xlsx"
            manual = {
                "categoryEvaluations": {"谷类": "偏多", "奶类": "偏少"},
                "energyEvaluation": "适宜",
                "calciumEvaluation": "偏少",
                "overallSuggestions": "人工建议一\n人工建议二",
                "interpretation": "该结论已由营养师人工复核。",
            }
            kwargs = {
                "user_overrides": {"age": 22, "gender": "female", "activityLevel": "2", "weight": 55},
                "auto_evaluation": False,
                "manual_evaluations": manual,
            }
            pdf_result = self.service.generate_pdf_from_ocr_person(sample_person(), pdf_path, **kwargs)
            self.service.generate_excel_from_ocr_person(sample_person(), excel_path, **kwargs)

            report = pdf_result["report_data"]
            self.assertEqual("偏多", report["category_result"]["谷类"]["evaluation"])
            self.assertEqual("偏少", report["nutrition"]["calcium"]["evaluation"])
            self.assertEqual(["人工建议一", "人工建议二"], report["suggestions"])
            self.assertEqual("该结论已由营养师人工复核。", report["interpretation"])
            self.assertIn("不包含维生素", report["nutrition_disclaimer"])

            with fitz.open(pdf_path) as document:
                text = "\n".join(page.get_text() for page in document)
            self.assertIn("该结论已由营养师人工复核。", text)
            self.assertIn("不包含维生素", text)

            workbook = load_workbook(excel_path, read_only=True, data_only=True)
            try:
                nutrition_sheet = workbook["营养总计"]
                note_row = next(row for row in range(2, nutrition_sheet.max_row + 1) if nutrition_sheet.cell(row, 1).value == "说明")
                self.assertIn("不包含维生素", nutrition_sheet.cell(note_row, 2).value)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
