import sys
import tempfile
import unittest
import uuid
from copy import copy
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import local_service as service  # noqa: E402


class MedicalNormalRangeTest(unittest.TestCase):
    def test_default_ranges_flag_only_out_of_range_values(self) -> None:
        ranges = {key: dict(value) for key, value in service.DEFAULT_MEDICAL_NORMAL_RANGES.items()}
        low = service.medical_range_assessment("白细胞\nWBC（×10^9/L）", "3.4", ranges)
        normal = service.medical_range_assessment("白细胞\nWBC（×10^9/L）", "3.5", ranges)
        high = service.medical_range_assessment("嗜碱性粒细胞百分数\nBASO%（%）", "1.1", ranges)
        self.assertEqual((True, "低于正常范围", "3.5–9.5"), (low["is_abnormal"], low["abnormal_reason"], low["normal_range"]))
        self.assertFalse(normal["is_abnormal"])
        self.assertEqual((True, "高于正常范围"), (high["is_abnormal"], high["abnormal_reason"]))

    def test_comparison_results_are_marked_only_when_definitely_abnormal(self) -> None:
        ranges = {"25羟维生素d": {"min": 20, "max": None}, "生长激素": {"min": None, "max": 8}}
        self.assertTrue(service.medical_range_assessment("25-羟维生素D", "<20", ranges)["is_abnormal"])
        self.assertFalse(service.medical_range_assessment("25-羟维生素D", "<21", ranges)["is_abnormal"])
        self.assertTrue(service.medical_range_assessment("生长激素", ">8", ranges)["is_abnormal"])
        self.assertFalse(service.medical_range_assessment("生长激素", ">7", ranges)["is_abnormal"])

    def test_saved_settings_can_intentionally_disable_a_default_range(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            config = Path(directory) / "medical_normal_ranges.json"
            with patch.object(service, "MEDICAL_NORMAL_RANGE_CONFIG", config):
                service.save_medical_normal_ranges({"白细胞": {"min": 4.0, "max": 9.0}})
                self.assertEqual({"白细胞": {"min": 4.0, "max": 9.0}}, service.load_medical_normal_ranges())

    def test_abnormal_excel_cell_is_red_and_can_be_restored(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "result.xlsx"
            workbook = Workbook()
            cell = workbook.active.cell(2, 2, "3.4")
            original = SimpleNamespace(font=copy(cell.font), fill=copy(cell.fill))
            service.apply_medical_range_style(cell, True)
            workbook.save(path)
            workbook.close()
            workbook = load_workbook(path)
            cell = workbook.active.cell(2, 2)
            self.assertEqual("FFFF0000", cell.font.color.rgb)
            service.apply_medical_range_style(cell, False, original)
            self.assertFalse(cell.font.color and cell.font.color.type == "rgb" and cell.font.color.rgb == "FFFF0000")
            workbook.close()

    def test_save_review_rechecks_and_marks_the_exported_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            template_path, output_path = root / "template.xlsx", root / "output.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.cell(1, 1, "序号")
            worksheet.cell(1, 2, "白细胞\nWBC（×10^9/L）")
            worksheet.cell(2, 1, 1)
            workbook.save(template_path)
            workbook.save(output_path)
            workbook.close()
            job_id = uuid.uuid4().hex
            service.JOBS[job_id] = {
                "status": "completed", "job_dir": root, "output_relative": Path("output.xlsx"),
                "template_path": template_path, "header_row": 1, "output_url": "/local-files/test/output.xlsx",
            }
            try:
                person = {"row": 2, "review_fields": [{"header": "白细胞\nWBC（×10^9/L）", "source": "白细胞", "value": "3.4"}]}
                with patch.object(service, "MEDICAL_NORMAL_RANGE_CONFIG", root / "normal_ranges.json"):
                    result = service.save_review(job_id, {"people": [person]})
                self.assertTrue(result["people"][0]["review_fields"][0]["is_abnormal"])
                workbook = load_workbook(output_path)
                self.assertEqual("FFFF0000", workbook.active.cell(2, 2).font.color.rgb)
                workbook.close()
            finally:
                service.JOBS.pop(job_id, None)


if __name__ == "__main__":
    unittest.main()
