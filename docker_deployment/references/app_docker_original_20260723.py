# -*- coding: utf-8 -*-
# @Time: 2026/3/31 16:07
# @Author: LHStudio
# @File: app.py
# @Software: PyCharm

from fastapi import FastAPI, UploadFile, File, HTTPException
from pathlib import Path
import shutil
import uuid
import json
import base64
import re
import uvicorn
from typing import List, Tuple, Optional

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from paddleocr import PaddleOCRVL

from fastapi.middleware.cors import CORSMiddleware

app = FastAPI(title="OCR Parse Service")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://127.0.0.1:3000",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# =========================================================
# OCR 初始化
# =========================================================
pipeline = PaddleOCRVL(
    vl_rec_backend="vllm-server",
    vl_rec_server_url="http://172.18.0.2:8118/v1"
)

# =========================================================
# 路径配置
# =========================================================
BASE_DIR = Path("/app/pdf_to_excel_service")
UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "outputs"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED_EXTS = {".pdf", ".png", ".jpg", ".jpeg", ".bmp", ".webp"}


# =========================================================
# 基础工具
# =========================================================
def safe_read_text(file_path: Path) -> str:
    try:
        return file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return file_path.read_text(encoding="utf-8", errors="ignore")


def normalize_text(text: str) -> str:
    if text is None:
        return ""
    return text.replace("\r\n", "\n").replace("\r", "\n")


def extract_element_info(element) -> str:
    """
    递归提取元素内的所有文本信息（忽略图片标签）
    """
    if isinstance(element, str):
        return element.strip()

    if getattr(element, "name", None) == "img":
        return ""

    result = []
    if hasattr(element, "children"):
        for child in element.children:
            text = extract_element_info(child)
            if text:
                result.append(text)

    return " ".join(result).strip()


# =========================================================
# HTML 表格解析
# =========================================================
def parse_html_table_to_excel(table_soup, ws, start_row: int) -> int:
    """
    解析 HTML 表格，保留合并网格和对齐
    """
    row_cursor = start_row
    occupied_cells = {}

    rows = table_soup.find_all("tr")
    if not rows:
        return row_cursor

    for tr in rows:
        col_cursor = 1
        cells = tr.find_all(["td", "th"], recursive=False)
        if not cells:
            cells = tr.find_all(["td", "th"])

        for cell in cells:
            while occupied_cells.get((row_cursor, col_cursor), False):
                col_cursor += 1

            rowspan = int(cell.get("rowspan", 1) or 1)
            colspan = int(cell.get("colspan", 1) or 1)

            raw_text = extract_element_info(cell)

            clean_text = clean_ocr_latex_text(raw_text)
            clean_text = re.sub(r"\*\*(.*?)\*\*", r"\1", clean_text).strip()

            excel_cell = ws.cell(row=row_cursor, column=col_cursor, value=clean_text)

            style_attr = (cell.get("style", "") or "").lower()
            if "text-align: center" in style_attr:
                excel_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                excel_cell.alignment = Alignment(vertical="center", wrap_text=True)

            if cell.name == "th":
                excel_cell.font = Font(bold=True)

            if rowspan > 1 or colspan > 1:
                ws.merge_cells(
                    start_row=row_cursor,
                    start_column=col_cursor,
                    end_row=row_cursor + rowspan - 1,
                    end_column=col_cursor + colspan - 1
                )
                for r in range(row_cursor, row_cursor + rowspan):
                    for c in range(col_cursor, col_cursor + colspan):
                        occupied_cells[(r, c)] = True
            else:
                occupied_cells[(row_cursor, col_cursor)] = True

            col_cursor += colspan

        row_cursor += 1

    return row_cursor


# =========================================================
# Markdown 表格解析（兜底）
# =========================================================
def split_md_row(line: str) -> List[str]:
    line = line.strip()
    if not line:
        return []

    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]

    return [part.strip() for part in line.split("|")]


def is_markdown_table_separator(line: str) -> bool:
    cells = split_md_row(line)
    if not cells:
        return False

    for cell in cells:
        test_cell = cell.replace(" ", "")
        if not re.fullmatch(r":?-{3,}:?", test_cell):
            return False
    return True


def is_markdown_table_block(lines: List[str]) -> bool:
    if len(lines) < 2:
        return False

    first = lines[0].strip()
    second = lines[1].strip()

    if "|" not in first or "|" not in second:
        return False

    return is_markdown_table_separator(second)


def parse_markdown_table_to_excel(table_lines: List[str], ws, start_row: int) -> int:
    if len(table_lines) < 2:
        return start_row

    header = split_md_row(table_lines[0])
    data_lines = table_lines[2:] if is_markdown_table_separator(table_lines[1]) else table_lines[1:]

    row_cursor = start_row

    for col_idx, value in enumerate(header, start=1):
        cell = ws.cell(row=row_cursor, column=col_idx, value=clean_ocr_latex_text(value))
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_cursor += 1

    for line in data_lines:
        cols = split_md_row(line)
        if not cols:
            continue
        for col_idx, value in enumerate(cols, start=1):
            cell = ws.cell(row=row_cursor, column=col_idx, value=clean_ocr_latex_text(value))
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        row_cursor += 1

    return row_cursor


# =========================================================
# 普通文本写入
# =========================================================
def write_text_line(ws, row_idx: int, text: str) -> int:
    clean_line = clean_ocr_latex_text(text).strip()
    if not clean_line:
        return row_idx

    indent_level = 0
    is_bold = False
    font_size = 11

    heading_match = re.match(r"^(#{1,6})\s+(.*)", clean_line)
    if heading_match:
        level = len(heading_match.group(1))
        clean_line = heading_match.group(2).strip()
        is_bold = True
        font_size = max(11, 16 - level)
        indent_level = max(0, level - 1)

    elif re.match(r"^(\d+\.|[a-zA-Z]\.|[-*+])\s+", clean_line):
        indent_level = 1

    if "**" in clean_line:
        clean_line = re.sub(r"\*\*(.*?)\*\*", r"\1", clean_line)
        is_bold = True

    cell = ws.cell(row=row_idx, column=1, value=clean_line)
    cell.font = Font(bold=is_bold, size=font_size)

    if indent_level > 0:
        cell.alignment = Alignment(indent=indent_level, vertical="center", wrap_text=True)
    else:
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    return row_idx + 1


def process_plain_text_block(block_text: str, ws, start_row: int) -> int:
    row_cursor = start_row
    lines = normalize_text(block_text).split("\n")

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        maybe_table = [line]
        j = i + 1
        while j < len(lines) and lines[j].strip():
            maybe_table.append(lines[j].rstrip())
            j += 1

        if is_markdown_table_block(maybe_table):
            row_cursor = parse_markdown_table_to_excel(maybe_table, ws, row_cursor)
            row_cursor += 1
            i = j
            continue

        row_cursor = write_text_line(ws, row_cursor, line)
        i += 1

    return row_cursor


# =========================================================
# 混合 markdown/html 内容处理（兜底）
# =========================================================
def process_markdown_text_to_sheet(md_text: str, ws, start_row: int = 1) -> int:
    md_text = normalize_text(md_text)
    current_row = start_row

    table_pattern = re.compile(r"(<table.*?</table>)", flags=re.IGNORECASE | re.DOTALL)
    parts = table_pattern.split(md_text)

    for part in parts:
        if not part or not part.strip():
            continue

        if re.match(r"^\s*<table", part, flags=re.IGNORECASE):
            soup = BeautifulSoup(part, "html.parser")
            table = soup.find("table")
            if table:
                current_row = parse_html_table_to_excel(table, ws, current_row)
                current_row += 1
            else:
                current_row = process_plain_text_block(part, ws, current_row)
            continue

        current_row = process_plain_text_block(part, ws, current_row)

    return current_row


# =========================================================
# JSON block 解析（核心修复）
# =========================================================
def load_json_objects_from_text(raw_text: str) -> List[dict]:
    """
    支持：
    1. 单个 JSON 对象
    2. 多个 JSON 对象拼接
    3. JSON 数组
    """
    raw_text = raw_text.strip()
    if not raw_text:
        return []

    # 先尝试整体解析
    try:
        data = json.loads(raw_text)
        if isinstance(data, list):
            return [item for item in data if isinstance(item, dict)]
        if isinstance(data, dict):
            return [data]
    except Exception:
        pass

    # 尝试按大括号逐个解析
    decoder = json.JSONDecoder()
    idx = 0
    results = []
    while idx < len(raw_text):
        while idx < len(raw_text) and raw_text[idx].isspace():
            idx += 1
        if idx >= len(raw_text):
            break
        try:
            obj, end = decoder.raw_decode(raw_text, idx)
            if isinstance(obj, dict):
                results.append(obj)
            elif isinstance(obj, list):
                results.extend([item for item in obj if isinstance(item, dict)])
            idx = end
        except Exception:
            break

    return results


def extract_blocks_from_json_text(json_text: str) -> List[dict]:
    """
    从 OCR json 文本中提取 block，按 block_order 排序
    兼容结构：
    {
        "parsing_res_list": [...]
    }
    关键修复：
    - 不再对所有 blocks 做全局排序；
    - 每个 JSON 对象/页面内部按 block_order 排序；
    - 然后按 JSON 原始出现顺序追加，避免多页 block_order 重复导致错序。
    或其他嵌套结构
    """
    objs = load_json_objects_from_text(json_text)
    ordered_blocks = []

    def sort_key(x):
        order = x.get("block_order")
        block_id = x.get("block_id", 10 ** 9)

        if order is None:
            order = 10 ** 9

        return order, block_id

    for obj_index, obj in enumerate(objs):
        if not isinstance(obj, dict):
            continue

        page_blocks = []

        # 标准结构
        if isinstance(obj.get("parsing_res_list"), list):
            for block in obj["parsing_res_list"]:
                if isinstance(block, dict):
                    page_blocks.append(block)

        # 兼容其他可能字段
        for key in ["blocks", "results", "data"]:
            value = obj.get(key)
            if isinstance(value, list):
                for item in value:
                    if isinstance(item, dict) and (
                            "block_label" in item or "block_content" in item
                    ):
                        page_blocks.append(item)

        # 只在当前 JSON 对象内部排序
        page_blocks.sort(key=sort_key)

        # 按 JSON 对象原始顺序追加
        ordered_blocks.extend(page_blocks)

    return ordered_blocks


def process_blocks_to_sheet(blocks: List[dict], ws, start_row: int = 1) -> int:
    """
    优先按 OCR block 顺序重建 Excel
    """
    current_row = start_row

    for block in blocks:
        block_label = str(block.get("block_label", "")).strip().lower()
        block_content = block.get("block_content", "")

        if not block_content:
            continue

        # 表格 block：最关键，优先处理
        if block_label == "table":
            content = str(block_content).strip()

            # HTML table
            if re.search(r"<table.*?</table>", content, flags=re.IGNORECASE | re.DOTALL):
                soup = BeautifulSoup(content, "html.parser")
                table = soup.find("table")
                if table:
                    current_row = parse_html_table_to_excel(table, ws, current_row)
                    current_row += 1
                    continue

            # markdown table
            lines = [line for line in normalize_text(content).split("\n") if line.strip()]
            if is_markdown_table_block(lines):
                current_row = parse_markdown_table_to_excel(lines, ws, current_row)
                current_row += 1
                continue

            # 表格块但解析失败，按文本兜底
            current_row = process_plain_text_block(content, ws, current_row)
            current_row += 1
            continue

        # 文本块
        if block_label in {"text", "title", "paragraph", "list"}:
            current_row = process_plain_text_block(str(block_content), ws, current_row)
            continue

        # 其他未知 block，兜底成文本
        current_row = process_plain_text_block(str(block_content), ws, current_row)

    return current_row


# =========================================================
# 样式
# =========================================================
def set_sheet_style(ws):
    max_col = max(ws.max_column, 7)

    for col_idx in range(1, max_col + 1):
        col_letter = get_column_letter(col_idx)
        if col_idx == 1:
            ws.column_dimensions[col_letter].width = 60
        else:
            ws.column_dimensions[col_letter].width = 18

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                current_alignment = cell.alignment
                if current_alignment:
                    cell.alignment = Alignment(
                        horizontal=current_alignment.horizontal,
                        vertical=current_alignment.vertical or "center",
                        wrap_text=True,
                        indent=current_alignment.indent or 0
                    )
                else:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)


# =========================================================
# 收集 OCR 输出
# =========================================================
def collect_saved_outputs(task_output_dir: Path) -> Tuple[str, str, List[str], List[str]]:
    md_files = sorted(task_output_dir.rglob("*.md"))
    json_files = sorted(task_output_dir.rglob("*.json"))

    markdown_parts = []
    json_parts = []

    for md_file in md_files:
        content = safe_read_text(md_file)
        if content.strip():
            markdown_parts.append(content)

    for json_file in json_files:
        content = safe_read_text(json_file)
        if content.strip():
            json_parts.append(content)

    markdown_text = "\n\n".join(markdown_parts).strip()
    json_text = "\n\n".join(json_parts).strip()

    return markdown_text, json_text, [str(p) for p in md_files], [str(p) for p in json_files]


# =========================================================
# Excel 构建
# =========================================================
def build_excel_from_ocr_outputs(markdown_text: str, json_text: str, excel_path: Path) -> dict:
    """
    根治方案：
    1. 优先 JSON block
    2. 不行再回退 markdown
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "问卷数据"

    used_mode = "markdown_fallback"
    block_count = 0

    blocks = extract_blocks_from_json_text(json_text)
    if blocks:
        process_blocks_to_sheet(blocks, ws, start_row=1)
        used_mode = "json_blocks"
        block_count = len(blocks)
    else:
        process_markdown_text_to_sheet(markdown_text, ws, start_row=1)

    set_sheet_style(ws)
    wb.save(excel_path)

    return {
        "used_mode": used_mode,
        "block_count": block_count
    }


def count_tables(markdown_text: str, json_text: str) -> int:
    """
    优先从 JSON block 统计 table 数量
    """
    blocks = extract_blocks_from_json_text(json_text)
    json_table_count = sum(1 for block in blocks if str(block.get("block_label", "")).lower() == "table")
    if json_table_count > 0:
        return json_table_count

    markdown_text = normalize_text(markdown_text)

    html_count = len(
        re.findall(r"<table.*?</table>", markdown_text, flags=re.DOTALL | re.IGNORECASE)
    )

    md_count = 0
    blocks_text = re.split(r"\n\s*\n", markdown_text)
    for block in blocks_text:
        lines = [line for line in block.split("\n") if line.strip()]
        if is_markdown_table_block(lines):
            md_count += 1

    return html_count + md_count

#LaTeX 清洗/还原函数
def clean_ocr_latex_text(text: str) -> str:
    if text is None:
        return ""

    text = str(text)

    # 去掉 markdown 数学公式包裹符号
    text = re.sub(r"\$\s*", "", text)
    text = re.sub(r"\s*\$", "", text)

    # \underline{\text{均匀}} -> 均匀
    text = re.sub(
        r"\\underline\s*\{\s*\\text\s*\{([^{}]*)\}\s*\}",
        r"\1",
        text
    )

    # \underline{56} -> 56
    text = re.sub(
        r"\\underline\s*\{([^{}]*)\}",
        r"\1",
        text
    )

    # \text{均匀} -> 均匀
    text = re.sub(
        r"\\text\s*\{([^{}]*)\}",
        r"\1",
        text
    )

    # \frac{38}{27} -> 38/27
    text = re.sub(
        r"\\frac\s*\{([^{}]*)\}\s*\{([^{}]*)\}",
        r"\1/\2",
        text
    )

    # \times -> ×
    text = text.replace(r"\times", "×")

    # 去掉 LaTeX 空格命令
    text = text.replace(r"\,", " ")
    text = text.replace(r"\ ", " ")

    # 处理多余空格
    text = re.sub(r"\s+", " ", text).strip()

    return text


# =========================================================
# API
# =========================================================
@app.get("/")
def health():
    return {"status": "ok"}


@app.post("/parse-file")
async def parse_file(file: UploadFile = File(...)):
    if not file.filename:
        raise HTTPException(status_code=400, detail="文件名为空")

    suffix = Path(file.filename).suffix.lower()
    if suffix not in ALLOWED_EXTS:
        raise HTTPException(
            status_code=400,
            detail="只支持 PDF、PNG、JPG、JPEG、BMP、WEBP 文件"
        )

    task_id = str(uuid.uuid4())
    input_path = UPLOAD_DIR / f"{task_id}{suffix}"
    task_output_dir = OUTPUT_DIR / task_id
    excel_path = task_output_dir / "result.xlsx"
    debug_path = task_output_dir / "debug.txt"
    merged_md_path = task_output_dir / "merged_markdown_debug.md"
    merged_json_path = task_output_dir / "merged_json_debug.json"

    task_output_dir.mkdir(parents=True, exist_ok=True)

    debug_lines = [
        f"task_id: {task_id}",
        f"filename: {file.filename}",
        f"suffix: {suffix}",
        f"input_path: {input_path}",
        f"task_output_dir: {task_output_dir}",
    ]

    try:
        # 1. 保存上传文件
        with open(input_path, "wb") as f:
            shutil.copyfileobj(file.file, f)

        # 2. OCR
        results = pipeline.predict(str(input_path))

        try:
            result_list = list(results)
        except TypeError:
            result_list = [results]

        debug_lines.append(f"result_count: {len(result_list)}")

        for idx, res in enumerate(result_list):
            debug_lines.append(f"\n===== result {idx} =====")
            debug_lines.append(f"type: {type(res)}")
            debug_lines.append(f"repr: {repr(res)}")

            if hasattr(res, "save_to_markdown"):
                try:
                    res.save_to_markdown(save_path=str(task_output_dir))
                    debug_lines.append("save_to_markdown: success")
                except Exception as e:
                    debug_lines.append(f"save_to_markdown: error -> {e}")

            if hasattr(res, "save_to_json"):
                try:
                    res.save_to_json(save_path=str(task_output_dir))
                    debug_lines.append("save_to_json: success")
                except Exception as e:
                    debug_lines.append(f"save_to_json: error -> {e}")

        # 3. 收集 md/json
        markdown_text, json_text, md_files, json_files = collect_saved_outputs(task_output_dir)

        # 4. 兜底
        if not markdown_text and json_text:
            markdown_text = json_text

        if not markdown_text and not json_text:
            markdown_text = "\n\n".join(str(res) for res in result_list)

        # 5. 调试落盘
        merged_md_path.write_text(markdown_text if markdown_text else "", encoding="utf-8")
        merged_json_path.write_text(json_text if json_text else "{}", encoding="utf-8")

        # 6. 生成 Excel（核心：JSON 优先）
        build_info = build_excel_from_ocr_outputs(markdown_text, json_text, excel_path)
        table_count = count_tables(markdown_text, json_text)

        debug_lines.append(f"md_files: {md_files}")
        debug_lines.append(f"json_files: {json_files}")
        debug_lines.append(f"table_count: {table_count}")
        debug_lines.append(f"excel_build_mode: {build_info['used_mode']}")
        debug_lines.append(f"json_block_count: {build_info['block_count']}")
        debug_lines.append(f"excel_path: {excel_path}")
        debug_lines.append(f"merged_md_path: {merged_md_path}")
        debug_lines.append(f"merged_json_path: {merged_json_path}")

        debug_path.write_text("\n".join(debug_lines), encoding="utf-8")

        # 7. 返回
        excel_b64 = base64.b64encode(excel_path.read_bytes()).decode("utf-8")

        return {
            "success": True,
            "task_id": task_id,
            "filename": file.filename,
            "file_type": suffix,
            "markdown": markdown_text,
            "json_text": json_text,
            "excel": excel_b64,
            "table_count": table_count,
            "md_files": md_files,
            "json_files": json_files,
            "excel_build_mode": build_info["used_mode"],
            "json_block_count": build_info["block_count"],
            "debug_path": str(debug_path),
            "merged_md_path": str(merged_md_path),
            "merged_json_path": str(merged_json_path),
        }

    except Exception as e:
        debug_lines.append(f"error: {repr(e)}")
        debug_path.write_text("\n".join(debug_lines), encoding="utf-8")
        raise HTTPException(status_code=500, detail=str(e))
#import uvicorn
#
#if __name__ == "__main__":
#    uvicorn.run(
#        "app:app",
#        host="0.0.0.0",
#        port=6006,
#        workers=1,
#        log_level="info"
#    )
