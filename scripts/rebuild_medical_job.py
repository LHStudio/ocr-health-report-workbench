from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import local_service as service  # noqa: E402


def rebuild(job_dir: Path, template_path: Path, output_path: Path, parse_mode: str) -> dict[str, Any]:
    excel_root = job_dir / "ocr_excel"
    markdown_root = job_dir / "ocr_markdown"
    if not excel_root.is_dir():
        raise FileNotFoundError(f"Missing OCR Excel directory: {excel_root}")
    if not template_path.is_file():
        raise FileNotFoundError(f"Missing template: {template_path}")

    grouped_excels: dict[str, list[Path]] = {}
    for excel_path in sorted(excel_root.rglob("*.xlsx")):
        relative = excel_path.relative_to(excel_root)
        grouped_excels.setdefault(relative.parent.as_posix(), []).append(excel_path)
    if not grouped_excels:
        raise RuntimeError(f"No OCR Excel files found under {excel_root}")

    workbook = load_workbook(template_path)
    worksheet = workbook.active
    header_row = service.detect_header_row(worksheet)
    headers = [worksheet.cell(header_row, column).value for column in range(1, worksheet.max_column + 1)]
    people_audit: list[dict[str, Any]] = []

    for person, excel_paths in sorted(grouped_excels.items()):
        person_fields: dict[str, str] = {}
        name_candidates: list[tuple[str, str]] = []
        source_files: list[str] = []

        for excel_path in excel_paths:
            relative = excel_path.relative_to(excel_root)
            markdown_path = markdown_root / relative.with_suffix(".md")
            markdown_text = markdown_path.read_text(encoding="utf-8") if markdown_path.is_file() else ""
            parsed_fields = service.parse_medical_page_fields(excel_path.read_bytes(), markdown_text, parse_mode)
            name_candidates.extend(service.extract_person_name_candidates(markdown_text))
            for key, value in parsed_fields.items():
                if service.normalized(key) == service.normalized("姓名"):
                    name = service.valid_person_name(value)
                    if name:
                        name_candidates.append((name, "saved_ocr"))
                else:
                    person_fields.setdefault(key, value)
            source_files.append(relative.as_posix())

        selected_name = service.choose_person_name(name_candidates)
        if selected_name:
            person_fields["姓名"] = selected_name

        target_row = service.first_empty_data_row(worksheet, header_row, headers)
        review_fields = service.create_review_fields(headers, person_fields)
        for field in review_fields:
            if not field["value"]:
                continue
            actual_column = next(
                index + 1 for index, header in enumerate(headers) if service.clean(header) == field["header"]
            )
            worksheet.cell(target_row, actual_column).value = field["value"]

        people_audit.append(
            {
                "person": person,
                "row": target_row,
                "source_files": source_files,
                "filled": sum(bool(field["value"]) for field in review_fields),
                "total": len(review_fields),
                "missing": [field["header"] for field in review_fields if not field["value"]],
                "fields": review_fields,
            }
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()

    result = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "job_dir": str(job_dir.resolve()),
        "template": str(template_path.resolve()),
        "output": str(output_path.resolve()),
        "parse_mode": parse_mode,
        "people_count": len(people_audit),
        "filled": sum(person["filled"] for person in people_audit),
        "total": sum(person["total"] for person in people_audit),
        "missing": sum(len(person["missing"]) for person in people_audit),
        "people": people_audit,
    }
    audit_path = output_path.with_suffix(".audit.json")
    audit_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    result["audit"] = str(audit_path.resolve())
    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Rebuild a medical summary from locally saved OCR artifacts.")
    parser.add_argument("--job-dir", type=Path, required=True)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--parse-mode", choices=("markdown", "excel"), default="markdown")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output_path = args.output
    if output_path is None:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = args.job_dir / "output" / f"{args.template.stem}_本地重解析_{timestamp}.xlsx"
    result = rebuild(args.job_dir.resolve(), args.template.resolve(), output_path.resolve(), args.parse_mode)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
