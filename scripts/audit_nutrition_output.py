from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import local_service as service  # noqa: E402


def sheet_rows(workbook, keyword: str) -> tuple[list[str], list[dict[str, str]]]:
    worksheet = next((sheet for sheet in workbook.worksheets if keyword in sheet.title), None)
    if worksheet is None:
        raise RuntimeError(f"Could not find worksheet containing {keyword!r}")
    header_row = service.detect_header_row(worksheet)
    headers = [service.clean(worksheet.cell(header_row, column).value) for column in range(1, worksheet.max_column + 1)]
    rows = []
    for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        row = {header: service.clean(values[index] if index < len(values) else "") for index, header in enumerate(headers) if header}
        if any(row.values()):
            rows.append(row)
    return headers, rows


def audit(workbook_path: Path) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    _, general_rows = sheet_rows(workbook, "人员")
    _, food_rows = sheet_rows(workbook, "食物")
    _, supplement_rows = sheet_rows(workbook, "保健")
    workbook.close()

    people = [row.get("人员文件夹", "") for row in general_rows if row.get("人员文件夹")]
    people_set = set(people)
    general_missing = {
        field: [row.get("人员文件夹", "") for row in general_rows if not row.get(field)]
        for field in service.NUTRITION_GENERAL_COLUMNS
        if field not in {"人员文件夹", "人工备注"}
    }
    meal_conflicts = []
    for row in general_rows:
        for field in ("早餐地点", "午餐地点", "晚餐地点"):
            value = row.get(field, "")
            if "不吃" in value and any(option in value for option in ("家", "学校食堂", "餐馆或街头")):
                meal_conflicts.append({"person": row.get("人员文件夹", ""), "field": field, "value": value})

    date_warnings = []
    for row in general_rows:
        value = row.get("调查日期", "")
        if value and not re.fullmatch(r"20\d{2}[./年-]\d{1,2}[./月-]\d{1,2}日?", value):
            date_warnings.append({"person": row.get("人员文件夹", ""), "value": value})

    food_by_person: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in food_rows:
        food_by_person[row.get("人员文件夹", "")].append(row)
    food_structure = []
    expected_codes = set(service.NUTRITION_FOOD_ITEMS)
    for person in sorted(people_set):
        rows = food_by_person.get(person, [])
        codes = [row.get("食物编号", "") for row in rows]
        counts = Counter(codes)
        missing = sorted(expected_codes - set(codes))
        extras = sorted(set(codes) - expected_codes)
        duplicates = sorted(code for code, count in counts.items() if code and count > 1)
        if len(rows) != len(expected_codes) or missing or extras or duplicates:
            food_structure.append({"person": person, "rows": len(rows), "missing": missing, "extras": extras, "duplicates": duplicates})

    food_review = [
        {
            "person": row.get("人员文件夹", ""), "code": row.get("食物编号", ""),
            "name": row.get("食物名称", ""), "quantity": row.get("平均每次食用量", ""),
            "count": row.get("次数", ""), "period": row.get("频率周期(请核对)", ""),
            "note": row.get("人工核对备注", ""), "raw": row.get("OCR原始行", ""),
        }
        for row in food_rows
        if service.nutrition_food_note_requires_review(row.get("人工核对备注"))
        or row.get("频率周期(请核对)") == "未识别"
    ]
    invalid_food_periods = [
        {"person": row.get("人员文件夹", ""), "code": row.get("食物编号", ""), "period": row.get("频率周期(请核对)", "")}
        for row in food_rows
        if row.get("频率周期(请核对)") not in {*service.NUTRITION_PERIODS, "未识别", ""}
    ]

    supplement_by_person: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in supplement_rows:
        supplement_by_person[row.get("人员文件夹", "")].append(row)
    expected_supplements = set(service.NUTRITION_SUPPLEMENT_ITEMS)
    supplement_structure = []
    for person in sorted(people_set):
        rows = supplement_by_person.get(person, [])
        names = [row.get("保健品种类", "") for row in rows]
        counts = Counter(names)
        missing = sorted(expected_supplements - set(names))
        extras = sorted(set(names) - expected_supplements)
        duplicates = sorted(name for name, count in counts.items() if name and count > 1)
        if len(rows) != len(expected_supplements) or missing or extras or duplicates:
            supplement_structure.append({"person": person, "rows": len(rows), "missing": missing, "extras": extras, "duplicates": duplicates})
    supplement_review = [
        {
            "person": row.get("人员文件夹", ""), "category": row.get("保健品种类", ""),
            "name": row.get("保健品名称", ""), "quantity": row.get("平均每次服用量", ""),
            "count": row.get("次数", ""), "period": row.get("频率周期(请核对)", ""),
            "note": row.get("备注", ""),
        }
        for row in supplement_rows
        if row.get("备注") or (row.get("频率周期(请核对)") == "未识别" and any(row.get(key) for key in ("保健品名称", "平均每次服用量", "次数")))
    ]

    return {
        "workbook": str(workbook_path.resolve()),
        "people_count": len(people),
        "duplicate_people": sorted(person for person, count in Counter(people).items() if count > 1),
        "general_filled": {field: len(people) - len(missing) for field, missing in general_missing.items()},
        "general_missing": general_missing,
        "meal_conflicts": meal_conflicts,
        "date_warnings": date_warnings,
        "food": {
            "rows": len(food_rows),
            "expected_rows": len(people) * len(expected_codes),
            "periods": dict(Counter(row.get("频率周期(请核对)", "") for row in food_rows)),
            "structural_issues": food_structure,
            "invalid_periods": invalid_food_periods,
            "review_count": len(food_review),
            "review_rows": food_review,
        },
        "supplements": {
            "rows": len(supplement_rows),
            "expected_rows": len(people) * len(expected_supplements),
            "periods": dict(Counter(row.get("频率周期(请核对)", "") for row in supplement_rows)),
            "structural_issues": supplement_structure,
            "review_count": len(supplement_review),
            "review_rows": supplement_review,
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Audit structure, missing fields, and review risks in a nutrition summary workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--json-out", type=Path)
    args = parser.parse_args()
    result = audit(args.workbook.resolve())
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(payload, encoding="utf-8")
    print(payload)


if __name__ == "__main__":
    main()
