# -*- coding: utf-8 -*-
# @Time: 2026/3/31 16:07
# @Author: LHStudio
# @File: app.py
# @Software: PyCharm

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from pathlib import Path
import shutil
import uuid
import json
import base64
import mimetypes
import os
import re
import threading
from time import perf_counter
from typing import List, Tuple, Optional

import cv2
import numpy as np
import pypdfium2 as pdfium
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from paddleocr import PaddleOCR, PaddleOCRVL

from fastapi.middleware.cors import CORSMiddleware

app = FastAPI(title="OCR Parse Service")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://127.0.0.1:3000",
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# =========================================================
# OCR 初始化
# =========================================================
# 部署兼容：
# - 当前裸机服务器默认使用 127.0.0.1:8118/v1 和 /root/pdf_to_excel_service；
# - 以下是对方 Docker 版的原始调用地址与目录，本服务器保留为注释，不能直接启用：
#     vl_rec_server_url="http://172.18.0.2:8118/v1"
#     BASE_DIR = Path("/app/pdf_to_excel_service")
# - Docker 环境如需复用本增强版，请设置这两个环境变量，而无需修改本文件：
#     PADDLEOCR_VL_SERVER_URL=http://172.18.0.2:8118/v1
#     OCR_SERVICE_BASE_DIR=/app/pdf_to_excel_service
VL_REC_SERVER_URL = os.getenv("PADDLEOCR_VL_SERVER_URL", "http://127.0.0.1:8118/v1").strip().rstrip("/")
if not VL_REC_SERVER_URL:
    raise RuntimeError("PADDLEOCR_VL_SERVER_URL 不能为空")

pipeline = PaddleOCRVL(
    vl_rec_backend="vllm-server",
    vl_rec_server_url=VL_REC_SERVER_URL,
)
classic_ocr_model = None
classic_ocr_lock = threading.Lock()

# =========================================================
# 路径配置
# =========================================================
BASE_DIR = Path(os.getenv("OCR_SERVICE_BASE_DIR", "/root/pdf_to_excel_service")).expanduser()
UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "outputs"
QUESTIONNAIRE_TEMPLATE = Path(__file__).resolve().parent / "questionnaire_templates" / "food_frequency_blank_reference.png"
SUNLIGHT_QUESTIONNAIRE_TEMPLATE = Path(__file__).resolve().parent / "questionnaire_templates" / "food_frequency_sunlight_blank_reference.png"
SUNLIGHT_REAL_SCAN_TEMPLATE = Path(__file__).resolve().parent / "questionnaire_templates" / "food_frequency_sunlight_real_scan_reference.png"

UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

ALLOWED_EXTS = {".pdf", ".png", ".jpg", ".jpeg", ".bmp", ".webp"}


# =========================================================
# 基础工具
# =========================================================
def safe_read_text(file_path: Path) -> str:
    try:
        return file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return file_path.read_text(encoding="utf-8", errors="ignore")


def normalize_text(text: str) -> str:
    if text is None:
        return ""
    return text.replace("\r\n", "\n").replace("\r", "\n")


def extract_element_info(element) -> str:
    """
    递归提取元素内的文本信息；对 OCR 已明确标注的勾选图片保留“☑”。

    说明：这不会凭空判断没有被模型识别到的手写勾选，而是避免在 HTML 表格转 Excel
    时把已有的 alt/title/aria-label 勾选信息丢掉，供下游人工核对和规则解析使用。
    """
    if isinstance(element, str):
        return element.strip()

    if getattr(element, "name", None) == "input" and str(element.get("type", "")).lower() == "checkbox":
        if element.has_attr("checked") or str(element.get("aria-checked", "")).lower() == "true":
            return "☑"
        return ""

    if getattr(element, "name", None) == "img":
        image_hint = " ".join(str(element.get(key, "")) for key in ("alt", "title", "aria-label", "class", "data-state"))
        if re.search(r"(?:unchecked|unselected|not[-_ ]?checked|未选|未勾|☐)", image_hint, re.I):
            return ""
        if str(element.get("aria-checked", "")).lower() == "false" or str(element.get("data-state", "")).lower() == "unchecked":
            return ""
        if str(element.get("aria-checked", "")).lower() == "true" or str(element.get("data-state", "")).lower() == "checked":
            return "☑"
        if re.search(r"(?:\bchecked\b|\bselected\b|\btick(?:ed)?\b|勾选|已选|√|✓|✔|☑)", image_hint, re.I):
            return "☑"
        return ""

    result = []
    if hasattr(element, "children"):
        for child in element.children:
            text = extract_element_info(child)
            if text:
                result.append(text)

    return " ".join(result).strip()


# =========================================================
# HTML 表格解析
# =========================================================
def parse_html_table_to_excel(table_soup, ws, start_row: int) -> int:
    """
    解析 HTML 表格，保留合并网格和对齐
    """
    row_cursor = start_row
    occupied_cells = {}

    rows = table_soup.find_all("tr")
    if not rows:
        return row_cursor

    for tr in rows:
        col_cursor = 1
        cells = tr.find_all(["td", "th"], recursive=False)
        if not cells:
            cells = tr.find_all(["td", "th"])

        for cell in cells:
            while occupied_cells.get((row_cursor, col_cursor), False):
                col_cursor += 1

            rowspan = int(cell.get("rowspan", 1) or 1)
            colspan = int(cell.get("colspan", 1) or 1)

            raw_text = extract_element_info(cell)

            clean_text = clean_ocr_latex_text(raw_text)
            clean_text = re.sub(r"\*\*(.*?)\*\*", r"\1", clean_text).strip()

            excel_cell = ws.cell(row=row_cursor, column=col_cursor, value=clean_text)

            style_attr = (cell.get("style", "") or "").lower()
            if "text-align: center" in style_attr:
                excel_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                excel_cell.alignment = Alignment(vertical="center", wrap_text=True)

            if cell.name == "th":
                excel_cell.font = Font(bold=True)

            if rowspan > 1 or colspan > 1:
                ws.merge_cells(
                    start_row=row_cursor,
                    start_column=col_cursor,
                    end_row=row_cursor + rowspan - 1,
                    end_column=col_cursor + colspan - 1
                )
                for r in range(row_cursor, row_cursor + rowspan):
                    for c in range(col_cursor, col_cursor + colspan):
                        occupied_cells[(r, c)] = True
            else:
                occupied_cells[(row_cursor, col_cursor)] = True

            col_cursor += colspan

        row_cursor += 1

    return row_cursor


# =========================================================
# Markdown 表格解析（兜底）
# =========================================================
def split_md_row(line: str) -> List[str]:
    line = line.strip()
    if not line:
        return []

    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]

    return [part.strip() for part in line.split("|")]


def is_markdown_table_separator(line: str) -> bool:
    cells = split_md_row(line)
    if not cells:
        return False

    for cell in cells:
        test_cell = cell.replace(" ", "")
        if not re.fullmatch(r":?-{3,}:?", test_cell):
            return False
    return True


def is_markdown_table_block(lines: List[str]) -> bool:
    if len(lines) < 2:
        return False

    first = lines[0].strip()
    second = lines[1].strip()

    if "|" not in first or "|" not in second:
        return False

    return is_markdown_table_separator(second)


def parse_markdown_table_to_excel(table_lines: List[str], ws, start_row: int) -> int:
    if len(table_lines) < 2:
        return start_row

    header = split_md_row(table_lines[0])
    data_lines = table_lines[2:] if is_markdown_table_separator(table_lines[1]) else table_lines[1:]

    row_cursor = start_row

    for col_idx, value in enumerate(header, start=1):
        cell = ws.cell(row=row_cursor, column=col_idx, value=clean_ocr_latex_text(value))
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_cursor += 1

    for line in data_lines:
        cols = split_md_row(line)
        if not cols:
            continue
        for col_idx, value in enumerate(cols, start=1):
            cell = ws.cell(row=row_cursor, column=col_idx, value=clean_ocr_latex_text(value))
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        row_cursor += 1

    return row_cursor


# =========================================================
# 普通文本写入
# =========================================================
def write_text_line(ws, row_idx: int, text: str) -> int:
    clean_line = clean_ocr_latex_text(text).strip()
    if not clean_line:
        return row_idx

    indent_level = 0
    is_bold = False
    font_size = 11

    heading_match = re.match(r"^(#{1,6})\s+(.*)", clean_line)
    if heading_match:
        level = len(heading_match.group(1))
        clean_line = heading_match.group(2).strip()
        is_bold = True
        font_size = max(11, 16 - level)
        indent_level = max(0, level - 1)

    elif re.match(r"^(\d+\.|[a-zA-Z]\.|[-*+])\s+", clean_line):
        indent_level = 1

    if "**" in clean_line:
        clean_line = re.sub(r"\*\*(.*?)\*\*", r"\1", clean_line)
        is_bold = True

    cell = ws.cell(row=row_idx, column=1, value=clean_line)
    cell.font = Font(bold=is_bold, size=font_size)

    if indent_level > 0:
        cell.alignment = Alignment(indent=indent_level, vertical="center", wrap_text=True)
    else:
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    return row_idx + 1


def process_plain_text_block(block_text: str, ws, start_row: int) -> int:
    row_cursor = start_row
    lines = normalize_text(block_text).split("\n")

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        maybe_table = [line]
        j = i + 1
        while j < len(lines) and lines[j].strip():
            maybe_table.append(lines[j].rstrip())
            j += 1

        if is_markdown_table_block(maybe_table):
            row_cursor = parse_markdown_table_to_excel(maybe_table, ws, row_cursor)
            row_cursor += 1
            i = j
            continue

        row_cursor = write_text_line(ws, row_cursor, line)
        i += 1

    return row_cursor


# =========================================================
# 混合 markdown/html 内容处理（兜底）
# =========================================================
def process_markdown_text_to_sheet(md_text: str, ws, start_row: int = 1) -> int:
    md_text = normalize_text(md_text)
    current_row = start_row

    table_pattern = re.compile(r"(<table.*?</table>)", flags=re.IGNORECASE | re.DOTALL)
    parts = table_pattern.split(md_text)

    for part in parts:
        if not part or not part.strip():
            continue

        if re.match(r"^\s*<table", part, flags=re.IGNORECASE):
            soup = BeautifulSoup(part, "html.parser")
            table = soup.find("table")
            if table:
                current_row = parse_html_table_to_excel(table, ws, current_row)
                current_row += 1
            else:
                current_row = process_plain_text_block(part, ws, current_row)
            continue

        current_row = process_plain_text_block(part, ws, current_row)

    return current_row


# =========================================================
# JSON block 解析（核心修复）
# =========================================================
def load_json_objects_from_text(raw_text: str) -> List[dict]:
    """
    支持：
    1. 单个 JSON 对象
    2. 多个 JSON 对象拼接
    3. JSON 数组
    """
    raw_text = raw_text.strip()
    if not raw_text:
        return []

    # 先尝试整体解析
    try:
        data = json.loads(raw_text)
        if isinstance(data, list):
            return [item for item in data if isinstance(item, dict)]
        if isinstance(data, dict):
            return [data]
    except Exception:
        pass

    # 尝试按大括号逐个解析
    decoder = json.JSONDecoder()
    idx = 0
    results = []
    while idx < len(raw_text):
        while idx < len(raw_text) and raw_text[idx].isspace():
            idx += 1
        if idx >= len(raw_text):
            break
        try:
            obj, end = decoder.raw_decode(raw_text, idx)
            if isinstance(obj, dict):
                results.append(obj)
            elif isinstance(obj, list):
                results.extend([item for item in obj if isinstance(item, dict)])
            idx = end
        except Exception:
            break

    return results


def extract_blocks_from_json_text(json_text: str) -> List[dict]:
    """
    从 OCR json 文本中提取 block，按 block_order 排序
    兼容结构：
    {
        "parsing_res_list": [...]
    }
    关键修复：
    - 不再对所有 blocks 做全局排序；
    - 每个 JSON 对象/页面内部按 block_order 排序；
    - 然后按 JSON 原始出现顺序追加，避免多页 block_order 重复导致错序。
    或其他嵌套结构
    """
    objs = load_json_objects_from_text(json_text)
    ordered_blocks = []

    def sort_key(x):
        order = x.get("block_order")
        block_id = x.get("block_id", 10 ** 9)

        if order is None:
            order = 10 ** 9

        return order, block_id

    for obj_index, obj in enumerate(objs):
        if not isinstance(obj, dict):
            continue

        page_blocks = []

        # 标准结构
        if isinstance(obj.get("parsing_res_list"), list):
            for block in obj["parsing_res_list"]:
                if isinstance(block, dict):
                    page_blocks.append(block)

        # 兼容其他可能字段
        for key in ["blocks", "results", "data"]:
            value = obj.get(key)
            if isinstance(value, list):
                for item in value:
                    if isinstance(item, dict) and (
                            "block_label" in item or "block_content" in item
                    ):
                        page_blocks.append(item)

        # 只在当前 JSON 对象内部排序
        page_blocks.sort(key=sort_key)

        # 按 JSON 对象原始顺序追加
        ordered_blocks.extend(page_blocks)

    return ordered_blocks


def process_blocks_to_sheet(blocks: List[dict], ws, start_row: int = 1) -> int:
    """
    优先按 OCR block 顺序重建 Excel
    """
    current_row = start_row

    for block in blocks:
        block_label = str(block.get("block_label", "")).strip().lower()
        block_content = block.get("block_content", "")

        if not block_content:
            continue

        # 表格 block：最关键，优先处理
        if block_label == "table":
            content = str(block_content).strip()

            # HTML table
            if re.search(r"<table.*?</table>", content, flags=re.IGNORECASE | re.DOTALL):
                soup = BeautifulSoup(content, "html.parser")
                table = soup.find("table")
                if table:
                    current_row = parse_html_table_to_excel(table, ws, current_row)
                    current_row += 1
                    continue

            # markdown table
            lines = [line for line in normalize_text(content).split("\n") if line.strip()]
            if is_markdown_table_block(lines):
                current_row = parse_markdown_table_to_excel(lines, ws, current_row)
                current_row += 1
                continue

            # 表格块但解析失败，按文本兜底
            current_row = process_plain_text_block(content, ws, current_row)
            current_row += 1
            continue

        # 文本块
        if block_label in {"text", "title", "paragraph", "list"}:
            current_row = process_plain_text_block(str(block_content), ws, current_row)
            continue

        # 其他未知 block，兜底成文本
        current_row = process_plain_text_block(str(block_content), ws, current_row)

    return current_row


# =========================================================
# 样式
# =========================================================
def set_sheet_style(ws):
    max_col = max(ws.max_column, 7)

    for col_idx in range(1, max_col + 1):
        col_letter = get_column_letter(col_idx)
        if col_idx == 1:
            ws.column_dimensions[col_letter].width = 60
        else:
            ws.column_dimensions[col_letter].width = 18

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                current_alignment = cell.alignment
                if current_alignment:
                    cell.alignment = Alignment(
                        horizontal=current_alignment.horizontal,
                        vertical=current_alignment.vertical or "center",
                        wrap_text=True,
                        indent=current_alignment.indent or 0
                    )
                else:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)


# =========================================================
# 收集 OCR 输出
# =========================================================
def collect_saved_outputs(task_output_dir: Path) -> Tuple[str, str, List[str], List[str]]:
    md_files = sorted(task_output_dir.rglob("*.md"))
    json_files = sorted(task_output_dir.rglob("*.json"))

    markdown_parts = []
    json_parts = []

    for md_file in md_files:
        content = safe_read_text(md_file)
        if content.strip():
            markdown_parts.append(content)

    for json_file in json_files:
        content = safe_read_text(json_file)
        if content.strip():
            json_parts.append(content)

    markdown_text = "\n\n".join(markdown_parts).strip()
    json_text = "\n\n".join(json_parts).strip()

    return markdown_text, json_text, [str(p) for p in md_files], [str(p) for p in json_files]


def valid_person_name(value: str) -> str:
    candidate = re.sub(r"\s+", "", str(value or "")).strip("_＿—-：:")
    if candidate in {"", "无", "未签", "未填写"} or "_" in candidate or "＿" in candidate:
        return ""
    if any(word in candidate for word in ("姓名", "调查", "日期", "编号", "签字", "签名", "受试者", "研究者", "患者", "性别", "年龄", "报告", "检验", "访视")):
        return ""
    return candidate if re.fullmatch(r"[\u3400-\u9fff·]{2,8}", candidate) else ""


def extract_identity_candidates(text: str, source: str) -> List[dict]:
    """从主 OCR 或整页二次 OCR 中提取姓名候选；不接受空白线和下一字段粘连。"""
    terminator = r"(?=\s*(?:调查日期|报告日期|访视号|性别|别[：:]?|年龄|龄[：:]?|病历号|病床号|样本编号|采样时间|日期|[，,；;。|<>\n]|$))"
    patterns = [
        ("name", rf"(?<!研究者)(?:(?:受试者|患者|受检者)\s*)?姓名\s*[：:]?\s*([\u3400-\u9fff·\s_＿—-]{{2,18}}?){terminator}"),
        ("misaligned_name", rf"(?<!姓)(?<!研究者)名\s*[：:]\s*([\u3400-\u9fff·\s_＿—-]{{2,18}}?){terminator}"),
        ("subject_signature", rf"(?:受试者|患者|受检者)\s*(?:签字|签名)\s*[：:]?\s*([\u3400-\u9fff·\s_＿—-]{{2,18}}?){terminator}"),
    ]
    result = []
    for method, pattern in patterns:
        for match in re.finditer(pattern, normalize_text(text or "")):
            name = valid_person_name(match.group(1))
            if name and not any(item["value"] == name and item["method"] == method for item in result):
                result.append({
                    "value": name,
                    "source": source,
                    "method": method,
                    "confidence": 0.96 if method == "name" and source == "full_page_ocr" else (0.9 if method == "name" else 0.82),
                })
    return result


def choose_identity_name(candidates: List[dict]) -> str:
    scores = {}
    for item in candidates:
        name = valid_person_name(item.get("value", ""))
        if name:
            scores[name] = scores.get(name, 0.0) + float(item.get("confidence", 0.5))
    if not scores:
        return ""
    ranking = sorted(scores.items(), key=lambda item: (-item[1], item[0]))
    if len(ranking) > 1 and abs(ranking[0][1] - ranking[1][1]) < 0.05:
        return ""
    return ranking[0][0]


def run_full_page_identity_ocr(input_path: Path, task_output_dir: Path) -> Tuple[str, List[dict]]:
    """主版面 OCR 漏掉姓名时，用整页上下文再识别一次；只抽取身份信息，不替换主表格。"""
    enhanced_dir = task_output_dir / "enhanced_identity"
    enhanced_dir.mkdir(parents=True, exist_ok=True)
    results = pipeline.predict(
        str(input_path),
        use_layout_detection=False,
        prompt_label="ocr",
        use_doc_orientation_classify=True,
        use_doc_unwarping=True,
        temperature=0.0,
    )
    for result in list(results):
        if hasattr(result, "save_to_markdown"):
            result.save_to_markdown(save_path=str(enhanced_dir))
        if hasattr(result, "save_to_json"):
            result.save_to_json(save_path=str(enhanced_dir))
    markdown = "\n\n".join(safe_read_text(path) for path in sorted(enhanced_dir.rglob("*.md"))).strip()
    return markdown, extract_identity_candidates(markdown, "full_page_ocr")


def run_food_header_identity_ocr(input_path: Path, task_output_dir: Path) -> Tuple[str, List[dict]]:
    """整页仍无有效姓名时，裁剪固定问卷页眉并放大识别，避免表格内容干扰手写姓名。"""
    if not QUESTIONNAIRE_TEMPLATE.exists():
        return "", []
    reference = cv2.imread(str(QUESTIONNAIRE_TEMPLATE), cv2.IMREAD_GRAYSCALE)
    if reference is None:
        return "", []
    image = render_document_gray(input_path, (reference.shape[1], reference.shape[0]))
    aligned, alignment_confidence = align_questionnaire(image, reference)
    if alignment_confidence < 0.35:
        return "", []
    height, width = aligned.shape
    crop = aligned[int(height * 0.025):int(height * 0.14), int(width * 0.27):int(width * 0.69)]
    crop = cv2.resize(crop, None, fx=3.0, fy=3.0, interpolation=cv2.INTER_CUBIC)
    enhanced_dir = task_output_dir / "enhanced_header_identity"
    enhanced_dir.mkdir(parents=True, exist_ok=True)
    crop_path = enhanced_dir / "questionnaire_header.png"
    cv2.imwrite(str(crop_path), crop)
    results = pipeline.predict(
        str(crop_path),
        use_layout_detection=False,
        prompt_label="ocr",
        use_doc_orientation_classify=False,
        use_doc_unwarping=False,
        temperature=0.0,
    )
    for result in list(results):
        if hasattr(result, "save_to_markdown"):
            result.save_to_markdown(save_path=str(enhanced_dir))
    markdown = "\n\n".join(
        safe_read_text(path)
        for path in sorted(enhanced_dir.rglob("*.md"))
        if path.name != crop_path.name
    ).strip()
    return markdown, extract_identity_candidates(markdown, "header_crop_ocr")


def run_body_composition_header_ocr(input_path: Path, task_output_dir: Path) -> Tuple[str, List[dict]]:
    """Crop and re-read the InBody header containing ID, height, age, and sex."""
    if input_path.suffix.lower() == ".pdf":
        document = pdfium.PdfDocument(str(input_path))
        try:
            array = document[0].render(scale=3.0).to_numpy()
        finally:
            document.close()
        gray = cv2.cvtColor(array, cv2.COLOR_RGBA2GRAY if array.shape[2] == 4 else cv2.COLOR_RGB2GRAY)
    else:
        raw = np.frombuffer(input_path.read_bytes(), dtype=np.uint8)
        gray = cv2.imdecode(raw, cv2.IMREAD_GRAYSCALE)
        if gray is None:
            raise ValueError("无法读取体成分报告图像")

    height, width = gray.shape
    crop = gray[:max(1, int(height * 0.18)), :]
    if width < 2400:
        crop = cv2.resize(crop, None, fx=1.5, fy=1.5, interpolation=cv2.INTER_CUBIC)
    enhanced_dir = task_output_dir / "enhanced_body_composition_header"
    enhanced_dir.mkdir(parents=True, exist_ok=True)
    crop_path = enhanced_dir / "body_composition_header.png"
    cv2.imwrite(str(crop_path), crop)
    results = pipeline.predict(
        str(crop_path),
        use_layout_detection=False,
        prompt_label="ocr",
        use_doc_orientation_classify=False,
        use_doc_unwarping=False,
        temperature=0.0,
    )
    for result in list(results):
        if hasattr(result, "save_to_markdown"):
            result.save_to_markdown(save_path=str(enhanced_dir))
        if hasattr(result, "save_to_json"):
            result.save_to_json(save_path=str(enhanced_dir))
    markdown = "\n\n".join(safe_read_text(path) for path in sorted(enhanced_dir.rglob("*.md"))).strip()
    return markdown, extract_identity_candidates(markdown, "body_composition_header_ocr")


def collect_image_assets(task_output_dir: Path, max_files: int = 20, max_total_bytes: int = 8 * 1024 * 1024) -> List[dict]:
    """增量返回 OCR 产生的小图资产；旧响应字段保持不变。"""
    assets = []
    total = 0
    image_paths = sorted(
        path for path in task_output_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
    )
    for path in image_paths:
        size = path.stat().st_size
        if len(assets) >= max_files or total + size > max_total_bytes:
            break
        raw = path.read_bytes()
        bbox_match = re.search(r"(?:box|bbox)_([0-9]+)_([0-9]+)_([0-9]+)_([0-9]+)", path.stem)
        assets.append({
            "path": path.relative_to(task_output_dir).as_posix(),
            "mime": mimetypes.guess_type(path.name)[0] or "application/octet-stream",
            "base64": base64.b64encode(raw).decode("ascii"),
            "bbox": [int(value) for value in bbox_match.groups()] if bbox_match else None,
        })
        total += size
    return assets


def render_document_gray(input_path: Path, target_size: Tuple[int, int]) -> np.ndarray:
    if input_path.suffix.lower() == ".pdf":
        document = pdfium.PdfDocument(str(input_path))
        array = document[0].render(scale=2.0).to_numpy()
        if array.shape[2] == 4:
            gray = cv2.cvtColor(array, cv2.COLOR_RGBA2GRAY)
        else:
            gray = cv2.cvtColor(array, cv2.COLOR_RGB2GRAY)
    else:
        raw = np.frombuffer(input_path.read_bytes(), dtype=np.uint8)
        image = cv2.imdecode(raw, cv2.IMREAD_GRAYSCALE)
        if image is None:
            raise ValueError("无法读取问卷图像")
        gray = image
    return cv2.resize(gray, target_size, interpolation=cv2.INTER_AREA)


def align_questionnaire(image: np.ndarray, reference: np.ndarray) -> Tuple[np.ndarray, float]:
    detector = cv2.ORB_create(5000)
    key_a, desc_a = detector.detectAndCompute(image, None)
    key_b, desc_b = detector.detectAndCompute(reference, None)
    if desc_a is None or desc_b is None:
        return image, 0.0
    matches = cv2.BFMatcher(cv2.NORM_HAMMING).knnMatch(desc_a, desc_b, k=2)
    good = [first for first, second in matches if first.distance < 0.72 * second.distance]
    if len(good) < 30:
        return image, 0.0
    source = np.float32([key_a[item.queryIdx].pt for item in good]).reshape(-1, 1, 2)
    target = np.float32([key_b[item.trainIdx].pt for item in good]).reshape(-1, 1, 2)
    matrix, mask = cv2.findHomography(source, target, cv2.RANSAC, 4.0)
    if matrix is None or mask is None or int(mask.sum()) < 20:
        return image, 0.0
    aligned = cv2.warpPerspective(image, matrix, (reference.shape[1], reference.shape[0]), borderValue=255)
    return aligned, float(mask.sum()) / max(len(good), 1)


def detect_food_questionnaire_checks(input_path: Path) -> Tuple[dict, List[dict]]:
    """固定版式食物频率问卷：模板对齐后检测问题 3-5 选项区新增手写墨迹。"""
    if not QUESTIONNAIRE_TEMPLATE.exists():
        return {}, []
    reference = cv2.imread(str(QUESTIONNAIRE_TEMPLATE), cv2.IMREAD_GRAYSCALE)
    if reference is None:
        return {}, []
    image = render_document_gray(input_path, (reference.shape[1], reference.shape[0]))
    aligned, alignment_confidence = align_questionnaire(image, reference)
    if alignment_confidence < 0.35:
        return {"warning": "questionnaire_alignment_low", "alignment_confidence": round(alignment_confidence, 3)}, []

    # 铅笔勾画常比印刷字浅；提高前景阈值，并扩大空白模板文字遮罩，保留括号附近的新笔画。
    current_ink = cv2.threshold(aligned, 225, 255, cv2.THRESH_BINARY_INV)[1]
    reference_ink = cv2.threshold(reference, 235, 255, cv2.THRESH_BINARY_INV)[1]
    reference_ink = cv2.dilate(reference_ink, np.ones((5, 5), np.uint8), iterations=1)
    novel_ink = cv2.bitwise_and(current_ink, cv2.bitwise_not(reference_ink))
    novel_ink = cv2.morphologyEx(novel_ink, cv2.MORPH_OPEN, np.ones((2, 2), np.uint8))

    question_rows = {"breakfast": (0.258, 0.284), "lunch": (0.313, 0.339), "dinner": (0.369, 0.395)}
    option_columns = {
        "home": (0.180, 0.295),
        "school": (0.295, 0.455),
        "restaurant": (0.455, 0.625),
        "skip": (0.625, 0.790),
    }
    option_labels = {"home": "家", "school": "学校食堂", "restaurant": "餐馆或街头", "skip": "不吃"}
    question_labels = {"breakfast": "早餐地点", "lunch": "午餐地点", "dinner": "晚餐地点"}
    height, width = reference.shape
    locations = {}
    visual_elements = []
    for question, (y_start, y_end) in question_rows.items():
        measurements = []
        for option, (x_start, x_end) in option_columns.items():
            x0, x1 = int(width * x_start), int(width * x_end)
            y0, y1 = int(height * y_start), int(height * y_end)
            roi = novel_ink[y0:y1, x0:x1]
            score = int(np.count_nonzero(roi))
            _, _, stats, _ = cv2.connectedComponentsWithStats(roi, connectivity=8)
            components = [item for item in stats[1:] if item[cv2.CC_STAT_AREA] >= 4]
            largest = max(components, key=lambda item: item[cv2.CC_STAT_AREA]) if components else [0, 0, 0, 0, 0]
            component = {
                "area": int(largest[cv2.CC_STAT_AREA]),
                "width": int(largest[cv2.CC_STAT_WIDTH]),
                "height": int(largest[cv2.CC_STAT_HEIGHT]),
                "count": len(components),
            }
            measurements.append((option, score, component, [x0, y0, x1, y1]))
        maximum = max((score for _, score, _, _ in measurements), default=0)
        maximum_area = max((component["area"] for _, _, component, _ in measurements), default=0)
        selected, uncertain = [], []
        for option, score, component, bbox in measurements:
            area = component["area"]
            component_width = component["width"]
            component_height = component["height"]
            aspect = component_width / max(component_height, 1)
            relative_area = area / max(maximum_area, 1)
            strong_stroke = (
                area >= 16 and component_width >= 4 and component_height >= 5
                and 0.3 <= aspect <= 6.0 and relative_area >= 0.33
            )
            weak_stroke = (
                area >= 8 and component_width >= 3 and component_height >= 3
                and 0.2 <= aspect <= 7.0 and relative_area >= 0.2
            )
            if strong_stroke:
                state = True
                selected.append(option_labels[option])
                confidence = min(0.98, 0.72 + area / 500)
            elif weak_stroke:
                state = None
                uncertain.append(option_labels[option])
                confidence = min(0.69, 0.38 + area / 300)
            else:
                state = False
                confidence = min(0.95, 0.72 + (16 - min(area, 16)) / 80)
            visual_elements.append({
                "label": f"{question_labels[question]}-{option_labels[option]}",
                "selected": state,
                "confidence": round(confidence * alignment_confidence, 3),
                "bbox": bbox,
                "score": score,
                "component": component,
                "method": "template_component_difference_v2",
            })
        locations[question] = {"selected": selected, "uncertain": uncertain, "max_score": maximum}
    return {"meal_locations": locations, "alignment_confidence": round(alignment_confidence, 3), "method": "template_component_difference_v2"}, visual_elements


def _sunlight_option_measurement(raw_ink: np.ndarray, novel_ink: np.ndarray, bbox: List[int]) -> dict:
    x0, y0, x1, y1 = bbox
    raw_roi = raw_ink[y0:y1, x0:x1]
    novel_roi = novel_ink[y0:y1, x0:x1]
    _, _, raw_stats, _ = cv2.connectedComponentsWithStats(raw_roi, connectivity=8)
    _, _, novel_stats, _ = cv2.connectedComponentsWithStats(novel_roi, connectivity=8)
    raw_components = [item for item in raw_stats[1:] if item[cv2.CC_STAT_AREA] >= 4]
    novel_components = [item for item in novel_stats[1:] if item[cv2.CC_STAT_AREA] >= 4]
    largest_raw = max(raw_components, key=lambda item: item[cv2.CC_STAT_AREA]) if raw_components else [0, 0, 0, 0, 0]
    largest_novel = max(novel_components, key=lambda item: item[cv2.CC_STAT_AREA]) if novel_components else [0, 0, 0, 0, 0]
    return {
        "raw_components": raw_components,
        "raw": {
            "area": int(largest_raw[cv2.CC_STAT_AREA]),
            "width": int(largest_raw[cv2.CC_STAT_WIDTH]),
            "height": int(largest_raw[cv2.CC_STAT_HEIGHT]),
            "count": len(raw_components),
        },
        "novel": {
            "score": int(np.count_nonzero(novel_roi)),
            "area": int(largest_novel[cv2.CC_STAT_AREA]),
            "width": int(largest_novel[cv2.CC_STAT_WIDTH]),
            "height": int(largest_novel[cv2.CC_STAT_HEIGHT]),
            "count": len(novel_components),
        },
    }


def _sunlight_mark_state(raw_components: List[np.ndarray], group: str) -> Tuple[Optional[bool], dict]:
    """将固定选项文字与手写勾/圈区分开；返回已选、待复核或未选。"""
    candidates = []
    for component in raw_components:
        area = int(component[cv2.CC_STAT_AREA])
        width = int(component[cv2.CC_STAT_WIDTH])
        height = int(component[cv2.CC_STAT_HEIGHT])
        if group == "time":
            strong = height >= 28
            weak = height >= 25
        else:
            strong = (width >= 29 and height >= 25) or height >= 32 or width >= 42
            weak = (width >= 26 and height >= 25) or height >= 28
        if strong:
            candidates.append((True, area, width, height))
        elif weak:
            candidates.append((None, area, width, height))
    if not candidates:
        return False, {"area": 0, "width": 0, "height": 0, "count": 0}
    state, area, width, height = max(candidates, key=lambda item: (item[0] is True, item[1], item[2] * item[3]))
    return state, {"area": area, "width": width, "height": height, "count": len(candidates)}


def detect_sunlight_questionnaire_checks(input_path: Path) -> Tuple[dict, List[dict]]:
    """第 3 页日照题：OpenCV 模板对齐后识别时段、皮肤暴露部位的手写勾选或圈选。"""
    if not SUNLIGHT_QUESTIONNAIRE_TEMPLATE.exists() and not SUNLIGHT_REAL_SCAN_TEMPLATE.exists():
        return {"warning": "sunlight_template_missing"}, []

    reference = (
        cv2.imread(str(SUNLIGHT_QUESTIONNAIRE_TEMPLATE), cv2.IMREAD_GRAYSCALE)
        if SUNLIGHT_QUESTIONNAIRE_TEMPLATE.exists()
        else None
    )
    real_scan_reference = (
        cv2.imread(str(SUNLIGHT_REAL_SCAN_TEMPLATE), cv2.IMREAD_GRAYSCALE)
        if SUNLIGHT_REAL_SCAN_TEMPLATE.exists()
        else None
    )
    if reference is None and real_scan_reference is None:
        return {"warning": "sunlight_template_unreadable"}, []

    template_variant = "legacy"
    if reference is not None:
        image = render_document_gray(input_path, (reference.shape[1], reference.shape[0]))
        aligned, alignment_confidence = align_questionnaire(image, reference)
    else:
        image = render_document_gray(input_path, (real_scan_reference.shape[1], real_scan_reference.shape[0]))
        reference = real_scan_reference
        aligned, alignment_confidence = align_questionnaire(image, reference)
        template_variant = "real_scan"

    # 旧模板已在历史问卷上验证稳定。只有旧模板明显对齐不佳，且实物扫描模板
    # 至少高出 0.05 时才切换，避免更新模板改变既有批次的检测结果。
    if template_variant == "legacy" and alignment_confidence < 0.45 and real_scan_reference is not None:
        real_scan_image = render_document_gray(
            input_path,
            (real_scan_reference.shape[1], real_scan_reference.shape[0]),
        )
        real_scan_aligned, real_scan_confidence = align_questionnaire(real_scan_image, real_scan_reference)
        if real_scan_confidence >= 0.35 and real_scan_confidence >= alignment_confidence + 0.05:
            reference = real_scan_reference
            aligned = real_scan_aligned
            alignment_confidence = real_scan_confidence
            template_variant = "real_scan"

    if alignment_confidence < 0.35:
        return {
            "warning": "sunlight_alignment_low",
            "alignment_confidence": round(alignment_confidence, 3),
            "template_variant": template_variant,
        }, []

    # 原图墨迹用于区分勾/圈与印刷文字；参考模板差分用于给人工复核提供可审计的新增笔迹分数。
    raw_ink = cv2.threshold(aligned, 205, 255, cv2.THRESH_BINARY_INV)[1]
    current_ink = cv2.threshold(aligned, 225, 255, cv2.THRESH_BINARY_INV)[1]
    reference_ink = cv2.threshold(reference, 235, 255, cv2.THRESH_BINARY_INV)[1]
    reference_ink = cv2.dilate(reference_ink, np.ones((17, 17), np.uint8), iterations=1)
    novel_ink = cv2.bitwise_and(current_ink, cv2.bitwise_not(reference_ink))
    novel_ink = cv2.morphologyEx(novel_ink, cv2.MORPH_OPEN, np.ones((2, 2), np.uint8))

    groups = (
        (
            "time",
            "晒太阳时段",
            False,
            (
                ("上午9点之前", (0.190, 0.225, 0.400, 0.290)),
                ("上午9点至下午3点", (0.360, 0.225, 0.680, 0.290)),
                ("下午3点之后", (0.620, 0.225, 0.900, 0.290)),
            ),
        ),
        (
            "body_parts",
            "皮肤暴露部位",
            True,
            (
                ("脸部", (0.180, 0.305, 0.320, 0.380)),
                ("颈部", (0.300, 0.305, 0.420, 0.380)),
                ("后背", (0.400, 0.305, 0.520, 0.380)),
                ("胳膊", (0.490, 0.305, 0.620, 0.380)),
                ("腿", (0.600, 0.305, 0.710, 0.380)),
                ("全身", (0.690, 0.305, 0.840, 0.380)),
            ),
        ),
    )
    height, width = reference.shape
    sunlight = {}
    visual_elements = []
    for source_key, field_label, multi_select, options in groups:
        selected, uncertain, measurements = [], [], []
        for option_label, (x_start, y_start, x_end, y_end) in options:
            bbox = [int(width * x_start), int(height * y_start), int(width * x_end), int(height * y_end)]
            measurement = _sunlight_option_measurement(raw_ink, novel_ink, bbox)
            state, component = _sunlight_mark_state(measurement["raw_components"], source_key)
            novel = measurement["novel"]
            if state is True:
                selected.append(option_label)
                confidence = min(0.98, 0.62 + component["height"] / 120 + novel["area"] / 800)
            elif state is None:
                uncertain.append(option_label)
                confidence = min(0.69, 0.38 + component["height"] / 100 + novel["area"] / 1000)
            else:
                confidence = 0.90
            measurements.append(novel["score"])
            visual_elements.append({
                "label": f"{field_label}-{option_label}",
                "question_label": field_label,
                "option_label": option_label,
                "multi_select": multi_select,
                "selected": state,
                "confidence": round(confidence * alignment_confidence, 3),
                "bbox": bbox,
                "score": novel["score"],
                "component": component,
                "raw_component": measurement["raw"],
                "method": "sunlight_template_aligned_handwritten_mark_v1",
                "template_variant": template_variant,
            })
        sunlight[source_key] = {"selected": selected, "uncertain": uncertain, "max_score": max(measurements, default=0)}
    return {
        "sunlight": sunlight,
        "alignment_confidence": round(alignment_confidence, 3),
        "method": "sunlight_template_aligned_handwritten_mark_v1",
        "template_variant": template_variant,
    }, visual_elements


# =========================================================
# Excel 构建
# =========================================================
def build_excel_from_ocr_outputs(markdown_text: str, json_text: str, excel_path: Path, enhanced_identity: Optional[dict] = None) -> dict:
    """
    根治方案：
    1. 优先 JSON block
    2. 不行再回退 markdown
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "问卷数据"

    used_mode = "markdown_fallback"
    block_count = 0

    blocks = extract_blocks_from_json_text(json_text)
    if blocks:
        process_blocks_to_sheet(blocks, ws, start_row=1)
        used_mode = "json_blocks"
        block_count = len(blocks)
    else:
        process_markdown_text_to_sheet(markdown_text, ws, start_row=1)

    set_sheet_style(ws)
    if enhanced_identity and enhanced_identity.get("selected_name"):
        identity_sheet = wb.create_sheet("增强识别信息")
        identity_sheet.append(["字段", "值", "识别方式"])
        identity_sheet.append([f"姓名：{enhanced_identity['selected_name']}", enhanced_identity["selected_name"], "整页二次 OCR / 多候选评分"])
        identity_sheet.column_dimensions["A"].width = 18
        identity_sheet.column_dimensions["B"].width = 24
        identity_sheet.column_dimensions["C"].width = 32
        for cell in identity_sheet[1]:
            cell.font = Font(bold=True)
    wb.save(excel_path)

    return {
        "used_mode": used_mode,
        "block_count": block_count
    }


def count_tables(markdown_text: str, json_text: str) -> int:
    """
    优先从 JSON block 统计 table 数量
    """
    blocks = extract_blocks_from_json_text(json_text)
    json_table_count = sum(1 for block in blocks if str(block.get("block_label", "")).lower() == "table")
    if json_table_count > 0:
        return json_table_count

    markdown_text = normalize_text(markdown_text)

    html_count = len(
        re.findall(r"<table.*?</table>", markdown_text, flags=re.DOTALL | re.IGNORECASE)
    )

    md_count = 0
    blocks_text = re.split(r"\n\s*\n", markdown_text)
    for block in blocks_text:
        lines = [line for line in block.split("\n") if line.strip()]
        if is_markdown_table_block(lines):
            md_count += 1

    return html_count + md_count

#LaTeX 清洗/还原函数
def clean_ocr_latex_text(text: str) -> str:
    if text is None:
        return ""

    text = str(text)

    # 去掉 markdown 数学公式包裹符号
    text = re.sub(r"\$\s*", "", text)
    text = re.sub(r"\s*\$", "", text)

    # \underline{\text{均匀}} -> 均匀
    text = re.sub(
        r"\\underline\s*\{\s*\\text\s*\{([^{}]*)\}\s*\}",
        r"\1",
        text
    )

    # \underline{56} -> 56
    text = re.sub(
        r"\\underline\s*\{([^{}]*)\}",
        r"\1",
        text
    )

    # \text{均匀} -> 均匀
    text = re.sub(
        r"\\text\s*\{([^{}]*)\}",
        r"\1",
        text
    )

    # \frac{38}{27} -> 38/27
    text = re.sub(
        r"\\frac\s*\{([^{}]*)\}\s*\{([^{}]*)\}",
        r"\1/\2",
        text
    )

    # \times -> ×
    text = text.replace(r"\times", "×")

    # 去掉 LaTeX 空格命令
    text = text.replace(r"\,", " ")
    text = text.replace(r"\ ", " ")

    # 处理多余空格
    text = re.sub(r"\s+", " ", text).strip()

    return text


# =========================================================
# API
# =========================================================
PROCESSING_MODES = {"fast", "accurate"}
DOCUMENT_KINDS = {"auto", "medical", "nutrition", "body_composition"}


def validate_parse_options(processing_mode: str, document_kind: str, page_index: int) -> Tuple[str, str, int]:
    """Normalize optional form fields while keeping legacy file-only requests valid."""
    normalized_mode = str(processing_mode or "accurate").strip().lower()
    normalized_kind = str(document_kind or "auto").strip().lower()
    if normalized_mode not in PROCESSING_MODES:
        raise HTTPException(status_code=400, detail="processing_mode must be fast or accurate")
    if normalized_kind not in DOCUMENT_KINDS:
        raise HTTPException(status_code=400, detail="document_kind must be auto, medical, nutrition, or body_composition")
    if page_index < 1:
        raise HTTPException(status_code=400, detail="page_index must be greater than or equal to 1")
    return normalized_mode, normalized_kind, page_index


@app.get("/")
def health():
    # 保持旧客户端所需的 status 字段；额外信息用于启动脚本和人工排查。
    return {
        "status": "ok",
        "api_version": "1.5-body-composition",
        "parse_file": "multipart/form-data: file 必填；processing_mode、document_kind、page_index、include_assets 可省略",
    }


@app.post("/classic-ocr")
def classic_ocr(file: UploadFile = File(...)):
    """Lightweight PP-OCRv5 endpoint for small handwritten questionnaire crops."""
    global classic_ocr_model
    suffix = Path(file.filename or "crop.png").suffix.lower()
    if suffix not in ALLOWED_EXTS - {".pdf"}:
        raise HTTPException(status_code=400, detail="classic-ocr only accepts image files")
    task_id = str(uuid.uuid4())
    input_path = UPLOAD_DIR / f"{task_id}{suffix}"
    with input_path.open("wb") as target:
        shutil.copyfileobj(file.file, target)
    request_started = perf_counter()
    try:
        with classic_ocr_lock:
            initialization_started = perf_counter()
            if classic_ocr_model is None:
                classic_ocr_model = PaddleOCR(
                    lang="ch",
                    ocr_version="PP-OCRv5",
                    device="cpu",
                    use_doc_orientation_classify=False,
                    use_doc_unwarping=False,
                    use_textline_orientation=False,
                    text_rec_score_thresh=0.0,
                )
            initialization_ms = round((perf_counter() - initialization_started) * 1000, 2)
            inference_started = perf_counter()
            results = list(classic_ocr_model.predict(str(input_path)))
            inference_ms = round((perf_counter() - inference_started) * 1000, 2)

        items = []
        for result in results:
            payload = getattr(result, "json", {})
            if callable(payload):
                payload = payload()
            data = payload.get("res", payload) if isinstance(payload, dict) else {}
            texts = data.get("rec_texts")
            scores = data.get("rec_scores")
            boxes = data.get("rec_boxes")
            texts = texts.tolist() if hasattr(texts, "tolist") else (texts or [])
            scores = scores.tolist() if hasattr(scores, "tolist") else (scores or [])
            boxes = boxes.tolist() if hasattr(boxes, "tolist") else (boxes or [])
            for index, text in enumerate(texts):
                items.append({
                    "text": str(text),
                    "score": round(float(scores[index]), 6) if index < len(scores) else 0.0,
                    "box": boxes[index] if index < len(boxes) else [],
                })
        return {
            "success": True,
            "task_id": task_id,
            "items": items,
            "text": "\n".join(item["text"] for item in items if item["text"]),
            "timings_ms": {
                "model_initialization": initialization_ms,
                "inference": inference_ms,
                "total": round((perf_counter() - request_started) * 1000, 2),
            },
            "api_version": "1.4-classic-crop",
        }
    except Exception as error:
        raise HTTPException(status_code=500, detail=str(error)) from error
    finally:
        input_path.unlink(missing_ok=True)


@app.post("/parse-file")
async def parse_file(
    file: UploadFile = File(...),
    processing_mode: str = Form("accurate"),
    document_kind: str = Form("auto"),
    page_index: int = Form(1),
    include_assets: bool = Form(True),
):
    request_started = perf_counter()
    timings_ms = {
        "upload_save": 0.0,
        "primary_ocr": 0.0,
        "ocr_output_save": 0.0,
        "output_collect": 0.0,
        "identity_ocr": 0.0,
        "header_identity_ocr": 0.0,
        "body_composition_header_ocr": 0.0,
        "checkbox_cv": 0.0,
        "debug_assets_write": 0.0,
        "excel_build": 0.0,
        "assets": 0.0,
        "response_encode": 0.0,
        "total": 0.0,
    }
    processing_mode, document_kind, page_index = validate_parse_options(
        processing_mode, document_kind, page_index
    )
    if not file.filename:
        raise HTTPException(status_code=400, detail="文件名为空")

    suffix = Path(file.filename).suffix.lower()
    if suffix not in ALLOWED_EXTS:
        raise HTTPException(
            status_code=400,
            detail="只支持 PDF、PNG、JPG、JPEG、BMP、WEBP 文件"
        )

    task_id = str(uuid.uuid4())
    input_path = UPLOAD_DIR / f"{task_id}{suffix}"
    task_output_dir = OUTPUT_DIR / task_id
    excel_path = task_output_dir / "result.xlsx"
    debug_path = task_output_dir / "debug.txt"
    merged_md_path = task_output_dir / "merged_markdown_debug.md"
    merged_json_path = task_output_dir / "merged_json_debug.json"

    task_output_dir.mkdir(parents=True, exist_ok=True)

    debug_lines = [
        f"task_id: {task_id}",
        f"filename: {file.filename}",
        f"suffix: {suffix}",
        f"processing_mode: {processing_mode}",
        f"document_kind: {document_kind}",
        f"page_index: {page_index}",
        f"include_assets: {include_assets}",
        f"input_path: {input_path}",
        f"task_output_dir: {task_output_dir}",
    ]

    try:
        # 1. 保存上传文件
        stage_started = perf_counter()
        with open(input_path, "wb") as f:
            shutil.copyfileobj(file.file, f)
        timings_ms["upload_save"] = round((perf_counter() - stage_started) * 1000, 2)

        # 2. OCR
        stage_started = perf_counter()
        try:
            results = pipeline.predict(str(input_path))
            try:
                result_list = list(results)
            except TypeError:
                result_list = [results]
        finally:
            # PaddleOCRVL may return a lazy iterator, so materialization is part of primary OCR time.
            timings_ms["primary_ocr"] = round((perf_counter() - stage_started) * 1000, 2)

        debug_lines.append(f"result_count: {len(result_list)}")

        stage_started = perf_counter()
        for idx, res in enumerate(result_list):
            debug_lines.append(f"\n===== result {idx} =====")
            debug_lines.append(f"type: {type(res)}")
            debug_lines.append(f"repr: {repr(res)}")

            if hasattr(res, "save_to_markdown"):
                try:
                    res.save_to_markdown(save_path=str(task_output_dir))
                    debug_lines.append("save_to_markdown: success")
                except Exception as e:
                    debug_lines.append(f"save_to_markdown: error -> {e}")

            if hasattr(res, "save_to_json"):
                try:
                    res.save_to_json(save_path=str(task_output_dir))
                    debug_lines.append("save_to_json: success")
                except Exception as e:
                    debug_lines.append(f"save_to_json: error -> {e}")
        timings_ms["ocr_output_save"] = round((perf_counter() - stage_started) * 1000, 2)

        # 3. 收集 md/json
        stage_started = perf_counter()
        markdown_text, json_text, md_files, json_files = collect_saved_outputs(task_output_dir)
        timings_ms["output_collect"] = round((perf_counter() - stage_started) * 1000, 2)

        # 4. 兜底
        if not markdown_text and json_text:
            markdown_text = json_text

        if not markdown_text and not json_text:
            markdown_text = "\n\n".join(str(res) for res in result_list)

        # 5. 身份信息增强：主版面结果缺姓名时，使用整页上下文二次 OCR。
        warnings = []
        enhanced_markdown = ""
        name_candidates = extract_identity_candidates(markdown_text, "layout_ocr")
        detected_food_page = "食物频率调查表" in markdown_text
        body_composition_page = page_index == 1 and document_kind == "body_composition"
        identity_relevant = (
            (page_index == 1 and document_kind in {"medical", "nutrition", "body_composition"})
            or any(keyword in markdown_text for keyword in ("姓名", "受试者", "食物频率调查表", "检验报告单"))
        )
        nutrition_first_page = page_index == 1 and (
            document_kind == "nutrition" or (document_kind == "auto" and detected_food_page)
        )
        full_identity_allowed = processing_mode == "accurate" or page_index == 1
        if not choose_identity_name(name_candidates) and identity_relevant and full_identity_allowed:
            stage_started = perf_counter()
            try:
                enhanced_markdown, enhanced_candidates = run_full_page_identity_ocr(input_path, task_output_dir)
                name_candidates.extend(enhanced_candidates)
                debug_lines.append(f"enhanced_identity_candidates: {enhanced_candidates}")
            except Exception as enhanced_error:
                warnings.append(f"整页身份信息增强失败：{enhanced_error}")
                debug_lines.append(f"enhanced_identity_error: {repr(enhanced_error)}")
            finally:
                timings_ms["identity_ocr"] = round((perf_counter() - stage_started) * 1000, 2)
        elif not choose_identity_name(name_candidates) and identity_relevant and processing_mode == "fast":
            debug_lines.append("enhanced_identity_skipped: fast mode only enhances page_index=1")

        if not choose_identity_name(name_candidates) and nutrition_first_page:
            stage_started = perf_counter()
            try:
                header_markdown, header_candidates = run_food_header_identity_ocr(input_path, task_output_dir)
                if header_markdown:
                    enhanced_markdown = "\n\n".join(item for item in (enhanced_markdown, header_markdown) if item)
                name_candidates.extend(header_candidates)
                debug_lines.append(f"header_identity_candidates: {header_candidates}")
            except Exception as header_error:
                warnings.append(f"问卷页眉姓名增强失败：{header_error}")
                debug_lines.append(f"header_identity_error: {repr(header_error)}")
            finally:
                timings_ms["header_identity_ocr"] = round((perf_counter() - stage_started) * 1000, 2)

        body_composition_header_markdown = ""
        if body_composition_page:
            stage_started = perf_counter()
            try:
                body_composition_header_markdown, body_header_candidates = run_body_composition_header_ocr(
                    input_path, task_output_dir
                )
                name_candidates.extend(body_header_candidates)
                if body_composition_header_markdown:
                    markdown_text = (
                        f"{markdown_text}\n\n<!-- enhanced_body_composition_header -->\n"
                        f"{body_composition_header_markdown}"
                    ).strip()
                    enhanced_markdown = "\n\n".join(
                        item for item in (enhanced_markdown, body_composition_header_markdown) if item
                    )
                debug_lines.append(f"body_composition_header_candidates: {body_header_candidates}")
            except Exception as body_header_error:
                warnings.append(f"体成分页眉增强失败：{body_header_error}")
                debug_lines.append(f"body_composition_header_error: {repr(body_header_error)}")
            finally:
                timings_ms["body_composition_header_ocr"] = round((perf_counter() - stage_started) * 1000, 2)

        selected_name = choose_identity_name(name_candidates)
        identity = {
            "selected_name": selected_name,
            "name_candidates": name_candidates,
            "enhancement_used": bool(enhanced_markdown),
        }
        if selected_name and selected_name not in markdown_text:
            markdown_text = f"{markdown_text}\n\n<!-- enhanced_identity -->\n姓名：{selected_name}".strip()

        questionnaire = {}
        visual_elements = []
        # 仅在固定版式的第 1、3 页运行对应 OpenCV 模板；续页食物表不使用第一页坐标。
        if nutrition_first_page:
            stage_started = perf_counter()
            try:
                questionnaire, visual_elements = detect_food_questionnaire_checks(input_path)
                if questionnaire.get("warning"):
                    warnings.append("食物问卷版式对齐置信度不足，勾选结果已保留为空")
                debug_lines.append(f"questionnaire: {questionnaire}")
            except Exception as checkbox_error:
                warnings.append(f"勾选增强失败：{checkbox_error}")
                debug_lines.append(f"checkbox_enhancement_error: {repr(checkbox_error)}")
            finally:
                timings_ms["checkbox_cv"] = round((perf_counter() - stage_started) * 1000, 2)
        elif page_index == 3 and document_kind == "nutrition":
            stage_started = perf_counter()
            try:
                questionnaire, visual_elements = detect_sunlight_questionnaire_checks(input_path)
                if questionnaire.get("warning"):
                    warnings.append("日照问卷版式对齐置信度不足，勾选结果已保留为空")
                debug_lines.append(f"sunlight_questionnaire: {questionnaire}")
            except Exception as sunlight_error:
                warnings.append(f"日照勾选增强失败：{sunlight_error}")
                debug_lines.append(f"sunlight_checkbox_error: {repr(sunlight_error)}")
            finally:
                timings_ms["checkbox_cv"] = round((perf_counter() - stage_started) * 1000, 2)

        structured = {
            "identity": identity,
            "questionnaire": questionnaire,
            "body_composition_header": body_composition_header_markdown,
        }

        # 6. 调试落盘
        stage_started = perf_counter()
        merged_md_path.write_text(markdown_text if markdown_text else "", encoding="utf-8")
        merged_json_path.write_text(json_text if json_text else "{}", encoding="utf-8")
        timings_ms["debug_assets_write"] = round((perf_counter() - stage_started) * 1000, 2)

        # 7. 生成 Excel（主表保持原规则，增强姓名写入独立工作表）
        stage_started = perf_counter()
        build_info = build_excel_from_ocr_outputs(markdown_text, json_text, excel_path, identity)
        timings_ms["excel_build"] = round((perf_counter() - stage_started) * 1000, 2)
        table_count = count_tables(markdown_text, json_text)
        stage_started = perf_counter()
        assets = collect_image_assets(task_output_dir) if include_assets else []
        timings_ms["assets"] = round((perf_counter() - stage_started) * 1000, 2)

        debug_lines.append(f"md_files: {md_files}")
        debug_lines.append(f"json_files: {json_files}")
        debug_lines.append(f"table_count: {table_count}")
        debug_lines.append(f"excel_build_mode: {build_info['used_mode']}")
        debug_lines.append(f"json_block_count: {build_info['block_count']}")
        debug_lines.append(f"excel_path: {excel_path}")
        debug_lines.append(f"merged_md_path: {merged_md_path}")
        debug_lines.append(f"merged_json_path: {merged_json_path}")
        debug_lines.append(f"assets_count: {len(assets)}")
        debug_lines.append(f"warnings: {warnings}")

        # 8. 返回。以下旧字段及类型保持不变；新增字段可被旧客户端安全忽略。
        stage_started = perf_counter()
        excel_b64 = base64.b64encode(excel_path.read_bytes()).decode("utf-8")
        timings_ms["response_encode"] = round((perf_counter() - stage_started) * 1000, 2)
        timings_ms["total"] = round((perf_counter() - request_started) * 1000, 2)
        debug_lines.append(f"timings_ms: {timings_ms}")
        debug_path.write_text("\n".join(debug_lines), encoding="utf-8")
        timings_ms["total"] = round((perf_counter() - request_started) * 1000, 2)

        return {
            "success": True,
            "task_id": task_id,
            "filename": file.filename,
            "file_type": suffix,
            "markdown": markdown_text,
            "json_text": json_text,
            "excel": excel_b64,
            "table_count": table_count,
            "md_files": md_files,
            "json_files": json_files,
            "excel_build_mode": build_info["used_mode"],
            "json_block_count": build_info["block_count"],
            "debug_path": str(debug_path),
            "merged_md_path": str(merged_md_path),
            "merged_json_path": str(merged_json_path),
            "api_version": "1.5-body-composition-compatible",
            "enhanced_markdown": enhanced_markdown,
            "structured": structured,
            "assets": assets,
            "visual_elements": visual_elements,
            "warnings": warnings,
            "processing_mode": processing_mode,
            "document_kind": document_kind,
            "page_index": page_index,
            "include_assets": include_assets,
            "timings_ms": timings_ms,
        }

    except Exception as e:
        timings_ms["total"] = round((perf_counter() - request_started) * 1000, 2)
        debug_lines.append(f"timings_ms: {timings_ms}")
        debug_lines.append(f"error: {repr(e)}")
        debug_path.write_text("\n".join(debug_lines), encoding="utf-8")
        raise HTTPException(status_code=500, detail=str(e))
